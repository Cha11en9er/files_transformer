"""Smoke checks for known bug scenarios."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.api.routes.shipments import _looks_like_18233
from app.parsing.pipeline import parse_file
from app.services.export_18233_templates import export_18233_from_templates
from app.services.materials_18233 import kit_sources
from app.services.profile_18233 import parse_bundle, reconcile_18233
from app.services.reconcile import items_to_dicts


def test_spec_alone_has_articles() -> None:
    src = kit_sources("626-1")["specification"]
    bundle = parse_bundle([src])
    items = [i for i in reconcile_18233(bundle) if i.article]
    assert len(items) == 13
    assert all(i.article for i in items)


def test_generic_reads_product_name() -> None:
    src = kit_sources("626-1")["specification"]
    doc = parse_file(str(src), allow_ocr=False)
    assert doc.lines
    assert sum(1 for line in doc.lines if line.normalized_article) >= 250


def test_auto_detect_18233_filenames() -> None:
    paths = list(kit_sources("626-1").values())
    assert _looks_like_18233(paths) is True


def test_ready_etalon_filename_detected() -> None:
    from app.api.routes.shipments import _is_ready_etalon_filename

    assert _is_ready_etalon_filename("18233 626-1 СПЕЦИФИКАЦИЯ С ЦВЕТАМИ.xlsx") is True
    assert _is_ready_etalon_filename("18233 ИНВОЙС 626-1.XLSX") is True
    assert _is_ready_etalon_filename("626-1-YS-RMB-EXW-Specification.xls") is False
    assert _is_ready_etalon_filename("626-1-YS-RMB-EXW-INVOICE.XLSX") is False


def test_cyrillic_role_classification() -> None:
    from app.api.routes.shipments import _classify_18233_filename

    assert _classify_18233_filename("18233 ИНВОЙС 626-1.XLSX") == "invoice"
    assert _classify_18233_filename("18233 ПАКИНГ 626-1.XLSX") == "packing"
    assert _classify_18233_filename("18233 626-1 СПЕЦИФИКАЦИЯ С ЦВЕТАМИ.xlsx") == "specification"


def test_beijing_profile_still_forced_to_18233_for_zhongfang(tmp_path: Path) -> None:
    # Simulate wrong UI profile: files are 18233 kit
    paths = list(kit_sources("626-1").values())
    assert _looks_like_18233(paths)
    items = items_to_dicts(reconcile_18233(parse_bundle(paths)))
    products = [
        i
        for i in items
        if (i.get("source_traces") or {}).get("invoice")
        or (i.get("source_traces") or {}).get("specification")
    ]
    assert len(products) == 13
    export_18233_from_templates(items, tmp_path / "out", header={"invoice_no": "ZFRMB26148-626-1"})
    assert list((tmp_path / "out").glob("*.xlsx"))


def test_6211_uses_invoice_amount_and_packing_weights() -> None:
    kit = (
        Path(__file__).resolve().parents[2]
        / "documents"
        / "18080 ЛЮ 621 ТМЛ"
        / "Исходные"
    )
    invoice = kit / "621-1-YS-RMB-EXW-INVOICE.XLSX"
    packing = kit / "621-1-YS-RMB-EXW-PL.XLSX"
    spec = kit / "621-1-YS-RMB-EXW-Specification.xls"
    if not invoice.exists():
        return
    items = reconcile_18233(parse_bundle([invoice, packing, spec]))
    products = [i for i in items if i.article]
    sherlock = next(i for i in products if "SHERLOCK" in (i.normalized_article or ""))
    assert abs(float(sherlock.commercial_data["price"]) - 10.4) < 0.01
    assert abs(float(sherlock.commercial_data["amount"]) - 19624.8) < 0.05
    assert abs(float(sherlock.packing_data["net_weight"]) - 811.41) < 0.02
    assert abs(float(sherlock.packing_data["gross_weight"]) - 851) < 0.02
    noble = next(i for i in products if i.normalized_article == "NOBLE110")
    assert abs(float(noble.commercial_data["amount"]) - 24598.08) < 0.05
    assert abs(float(noble.packing_data["net_weight"]) - 533.52) < 0.05
    pl = sherlock.source_traces.get("packing_list_group") or {}
    assert "SOFA FABRIC" in str(pl.get("design") or "").upper()
    assert "SHERLOCK" in str(pl.get("design") or "").upper()
    assert "(15+30)" not in str(pl.get("design") or "")


def test_6211_export_uses_etalon_columns_and_totals(tmp_path: Path) -> None:
    kit = (
        Path(__file__).resolve().parents[2]
        / "documents"
        / "18080 ЛЮ 621 ТМЛ"
        / "Исходные"
    )
    invoice = kit / "621-1-YS-RMB-EXW-INVOICE.XLSX"
    packing = kit / "621-1-YS-RMB-EXW-PL.XLSX"
    spec = kit / "621-1-YS-RMB-EXW-Specification.xls"
    if not invoice.exists():
        return
    items = items_to_dicts(reconcile_18233(parse_bundle([invoice, packing, spec])))
    header = {"invoice_no": "ZFRMB26125-621-1", "invoice_date": "Jul.18,2026"}
    paths = export_18233_from_templates(items, tmp_path / "out", header=header, shipment_title="18080")
    inv = next(p for p in paths if "ИНВОЙС" in p.name.upper())
    pl = next(p for p in paths if "ПАКИНГ" in p.name.upper())
    sp = next(p for p in paths if "СПЕЦИФ" in p.name.upper())
    from openpyxl import load_workbook

    inv_ws = load_workbook(inv).active
    headers = [str(inv_ws.cell(24, c).value or "") for c in range(1, 11)]
    assert headers[1].upper() == "DESIGN"
    assert "ROLLS" in headers[3].upper()
    assert "TOTAL M2" in headers[5].upper()
    letterhead = " ".join(str(inv_ws.cell(r, 1).value or "") for r in range(1, 23))
    assert "INVOICE" in letterhead.upper()
    assert "ZFRMB26125-621-1" in " ".join(str(inv_ws.cell(r, c).value or "") for r in range(1, 23) for c in range(1, 11))
    designs = [str(inv_ws.cell(r, 2).value or "") for r in range(25, inv_ws.max_row + 1)]
    assert any("Sherlock 980" in d for d in designs)
    assert any("SOFA FABRIC" in d and "Sherlock" in d for d in designs)
    assert any(d.upper().startswith("TOTAL") for d in designs)
    amounts = [inv_ws.cell(r, 10).value for r in range(25, inv_ws.max_row + 1)]
    assert any(abs(float(a) - 19624.8) < 0.05 for a in amounts if a not in (None, ""))

    pl_ws = load_workbook(pl).active
    pl_headers = [str(pl_ws.cell(9, c).value or "") for c in range(1, 9)]
    assert pl_headers[1].upper() == "DESIGN"
    assert "G/M" in pl_headers[2].upper()
    pl_designs = [str(pl_ws.cell(r, 2).value or "") for r in range(10, pl_ws.max_row + 1)]
    assert any("SOFA FABRIC" in d and "Sherlock" in d for d in pl_designs)
    assert not any("(15+30)" in d for d in pl_designs)
    assert any(str(pl_ws.cell(r, 1).value or "").upper() == "TOTAL" for r in range(10, pl_ws.max_row + 1))
    pack_letter = " ".join(str(pl_ws.cell(r, 1).value or "") for r in range(1, 9))
    assert "PACKING LIST" in pack_letter.upper()

    spec_ws = load_workbook(sp)[load_workbook(sp).sheetnames[0]]
    spec_blob = " ".join(str(spec_ws.cell(r, c).value or "") for r in range(1, 23) for c in range(1, 16))
    assert "SPECIFICATION" in spec_blob.upper()
    assert "ZFRMB26125-621-1" in spec_blob
    spec_headers = [str(spec_ws.cell(22, c).value or "") for c in range(2, 12)]
    assert any("Widht" in h or "Ширина" in h for h in spec_headers)
    assert any("Q-ty m2" in h or "кв.м" in h for h in spec_headers)
    assert any("total" in str(spec_ws.cell(r, 2).value or "").lower() for r in range(23, spec_ws.max_row + 1))
