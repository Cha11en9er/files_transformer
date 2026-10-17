"""Новое чтение на сайте: столбцы таблицы и справочник вне лотов."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.opencode_review import _model_parts
from app.services.export import write_recognition_book
from app.services.prepare_site import (
    apply_verdict,
    assess_draft,
    fee_rows,
    filled_base_columns,
    header_changes,
    header_diff,
    header_from_draft,
    is_reference_name,
    lots_to_rows,
    needs_second_model,
    overlay_model_header,
    review_files,
    short_prompt,
    split_uploads,
    verdict_notes,
    vision_paths,
)


def _lot(**fields):
    base = {"freight": False, "vendor": "", "description": "", "pieces": None, "packages": None}
    base.update(fields)
    return base


def test_partial_column_counts_as_filled() -> None:
    lots = [
        _lot(vendor="A", pieces=1, price=2, amount=2, hs="111"),
        _lot(vendor="B", pieces=3),
    ]
    # артикул, количество, цена, сумма, HS и тот же код в ТН ВЭД
    assert filled_base_columns(lots) == 6
    assert needs_second_model(lots) is False


def test_four_filled_columns_call_the_second_model() -> None:
    lots = [_lot(vendor="A", pieces=10, price=1, amount=10)]
    assert filled_base_columns(lots) == 4
    assert needs_second_model(lots) is True


def test_empty_cells_do_not_fill_a_column() -> None:
    lots = [_lot(vendor="A", description="   ", pieces=None, price="")]
    assert filled_base_columns(lots) == 1


def test_verdict_keeps_draft_numbers() -> None:
    draft = [_lot(vendor="A", pieces=10, price=1.5, amount=15)]
    payload = {"lots": [{"index": 0, "action": "keep", "fields": {}}]}
    out = apply_verdict(draft, payload)
    assert out[0]["pieces"] == 10
    assert out[0]["price"] == 1.5


def test_fix_does_not_wipe_a_number_with_empty() -> None:
    draft = [_lot(vendor="A", pieces=10, price=1.5)]
    payload = {"lots": [{"index": 0, "action": "fix", "fields": {"price": "", "pieces": 12}}]}
    out = apply_verdict(draft, {})
    assert out[0]["pieces"] == 10
    out = apply_verdict(draft, payload)
    assert out[0]["pieces"] == 12
    assert out[0]["price"] == 1.5


def test_grok_variant_is_not_part_of_the_model_id() -> None:
    provider, model_id, variant = _model_parts("openrouter/x-ai/grok-4.7:high")
    assert provider == "openrouter"
    assert model_id == "x-ai/grok-4.7"
    assert variant == "high"
    provider, model_id, variant = _model_parts("openrouter/x-ai/grok-4.7")
    assert model_id == "x-ai/grok-4.7"
    assert variant is None


def test_hs_is_digits_without_dots() -> None:
    rows = lots_to_rows(
        [_lot(vendor="BLOOM", hs="54.07.73.00.90.11", hs_alt="5407699000", pieces=49)],
        ["hs_conflict"],
    )
    customs = rows[0]["customs_data"]
    assert customs["hs_code"] == "540773009011"
    assert customs["tnved_code"] == "5407699000"
    assert "." not in customs["hs_code"]
    assert rows[0]["validation_errors"] == []


def test_missing_currency_does_not_become_yuan() -> None:
    from app.parsing.header_extract import export_currency_label, export_currency_word_ru

    assert export_currency_label(None, hangzhou_style=True) == ""
    assert export_currency_label("USD", hangzhou_style=True) == "USD"
    assert export_currency_label("CNY", hangzhou_style=True) == "RMB"
    assert export_currency_word_ru(None) == ""
    assert export_currency_word_ru("USD") == "долл. США"


def test_different_factories_are_not_joined_into_the_header() -> None:
    header = header_from_draft({}, [_lot(producer="Factory One"), _lot(producer="Factory Two")])
    assert header["manufacturer"] == ""
    single = header_from_draft({}, [_lot(producer="Only Plant")])
    assert single["manufacturer"] == "Only Plant"
    listed = header_from_draft(
        {"consignee": "Master LLC", "consignee_address": "Zastavskaya 22"},
        [_lot(producer="Utmaster Co., Ltd / ZHEJIANG WOXIN / CIXI YUXIAO / Anhui Jianghuai")],
    )
    assert listed["manufacturer"] == ""
    assert listed["consignee"] == "Master LLC"
    assert listed["consignee_address"] == "Zastavskaya 22"


def test_color_brand_and_size_reach_the_row() -> None:
    row = lots_to_rows(
        [_lot(finish="Matte black", brand="Micama", size="180-185, XL:2", package_type="мешок", producer="Plant A")],
        [],
    )[0]
    assert row["commercial_data"]["color"] == "Matte black"
    assert row["commercial_data"]["size"] == "180-185, XL:2"
    assert row["customs_data"]["brand"] == "Micama"
    assert row["customs_data"]["manufacturer"] == "Plant A"
    assert row["packing_data"]["package_type"] == "мешок"


def test_draft_currency_reaches_the_header() -> None:
    # Валюта у колонки цены — USD. Она не должна теряться перед экспортом (иначе профиль ставит RMB).
    draft = {"currency": "USD", "seller": "ACME", "contract": "C-1"}
    header = header_from_draft(draft, [])
    assert header["currency"] == "USD"


def test_model_can_correct_the_currency() -> None:
    header = {"currency": "CNY"}
    payload = {"header": {"currency": "USD"}}
    out = overlay_model_header(header, payload)
    assert out["currency"] == "USD"


def test_empty_model_currency_does_not_wipe_the_draft() -> None:
    header = {"currency": "USD"}
    payload = {"header": {"currency": ""}}
    out = overlay_model_header(header, payload)
    assert out["currency"] == "USD"


def test_model_can_set_consignee_apart_from_buyer() -> None:
    header = {"buyer": "Logiya DV LLC", "consignee": "Logiya DV LLC"}
    payload = {"header": {"consignee": "Master LLC", "consignee_address": "Zastavskaya 22"}}
    out = overlay_model_header(header, payload)
    assert out["consignee"] == "Master LLC"
    assert out["consignee_address"] == "Zastavskaya 22"
    assert out["buyer"] == "Logiya DV LLC"


def test_reference_file_stays_out_of_goods() -> None:
    saved = [
        ("c:/in/invoice.pdf", "invoice.pdf"),
        ("c:/in/сводная.xlsx", "сводная.xlsx"),
        ("c:/in/letter.pdf", "letter.pdf"),
    ]
    goods, references = split_uploads(saved, {"letter.pdf"})
    assert [name for _path, name in goods] == ["invoice.pdf"]
    assert {name for _path, name in references} == {"сводная.xlsx", "letter.pdf"}
    assert is_reference_name("справочник.xlsx", set())


def test_pallets_and_package_type_stay_apart_from_places() -> None:
    lot = _lot(vendor="A1", packages=2150, package_type="carton box", pallet_count=20, gross=21750.0, gross_with_pallet=22000.0, pallet_weight=250.0)
    row = lots_to_rows([lot], [])[0]
    assert row["packing_data"]["rolls"] == 2150
    assert row["packing_data"]["package_type"] == "carton box"
    assert row["packing_data"]["pallets"] == 20
    assert row["packing_data"]["gross_weight_with_pallet"] == 22000.0
    assert any(flag["field_name"] == "gross_weight" for flag in row["validation_errors"])


def test_shipper_code_goes_to_hs_and_tnved_stays_main() -> None:
    lot = _lot(vendor="A1", hs="9401908009", hs_alt="9401909090", hs_alt_shipper=True)
    customs = lots_to_rows([lot], [])[0]["customs_data"]
    assert customs["tnved_code"] == "9401908009"
    assert customs["hs_code"] == "9401909090"
    plain = _lot(vendor="A2", hs="5903101000", hs_alt="590310909000")
    customs = lots_to_rows([plain], [])[0]["customs_data"]
    assert customs["hs_code"] == "5903101000"
    assert customs["tnved_code"] == "590310909000"


def test_fees_become_rows_without_quantity() -> None:
    rows = fee_rows([{"description": "Packing fee", "amount": 1084.8, "freight": True}, {"description": "no sum"}])
    assert len(rows) == 1
    assert rows[0]["commercial_data"] == {"amount": 1084.8}
    assert "qty" not in rows[0]["commercial_data"]


def test_model_changes_and_drops_are_visible() -> None:
    draft = [_lot(vendor="A1", pieces=10, net=4.0), _lot(vendor="B2", pieces=5)]
    payload = {
        "lots": [
            {"index": 0, "action": "fix", "fields": {"net": 40.0}, "reason": "на фото 40"},
            {"index": 1, "action": "drop", "reason": "это итог"},
        ]
    }
    out = apply_verdict(draft, payload)
    assert len(out) == 1
    flags = lots_to_rows(out, [])[0]["validation_errors"]
    assert any(flag["field_name"] == "net_weight" and "было 4" in flag["message"] for flag in flags)
    notes = verdict_notes(draft, payload)
    assert len(notes) == 1 and "B2" in notes[0]


def test_conflicting_number_of_another_document_is_flagged_not_chosen() -> None:
    lot = _lot(vendor="A1", net=100.0, conflicts={"net": 120.0})
    row = lots_to_rows([lot], [])[0]
    assert row["packing_data"]["net_weight"] == 100.0
    assert any("другое число" in flag["message"] for flag in row["validation_errors"])


def test_model_header_replaces_draft_and_change_is_reported() -> None:
    header = {"invoice_date": "25.02.2014", "contract_date": "", "buyer_address": "Moscow"}
    payload = {"header": {"invoice_date": "31.03.2014", "contract_date": "25.02.2014", "buyer_address": "moscow"}}
    out = overlay_model_header(header, payload)
    assert out["invoice_date"] == "31.03.2014"
    assert out["contract_date"] == "25.02.2014"
    notes = header_changes(header, out)
    assert len(notes) == 1 and "дата инвойса" in notes[0]


def test_weight_sold_goods_get_quantity_from_net() -> None:
    lot = _lot(vendor="A1", unit="kg", price=0.85, amount=17850.0, net=21000.0, packages=2150)
    row = lots_to_rows([lot], [])[0]
    assert row["commercial_data"]["qty"] == 21000.0
    other = _lot(vendor="A2", unit="kg", price=0.85, amount=999.0, net=21000.0)
    assert "qty" not in lots_to_rows([other], [])[0]["commercial_data"]


def test_cosmetic_header_differences_are_not_changes() -> None:
    header = {
        "buyer": "\u201cSM REGIONTEKSTIL\u201d LLC",
        "invoice_date": "13/01/2026",
        "contract_no": "NE\u0421-01/10",
        "delivery_terms": "EXW ISTANBUL//Turkey",
        "currency": "RMB",
        "contract_date": "2018-05-23",
    }
    payload = {
        "header": {
            "buyer": '"SM REGIONTEKSTIL" LLC',
            "invoice_date": "13.01.2026",
            "contract": "NEC-01/10",
            "delivery": "EXW ISTANBUL/Turkey",
            "currency": "CNY",
            "contract_date": "23.05.2018",
        }
    }
    out = overlay_model_header(header, payload)
    assert header_diff(header, out) == []
    assert out["contract_no"] == "NE\u0421-01/10"
    assert out["invoice_date"] == "13/01/2026"


def test_real_header_change_is_listed_with_before_and_after() -> None:
    header = {"seller_address": "Hadimkoy Mah. No:7", "currency": ""}
    payload = {"header": {"seller_address": "Hadimkoy Mah. No:7 Arnavutkoy-Istanbul", "currency": "USD"}}
    out = overlay_model_header(header, payload)
    diff = {item["field"]: item for item in header_diff(header, out)}
    assert diff["seller_address"]["kind"] == "replaced"
    assert diff["seller_address"]["before"] == "Hadimkoy Mah. No:7"
    assert diff["currency"]["kind"] == "filled"
    # В строки сообщений попадают только замены чужого значения.
    assert len(header_changes(header, out)) == 1


def test_printed_currency_goes_to_header() -> None:
    assert header_from_draft({"currency": "CNY", "currency_printed": "RMB"}, [])["currency"] == "RMB"


def test_price_list_filename_is_a_catalog() -> None:
    assert is_reference_name("Price list.pdf", set())
    assert is_reference_name("прайс механизмов.pdf", set())
    assert not is_reference_name("Invoice.pdf", set())


def test_review_tab_shows_rows_of_its_own_file() -> None:
    tables = {
        "pack.xlsx / Page1": {"rows": [{"article": "A", "rolls": 1}], "note": "", "text": "", "total_rows": 1},
        "scan.jpg": {"rows": [], "note": "Код не прочитал этот файл", "text": "сырой текст", "total_rows": 0},
    }
    docs = [
        {"name": "pack.xlsx / Page1", "role": "packing"},
        {"name": "scan.jpg", "role": "image"},
    ]
    context = review_files(docs, [{"article": "ALL"}], tables)
    by_name = {entry["filename"]: entry for entry in context["excel"] + context["pdfs"]}
    assert by_name["pack.xlsx / Page1"]["table"] == [{"article": "A", "rolls": 1}]
    assert by_name["pack.xlsx / Page1"]["kind"] == "excel"
    assert "scan.jpg" not in by_name
    assert "сырой текст" not in str(context)


def test_model_european_decimals_are_parsed() -> None:
    lots = [_lot(vendor="A1", pieces=10)]
    payload = {
        "lots": [
            {
                "index": 0,
                "action": "fix",
                "fields": {"net": "6,17", "gross": "6,50", "price": "3,85", "amount": "793,10"},
            }
        ]
    }
    out = apply_verdict(lots, payload)
    assert out[0]["net"] == 6.17
    assert out[0]["gross"] == 6.5
    assert out[0]["price"] == 3.85
    assert out[0]["amount"] == 793.1
    rows = lots_to_rows(out, [])
    assert rows[0]["packing_data"]["net_weight"] == 6.17


def test_model_add_with_comma_decimals_builds_rows() -> None:
    payload = {
        "lots": [
            {
                "action": "add",
                "fields": {
                    "vendor": "DYER 789",
                    "pieces": "206",
                    "packages": "5",
                    "price": "3,85",
                    "amount": "793,10",
                    "net": "64,50",
                    "gross": "67,95",
                    "unit": "meters",
                },
            }
        ]
    }
    out = apply_verdict([], payload)
    rows = lots_to_rows(out, [])
    assert len(rows) == 1
    assert rows[0]["commercial_data"]["price"] == 3.85
    assert rows[0]["packing_data"]["net_weight"] == 64.5


def test_dt_sheet_follows_the_scheme_and_leaves_gaps_empty() -> None:
    from app.services.export_tsd import _dt_key, dt_headers, dt_rows

    lots = [
        _lot(
            vendor="A1",
            model="M-1",
            description="Long name",
            producer="Factory One",
            brand="Micama",
            pieces=10,
            unit="pcs",
            price=2,
            amount=20,
            packages=3,
            package_type="картонная коробка",
            pallet_count=2,
            net=4,
            gross=5,
            hs="1111111111",
            origin="CN",
            size="180-185, XL:2",
        ),
        _lot(
            vendor="A2",
            description="Other",
            producer="Factory Two",
            pieces=1,
            unit="set",
            price=4,
            amount=4,
            packages=1,
            package_type="мешок",
            net=1,
            gross=1.5,
            hs="2222222222",
        ),
        _lot(vendor="A3", pieces=1, packages=1),
    ]
    items = lots_to_rows(lots, [])
    headers = dt_headers()
    rows = dt_rows(
        items,
        {
            "invoice_no": "INV-9",
            "invoice_date": "2026-01-05",
            "currency": "USD",
            "container": "CONT1",
            "manufacturer": "Factory One / Factory Two",
        },
    )
    keys = [_dt_key(title) for title in headers]
    assert keys[:6] == ["№", "Код ТН ВЭД", "Описание", "Описание в группе", "Изготовитель", "ТЗ"]
    assert keys.count("Описание") == 2
    assert keys.count("") == 3
    assert "ИТС по запросу" in keys and "Декларация соответствия" in keys
    assert "CT" not in keys
    body, total = rows[:-1], rows[-1]

    def col(name, row):
        return row[keys.index(name)]

    assert [col("Изготовитель", row) for row in body] == ["Factory One", "Factory Two", None]
    assert col("Код упаковки", body[0]) == "картонная коробка"
    assert col("Код упаковки", body[1]) == "мешок"
    assert all(col("Кол-во упак", row) != "CT" and col("Код упаковки", row) != "CT" for row in body)
    assert col("кол-во мест", body[0]) == col("Кол-во упак", body[0]) == 3
    assert col("Описание в группе", body[0]) is None
    assert col("Описание", body[0]) == "Long name"
    second_description = [i for i, name in enumerate(keys) if name == "Описание"][1]
    assert body[0][second_description] is None
    assert col("ТЗ", body[0]) == "Micama"
    assert col("марка", body[0]) is None
    assert col("модель", body[0]) == "M-1"
    assert col("модель", body[1]) is None
    assert col("поддоны кол-во", body[0]) == 2
    assert col("поддоны кол-во", body[1]) is None
    assert col("серийный номер", body[0]) is None
    assert col("номер инвойса", body[0]) is None
    assert col("дата инвойса", body[0]) is None
    assert col("Валюта инвойса", body[0]) is None
    assert col("контейнер", body[0]) is None
    assert col("номера контейнеры", body[0]) is None
    assert col("нетто без упаковки", body[0]) is None
    assert col("доп.единица", body[0]) is None
    assert col("Код наличия упаковки", body[0]) is None
    assert col("ИТС по запросу", body[0]) is None
    assert col("ИТС цена/нетто", body[0]) is None
    assert col("Декларация соответствия", body[0]) is None
    assert col("кол-во товара", total) == 12
    assert col("кол-во мест", total) == 5
    assert col("нетто", total) == 5
    assert col("брутто", total) == 6.5
    assert col("ст-ть товара", total) == 24
    assert col("цена за ед. товара", total) is None
    assert col("артикул", total) is None
    assert col("Код ТН ВЭД", total) is None
    assert col("Описание", total) is None
    assert col("страна происхождения", total) is None
    assert "Factory One / Factory Two" not in [col("Изготовитель", row) for row in rows]


def _draft_with(**extra):
    lot = _lot(vendor="A", pieces=2, packages=1, price=3, amount=6, net=1, gross=1.2, hs="11111111")
    draft = {
        "lots": [lot],
        "flags": [],
        "stated": {"pieces": 2, "amount": 6},
        "documents": [
            {
                "name": "inv.xlsx / Sheet",
                "role": "invoice",
                "line_count": 1,
                "readable": True,
                "stated": {"pieces": 2, "amount": 6},
            }
        ],
        "document_tables": {
            "inv.xlsx / Sheet": {"rows": [{"article": "A", "qty": 2, "amount": 6, "rolls": 1}]}
        },
    }
    draft.update(extra)
    return draft


def test_known_form_skips_the_model() -> None:
    route = assess_draft(_draft_with())
    assert route["path"] == "known"
    assert route["leftovers"] == []
    assert route["checks"][0]["gaps"] == []


def test_gap_against_printed_total_is_unknown() -> None:
    draft = _draft_with()
    draft["documents"][0]["stated"] = {"pieces": 9, "amount": 6}
    route = assess_draft(draft)
    assert route["path"] == "unknown"
    assert route["checks"][0]["gaps"][0]["field"] == "pieces"
    assert route["checks"][0]["gaps"][0]["printed"] == 9
    assert route["checks"][0]["gaps"][0]["rows"] == 2


def test_sheet_that_did_not_enter_lots_is_kept() -> None:
    draft = _draft_with()
    draft["documents"].append(
        {"name": "other.xlsx / Sheet", "role": "invoice", "line_count": 1, "readable": True, "stated": {}}
    )
    draft["document_tables"]["other.xlsx / Sheet"] = {"rows": [{"article": "B", "qty": 4, "amount": 8}]}
    route = assess_draft(draft)
    assert route["path"] == "unknown"
    assert route["leftovers"][0]["name"] == "other.xlsx / Sheet"
    assert route["leftovers"][0]["blocking"] is True


def test_scan_without_text_is_unknown() -> None:
    draft = _draft_with()
    draft["documents"].append({"name": "scan.pdf", "role": "scan", "line_count": 0, "readable": False, "stated": {}})
    route = assess_draft(draft)
    assert route["path"] == "unknown"
    assert route["unreadable"][0]["name"] == "scan.pdf"


def test_sheet_without_printed_total_is_similar() -> None:
    draft = _draft_with()
    draft["documents"][0]["stated"] = {}
    draft["stated"] = {}
    route = assess_draft(draft)
    assert route["path"] == "similar"


def test_quiet_draft_sheet_does_not_call_the_model() -> None:
    draft = _draft_with()
    draft["documents"].append(
        {"name": "calc.xlsx / Sheet", "role": "draft", "line_count": 1, "readable": True, "stated": {}}
    )
    draft["document_tables"]["calc.xlsx / Sheet"] = {"rows": [{"article": "Z", "qty": 1, "amount": 1}]}
    route = assess_draft(draft)
    assert route["path"] == "known"
    assert route["leftovers"][0]["blocking"] is False


def test_short_prompt_has_no_folder_names_and_keeps_both_numbers() -> None:
    draft = _draft_with()
    draft["sheet_checks"] = [{"name": "inv.xlsx / Sheet", "gaps": [{"field": "pieces", "rows": 2, "printed": 9}]}]
    draft["leftovers"] = [{"name": "other.xlsx / Sheet", "why": "строки с количеством не вошли в лоты"}]
    text = short_prompt(draft)
    assert "Шаблон" not in text
    assert "не выбирай" in text
    assert "leftovers" in text
    assert "9" in text


def test_vision_paths_keep_only_the_problem_file() -> None:
    reading = [(Path("inv.xlsx"), "inv.xlsx"), (Path("scan.pdf"), "scan.pdf")]
    assessment = {"unreadable": [{"name": "scan.pdf", "blocking": True}], "leftovers": [], "checks": []}
    picked = vision_paths(reading, assessment)
    assert [path.name for path in picked] == ["scan.pdf"]


def test_check_book_writes_both_numbers(tmp_path: Path) -> None:
    target = tmp_path / "ПРОВЕРКА.xlsx"
    saved = write_recognition_book(
        target,
        {
            "route": "unknown",
            "reasons": ["сумма строк не сходится с итогом листа"],
            "checks": [
                {
                    "name": "inv.xlsx / Sheet",
                    "role": "invoice",
                    "sums": {"pieces": 2},
                    "printed": {"pieces": 9},
                    "gaps": [{"field": "pieces", "rows": 2, "printed": 9}],
                }
            ],
            "leftovers": [{"name": "other.xlsx / Sheet", "role": "invoice", "why": "строки с количеством не вошли в лоты"}],
            "unreadable": [{"name": "scan.pdf", "role": "scan"}],
        },
    )
    assert saved == target
    from openpyxl import load_workbook

    book = load_workbook(target)
    sheet = book.active
    assert sheet.title == "Проверка"
    text = " ".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert "9" in text
    assert "2" in text
    assert "одно не выбрано" in text
    assert "только картинка" in text
    assert write_recognition_book(tmp_path / "empty.xlsx", {}) is None
