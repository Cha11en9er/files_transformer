from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.services.export import export_18233
from app.transform.extract import classify_role
from app.transform.service import transform_paths


def _item() -> dict:
    return {
        "article": "Melange 928",
        "commercial_data": {"qty": 302.9, "unit": "meters", "price": 22.5, "amount": 6815.25},
        "packing_data": {
            "rolls": 8,
            "meters": 302.9,
            "area": 427.089,
            "net_weight": 209,
            "gross_weight": 219,
            "width": 1.41,
            "gm": 690,
        },
        "customs_data": {"hs_code": "5407530000", "tnved_code": "5407530000"},
    }


def test_packing_letterhead_is_compact_and_inv_in_last_columns(tmp_path: Path) -> None:
    header = {
        "buyer": '"SM REGIONTEKSTIL\'" LLC',
        "buyer_address": "143421, Moscow region\nOGRN 1175024014472    TIN 5032281280    KPP 502401001",
        "seller": "HANGZHOU ZHONGFANG TEXTILE IMP./EXP. CO., LTD.",
        "invoice_no": "ZFRMB26148-626-1",
        "invoice_date": "Aug.19,2026",
        "contract_no": "SM-LU2",
        "contract_date": "23/11/2018",
        "delivery_terms": "EX-WORK HANGZHOU",
    }
    paths = export_18233([_item()], tmp_path, header=header, shipment_title="ZFRMB26148-626-1")
    packing = next(p for p in paths if "ПАКИНГ" in p.name.upper())
    invoice = next(p for p in paths if "ИНВОЙС" in p.name.upper())
    spec = next(p for p in paths if "СПЕЦИФ" in p.name.upper())

    pl = load_workbook(packing).active
    assert pl.auto_filter.ref in (None, "")
    blob = " | ".join(str(pl.cell(r, c).value or "") for r in range(1, 12) for c in range(1, 10))
    assert "REGIONTEKSTIL" in blob.upper()
    assert "143421" in blob
    assert "SELLER:" not in blob.upper()
    assert "MANUFACTURER" not in blob.upper()
    assert "SPECIFICATION" not in blob.upper()
    found_inv = False
    found_date = False
    for r in range(1, 12):
        for c in range(1, 9):
            val = str(pl.cell(r, c).value or "").upper()
            if val in {"INV.NO.", "INV.NO"}:
                found_inv = True
                assert c == 7
                assert str(pl.cell(r, c + 1).value) == "ZFRMB26148-626-1"
            if val == "DATE:":
                found_date = True
                assert c == 7
                assert "Aug.19" in str(pl.cell(r, c + 1).value)
    assert found_inv and found_date

    header_row = next(r for r in range(1, 20) if str(pl.cell(r, 1).value or "").upper() in {"NO.", "NO"})
    data_row = header_row + 1
    assert pl.cell(header_row, 1).font.bold
    assert not (pl.cell(data_row, 7).font.bold or False)

    inv = load_workbook(invoice).active
    inv_blob = " | ".join(str(inv.cell(r, c).value or "") for r in range(1, 20) for c in range(1, 11))
    assert "REGIONTEKSTIL" in inv_blob.upper()
    assert "143421" in inv_blob
    assert "Aug.19" in inv_blob
    assert inv.auto_filter.ref in (None, "")

    sp = load_workbook(spec).active
    spec_blob = " | ".join(str(sp.cell(r, c).value or "") for r in range(1, 18) for c in range(1, 14))
    assert "REGIONTEKSTIL" in spec_blob.upper()
    assert "Aug.19" in spec_blob
    assert sp.auto_filter.ref in (None, "")


def test_catalog_filename_is_not_goods(tmp_path: Path) -> None:
    assert classify_role("specification art quantity hs", {"qty", "article"}, source="справочник_сводная.xlsx") == "catalog"
    book = Workbook()
    sheet = book.active
    sheet.title = "Specification"
    sheet.append(["Art No.", "Quantity", "HS code", "Description"])
    sheet.append(["ZZ-99", 5, "5407610000", "sofa fabric"])
    catalog = tmp_path / "справочник_сводная.xlsx"
    book.save(catalog)
    invoice = tmp_path / "invoice.xlsx"
    inv = Workbook()
    ws = inv.active
    ws.title = "Invoice"
    ws.append(["DATE:", "Aug.19,2026"])
    ws.append(["Art No.", "Quantity", "Unit", "Price", "Amount"])
    ws.append(["ABC-1", 10, "pcs", 2, 20])
    inv.save(invoice)
    result = transform_paths(
        [(str(invoice), invoice.name), (str(catalog), catalog.name)],
    )
    roles = {f.filename: f.role_summary for f in result.files}
    assert "catalog" in (roles.get(catalog.name) or "")
    assert all(it.article != "ZZ-99" for it in result.items)
