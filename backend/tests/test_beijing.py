from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.parsing.pipeline import parse_upload
from app.parsing.product_row import is_plausible_article
from app.services.export_beijing import export_beijing_book, spec_headers
from app.services.reconcile import items_to_dicts, reconcile_documents

DOCS = Path(__file__).resolve().parents[2] / "documents"
KIT_002 = (
    DOCS
    / "17974 Шенжень 25.06 FESU5421074"
    / "Исходные"
    / "Invoice n Packing list（002-26E,20260629).xlsx"
)
KIT_003 = (
    DOCS
    / "17974 Шенжень 25.06 FESU5421074"
    / "Исходные"
    / "Invoice n Packing list（003-26E,20260703)-00.xlsx"
)


def test_digit_letter_codes_are_articles() -> None:
    assert is_plausible_article("767B")
    assert is_plausible_article("711B-1")
    assert is_plausible_article("1192")
    assert not is_plausible_article("12")


def test_merged_article_keeps_separate_qty_lots(tmp_path: Path) -> None:
    path = tmp_path / "goldluck.xlsx"
    book = Workbook()
    invoice = book.active
    invoice.title = "Invoice + Packing list"
    invoice.append(["Commercial invoice"])
    invoice.append(["Art No.", "Color", "Quantity", "Unit", "Price", "Amount", "Net Wt.", "Gross Wt."])
    invoice.append(["I2388-120", "Sand Black", 1080, "pcs", 5.54, 5983.2, 351, 372.5])
    invoice.append([None, "Sand Black", 18, "pcs", 5.54, 99.72, 6, 6.5])
    invoice.merge_cells("A3:A4")
    invoice.append(["D680", "Black", 20, "sets", 163.4, 3268, 100, 105])
    invoice.append([None, None, None, None, None, None, 30, 32])
    invoice.append([None, None, None, None, None, None, 33, 32])
    invoice.merge_cells("A5:A7")
    invoice.merge_cells("C5:C7")
    invoice.merge_cells("E5:E7")
    invoice.merge_cells("F5:F7")
    packing = book.create_sheet("Packing list")
    packing.append(["Packing list"])
    packing.append(["Art No.", "Quantity", "Cartons", "Net Wt.", "Gross Wt.", "Volume"])
    packing.append(["I2388-120", 1080, 27, 351, 372.5, 0.575])
    packing.append(["I2388-120", 18, 1, 6, 6.5, 0.021])
    packing.append(["D680", 20, 8, 163, 169, 0.56])
    book.save(path)

    docs = parse_upload(str(path), filename=path.name, allow_ocr=False)
    types = {doc.doc_type.value for doc in docs if doc.doc_type}
    assert "INVOICE" in types
    assert "PACKING_LIST" in types
    items = items_to_dicts(reconcile_documents(docs))
    lots = [item for item in items if item["article"] == "I2388-120"]
    assert sorted(item["commercial_data"]["qty"] for item in lots) == [18, 1080]
    amounts = {item["commercial_data"]["qty"]: item["commercial_data"]["amount"] for item in lots}
    assert amounts[1080] == 5983.2
    assert amounts[18] == 99.72
    d680 = next(item for item in items if item["article"] == "D680")
    assert d680["commercial_data"]["qty"] == 20
    assert d680["packing_data"]["net_weight"] == 163
    assert d680["commercial_data"]["amount"] == 3268


def test_beijing_export_has_no_and_spec_without_color(tmp_path: Path) -> None:
    items = [
        {
            "article": "KD020",
            "normalized_article": "KD020",
            "commercial_data": {"qty": 5040, "unit": "sets", "price": 3.35, "amount": 16884, "color": "Zinc"},
            "packing_data": {
                "boxes": 42,
                "net_weight": 918.0,
                "gross_weight": 935.0,
                "volume": 0.59,
                "measurement": "29.5*25*19",
            },
            "customs_data": {},
            "source_traces": {"invoice": {"qty": 5040}},
        }
    ]
    out = tmp_path / "beijing.xlsx"
    export_beijing_book(items, out, {"invoice_no": "26BEET-002", "contract_no": "DJO-3", "buyer": "ELEMENT LLC"})
    wb = load_workbook(out)
    assert wb.sheetnames == ["Invoice", "Packing list", "Specification", "описание"]
    invoice = wb["Invoice"]
    spec = wb["Specification"]
    inv_headers = [cell.value for cell in next(invoice.iter_rows(min_row=1, max_row=20)) if False]
    inv_header_row = None
    spec_header_row = None
    for row in invoice.iter_rows(min_row=1, max_row=20, values_only=True):
        if row and row[0] == "No":
            inv_header_row = list(row)
            break
    for row in spec.iter_rows(min_row=1, max_row=20, values_only=True):
        if row and row[0] == "No":
            spec_header_row = list(row)
            break
    assert inv_header_row[0] == "No"
    assert "Color" in inv_header_row
    assert spec_header_row[0] == "No"
    assert spec_headers() == [h for h in spec_header_row if h]
    assert "Color" not in spec_header_row
    body = None
    for row in invoice.iter_rows(min_row=1, max_row=30, values_only=True):
        if row and row[3] == "KD020":
            body = row
            break
    assert body is not None
    assert body[0] == 1
    assert type(body[0]) is int
    assert invoice.freeze_panes is None
    assert spec.freeze_panes is None
    assert body[4] == "Zinc"
    assert body[5] == 5040
    assert body[7] == 3.35
    assert "Color" not in "".join(str(h or "") for h in spec_header_row)


def test_shipment_title_defaults_to_invoice_no() -> None:
    from app.services.export_style import resolve_shipment_title

    assert resolve_shipment_title("", {"invoice_no": "INV_WAY04"}) == "INV_WAY04"
    assert resolve_shipment_title("Комплект документов", {"invoice_no": "26BEET-003"}) == "26BEET-003"
    assert resolve_shipment_title("export", {"invoice_no": "ZFRMB26148-626-1"}) == "ZFRMB26148-626-1"
    assert resolve_shipment_title("Goldluck", {"invoice_no": "26BEET-003"}) == "Goldluck"
    assert resolve_shipment_title("", {}) == "export"


def test_export_preview_uses_invoice_no_when_title_blank() -> None:
    from fastapi.testclient import TestClient

    from app.main import app

    client = TestClient(app)
    preview = client.post(
        "/api/v1/shipments/export/preview",
        json={
            "title": "",
            "profile_type": "BEIJING",
            "header_fields": {"invoice_no": "INV_WAY04"},
            "items": [],
        },
    )
    assert preview.status_code == 200, preview.text
    filename = preview.json()["files"][0]["filename"]
    assert filename.startswith("INV_WAY04")

    custom = client.post(
        "/api/v1/shipments/export/preview",
        json={
            "title": "Моя поставка",
            "profile_type": "BEIJING",
            "header_fields": {"invoice_no": "INV_WAY04"},
            "items": [],
        },
    )
    assert custom.status_code == 200, custom.text
    assert custom.json()["files"][0]["filename"].startswith("Моя поставка")


def test_beijing_export_filename_and_table_borders(tmp_path: Path) -> None:
    from app.services.export_style import safe_export_stem

    assert safe_export_stem("WAY04 / Bestway") == "WAY04 _ Bestway"
    items = [
        {
            "article": "KD020",
            "normalized_article": "KD020",
            "commercial_data": {"qty": 10, "unit": "pcs", "price": 1, "amount": 10},
            "packing_data": {"boxes": 1, "net_weight": 2.0, "gross_weight": 3.0},
            "customs_data": {},
            "source_traces": {},
        }
    ]
    out = tmp_path / f"{safe_export_stem('Моя поставка')} для ЭД.xlsx"
    export_beijing_book(items, out, {"invoice_no": "1"})
    wb = load_workbook(out)
    invoice = wb["Invoice"]
    header_row = None
    for row in invoice.iter_rows(min_row=1, max_row=20):
        if row and row[0].value == "No":
            header_row = row
            break
    assert header_row is not None
    assert header_row[0].border.left.style == "thin"
    assert header_row[0].font.bold is True
    assert invoice.cell(1, 1).font.bold is True



def test_17974_002_keeps_digit_articles_and_lots() -> None:
    if not KIT_002.exists():
        return
    docs = parse_upload(str(KIT_002), filename=KIT_002.name, allow_ocr=False)
    items = items_to_dicts(reconcile_documents(docs))
    by_art = {}
    for item in items:
        by_art.setdefault(item["article"], []).append(item)
    assert "767B" in by_art
    assert "711B-1" in by_art
    lots = by_art["I2388-120"]
    assert sorted(item["commercial_data"]["qty"] for item in lots) == [18, 1080]
    assert by_art["KD020"][0]["commercial_data"]["amount"] == 16884
    assert by_art["KD020"][0]["commercial_data"]["color"] == "Zinc"
    types = {doc.doc_type.value for doc in docs if doc.doc_type}
    assert types == {"INVOICE", "PACKING_LIST"}


def test_17974_003_set_and_catalog_example() -> None:
    if not KIT_003.exists():
        return
    docs = parse_upload(str(KIT_003), filename=KIT_003.name, allow_ocr=False)
    items = items_to_dicts(reconcile_documents(docs))
    by_art = {item["article"]: item for item in items}
    assert by_art["MD 812"]["commercial_data"]["qty"] == 200100
    assert by_art["MD 812"]["commercial_data"]["price"] == 0.0592
    assert by_art["MD 812"]["commercial_data"]["amount"] == 11845.92
    assert by_art["D680"]["commercial_data"]["qty"] == 20
    assert by_art["D680"]["packing_data"]["net_weight"] == 163
    lots = [item for item in items if item["article"] == "NO228"]
    assert sorted(item["commercial_data"]["qty"] for item in lots) == [20, 1480]
