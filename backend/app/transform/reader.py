"""Unified matrix reader.

One job: turn a workbook (.xlsx/.xlsm/.xls, even an .html renamed to .xls) into a
list of `Sheet`, where each sheet is a rectangular grid with merged cells already
expanded (the origin value copied into every covered cell). Losing merged article
cells on .xls was the single biggest source of "values from the wrong row".
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.transform.errors import FileReadError

EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}


@dataclass
class Sheet:
    name: str
    grid: list[list[Any]]                       # merges already expanded
    inherited: set[tuple[int, int]] = field(default_factory=set)  # (row, col) filled from a merge
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)  # 0-based inclusive ranges
    source: str = ""                            # file name

    @property
    def nrows(self) -> int:
        return len(self.grid)

    @property
    def ncols(self) -> int:
        return max((len(r) for r in self.grid), default=0)


def _rectangular(grid: list[list[Any]]) -> list[list[Any]]:
    width = max((len(r) for r in grid), default=0)
    for row in grid:
        if len(row) < width:
            row.extend([None] * (width - len(row)))
    return grid


def _fill_merges(grid: list[list[Any]], ranges: list[tuple[int, int, int, int]]) -> set[tuple[int, int]]:
    """ranges are 0-based (r0, c0, r1, c1) inclusive. Copy origin into covered cells."""
    inherited: set[tuple[int, int]] = set()
    for r0, c0, r1, c1 in ranges:
        if r0 >= len(grid) or c0 >= len(grid[r0]):
            continue
        origin = grid[r0][c0]
        for r in range(r0, min(r1, len(grid) - 1) + 1):
            for c in range(c0, min(c1, len(grid[r]) - 1) + 1):
                if (r, c) == (r0, c0):
                    continue
                if grid[r][c] in (None, ""):
                    grid[r][c] = origin
                    inherited.add((r, c))
    return inherited


def _sniff(path: str) -> str:
    try:
        head = Path(path).read_bytes()[:512]
    except OSError as exc:
        raise FileReadError("Файл не найден. Загрузи его ещё раз.", technical=str(exc)) from exc
    if head.startswith(b"PK"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    low = head.lstrip().lower()
    if low.startswith((b"<html", b"<!doctype", b"<?xml", b"<table")) or b"<table" in low:
        return "html"
    suffix = Path(path).suffix.lower()
    if suffix == ".xls":
        return "xls"
    return "xlsx"


def _read_xlsx(path: str) -> list[Sheet]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - mapped to a friendly message upstream
        raise FileReadError(_xlsx_hint(exc), technical=repr(exc)) from exc
    sheets: list[Sheet] = []
    try:
        for ws in wb.worksheets:
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            if not max_row or not max_col:
                continue
            grid = [[ws.cell(r, c).value for c in range(1, max_col + 1)] for r in range(1, max_row + 1)]
            ranges = [
                (m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1)
                for m in ws.merged_cells.ranges
            ]
            inherited = _fill_merges(grid, ranges)
            sheets.append(Sheet(
                name=str(ws.title),
                grid=_rectangular(grid),
                inherited=inherited,
                merges=ranges,
                source=Path(path).name,
            ))
    finally:
        wb.close()
    return sheets


def _xlsx_hint(exc: Exception) -> str:
    low = repr(exc).lower()
    if "not a zip" in low or "badzip" in low or "central directory" in low:
        return "Файл повреждён или это не настоящий .xlsx. Открой его в Excel и пересохрани как .xlsx."
    if "password" in low or "encrypt" in low:
        return "Файл защищён паролем. Сними защиту в Excel и загрузи снова."
    return "Не удалось открыть .xlsx. Пересохрани файл в Excel и загрузи снова."


def _read_xls(path: str) -> list[Sheet]:
    import xlrd

    try:
        book = xlrd.open_workbook(path, formatting_info=True)
    except Exception:
        # formatting_info is unavailable for some .xls variants; retry without merges
        try:
            book = xlrd.open_workbook(path)
        except Exception as exc:  # noqa: BLE001
            raise FileReadError(
                "Не удалось открыть .xls. Пересохрани файл в Excel как .xlsx и загрузи снова.",
                technical=repr(exc),
            ) from exc
    sheets: list[Sheet] = []
    for ws in book.sheets():
        if ws.nrows == 0:
            continue
        grid = [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]
        grid = [[(None if v == "" else v) for v in row] for row in grid]
        ranges = [(r0, c0, r1 - 1, c1 - 1) for (r0, r1, c0, c1) in getattr(ws, "merged_cells", [])]
        inherited = _fill_merges(grid, ranges)
        sheets.append(Sheet(
            name=str(ws.name),
            grid=_rectangular(grid),
            inherited=inherited,
            merges=ranges,
            source=Path(path).name,
        ))
    return sheets


def _read_html(path: str) -> list[Sheet]:
    import pandas as pd

    from app.parsing.languages import decode_bytes

    raw = Path(path).read_bytes()
    text, _enc = decode_bytes(raw)
    try:
        frames = pd.read_html(io.StringIO(text), header=None)
    except Exception as exc:  # noqa: BLE001
        raise FileReadError(
            "Файл выглядит как .html, а не Excel. Открой его в Excel и пересохрани как .xlsx.",
            technical=repr(exc),
        ) from exc
    sheets: list[Sheet] = []
    for i, frame in enumerate(frames):
        grid = frame.where(frame.notna(), None).values.tolist()
        sheets.append(Sheet(name=f"Sheet{i + 1}", grid=_rectangular(grid), source=Path(path).name))
    return sheets


def read_workbook(path: str) -> list[Sheet]:
    """Read every sheet of an Excel workbook with merged cells expanded."""
    kind = _sniff(path)
    if kind == "html":
        sheets = _read_html(path)
    elif kind == "xls":
        try:
            sheets = _read_xls(path)
        except FileReadError:
            # some .xls are actually xlsx/html mislabelled
            try:
                sheets = _read_xlsx(path)
            except FileReadError:
                sheets = _read_html(path)
    else:
        try:
            sheets = _read_xlsx(path)
        except FileReadError:
            sheets = _read_xls(path)
    sheets = [s for s in sheets if s.nrows]
    if not sheets:
        raise FileReadError("В файле нет данных для обработки.")
    return sheets
