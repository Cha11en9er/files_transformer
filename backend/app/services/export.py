"""Excel export workers — Word spec §3.2 / §4.1 / §7.

Profile 18233: three workbooks (invoice, packing, specification with colors).
PDF-описания в первом этапе не формируются.
Profile BEIJING: one workbook with sheets Invoice, Packing list, Specification, Description.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.services.export_18233_templates import (
    _detect_kit,
    _layout_kit,
    invoice_table_rows,
    is_fabric_layout,
    packing_table_rows,
    preview_headers,
    spec_table_rows,
)
from app.services.export_beijing import (
    description_headers,
    description_rows,
    export_beijing_book,
    invoice_headers,
    invoice_rows,
    packing_headers,
    packing_rows,
    spec_headers,
    spec_rows,
)
from app.services.export_tsd import export_tsd_book, is_tsd_layout, tsd_preview
from app.services.export_style import (
    merge_row,
    safe_export_stem,
    style_data_table,
    style_letterhead_row,
    style_split_letterhead,
    unfreeze_workbook,
)


def _write_sheet(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)


def _product_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    products = [
        item
        for item in items
        if (item.get("source_traces") or {}).get("invoice")
        or (item.get("source_traces") or {}).get("specification")
        or item.get("article")
    ]
    if not products:
        products = list(items)

    def _sort_key(item: dict[str, Any]) -> int:
        traces = item.get("source_traces") or {}
        inv = traces.get("invoice") or {}
        no = inv.get("no")
        try:
            return int(no)
        except (TypeError, ValueError):
            return 999

    products.sort(key=_sort_key)
    return products


def _kit_label(items: list[dict[str, Any]], header: dict[str, Any] | None) -> str:
    kit = _detect_kit(items, header)
    if kit:
        return kit
    inv = str((header or {}).get("invoice_no") or "")
    if inv:
        tail = inv.split("-")[-1].strip()
        if tail and tail.lower() != "all":
            return tail
    return ""


def _description_text(item: dict[str, Any]) -> str:
    customs = item.get("customs_data") or {}
    en = customs.get("description_en")
    ru = customs.get("description_ru")
    if en and ru and not is_factory_note(en) and not is_factory_note(ru):
        return f"{en}/{ru}"
    for key in ("description", "description_en", "description_ru"):
        value = customs.get(key)
        if value and not is_factory_note(value):
            return str(value)
    return ""


def _qty_of(item: dict[str, Any]) -> Any:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    if commercial.get("qty") not in (None, ""):
        return commercial.get("qty")
    return packing.get("meters")


def _unit_of(item: dict[str, Any]) -> str:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    return str(commercial.get("unit") or packing.get("unit") or "")


def _places_of(item: dict[str, Any]) -> Any:
    packing = item.get("packing_data") or {}
    traces = item.get("source_traces") or {}
    pl = traces.get("packing_list_group") or traces.get("invoice") or {}
    return packing.get("rolls") or packing.get("boxes") or pl.get("rolls")


def _packing_design(item: dict[str, Any]) -> str:
    traces = item.get("source_traces") or {}
    pl = traces.get("packing_list_group") or {}
    design = pl.get("design")
    if design and not is_factory_note(design):
        return str(design).replace(" / ", "\n")
    article = item.get("article") or ""
    category = pl.get("category")
    if category and not is_factory_note(category):
        family = article.split()[0] if article else article
        return f"{category}\n{family}".strip()
    return str(article)


def _item_rows(items: list[dict[str, Any]], kind: str) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for item in items:
        commercial = item.get("commercial_data") or {}
        packing = item.get("packing_data") or {}
        customs = item.get("customs_data") or {}
        article = item.get("article") or item.get("normalized_article") or ""
        if kind == "invoice":
            rows.append(
                [
                    article,
                    item.get("model"),
                    commercial.get("color"),
                    _qty_of(item),
                    _unit_of(item),
                    commercial.get("price"),
                    commercial.get("amount"),
                    commercial.get("currency"),
                    customs.get("hs_code"),
                ]
            )
        elif kind == "packing":
            rows.append(
                [
                    article,
                    packing.get("rolls") or packing.get("boxes"),
                    packing.get("meters") if packing.get("meters") is not None else commercial.get("qty"),
                    packing.get("area"),
                    packing.get("net_weight"),
                    packing.get("gross_weight"),
                    packing.get("volume"),
                ]
            )
        elif kind == "specification":
            rows.append(
                [
                    article,
                    commercial.get("color"),
                    packing.get("rolls") or packing.get("boxes"),
                    packing.get("meters") if packing.get("meters") is not None else commercial.get("qty"),
                    packing.get("area"),
                    packing.get("net_weight"),
                    packing.get("gross_weight"),
                    customs.get("tnved_code") or customs.get("hs_code"),
                    customs.get("description_en"),
                    customs.get("description_ru"),
                ]
            )
        elif kind == "description":
            rows.append(
                [
                    article,
                    customs.get("tnved_code"),
                    customs.get("description_en"),
                    customs.get("description_ru"),
                    customs.get("country"),
                    customs.get("manufacturer"),
                ]
            )
    return rows


def _aligned_invoice_rows(items: list[dict[str, Any]], fabric: bool = True) -> tuple[list[str], list[list[Any]]]:
    headers = preview_headers("626-1" if fabric else "18312")["invoice"]
    return headers, invoice_table_rows(items, fabric)


def _aligned_packing_rows(items: list[dict[str, Any]], fabric: bool = True) -> tuple[list[str], list[list[Any]]]:
    headers = preview_headers("626-1" if fabric else "18312")["packing"]
    return headers, packing_table_rows(items, fabric)


def _aligned_spec_rows(items: list[dict[str, Any]], fabric: bool = True) -> tuple[list[str], list[list[Any]]]:
    headers = preview_headers("626-1" if fabric else "18312")["specification"]
    return headers, spec_table_rows(items, fabric)


def build_export_preview(
    items: list[dict[str, Any]],
    *,
    profile_type: str,
    header: dict[str, Any] | None = None,
    shipment_title: str = "export",
) -> dict[str, Any]:
    """JSON preview of files that export() will write. No PDF/txt."""
    if str(profile_type).upper() in {"BEIJING", "PROFILETYPE.BEIJING"}:
        products = _product_items(items)
        stem = safe_export_stem(shipment_title)
        if is_tsd_layout(products, header):
            return {
                "profile_type": "BEIJING",
                "files": [tsd_preview(products, header, f"ТСД {stem}.xlsx")],
            }
        return {
            "profile_type": "BEIJING",
            "files": [
                {
                    "filename": f"{stem} для ЭД.xlsx",
                    "sheets": [
                        {"title": "Invoice", "headers": invoice_headers(), "rows": invoice_rows(products)},
                        {"title": "Packing list", "headers": packing_headers(), "rows": packing_rows(products)},
                        {"title": "Specification", "headers": spec_headers(), "rows": spec_rows(products)},
                        {"title": "описание", "headers": description_headers(), "rows": description_rows(products, header)},
                    ],
                }
            ],
        }

    products = _product_items(items)
    kit = _kit_label(products, header)
    layout = _layout_kit(kit, products, header)
    fabric = is_fabric_layout(layout)
    suffix = f" {kit}" if kit else ""
    stem = safe_export_stem(shipment_title)
    inv_h, inv_r = _aligned_invoice_rows(products, fabric)
    pl_h, pl_r = _aligned_packing_rows(products, fabric)
    spec_h, spec_r = _aligned_spec_rows(products, fabric)
    return {
        "profile_type": "18233",
        "files": [
            {
                "filename": f"{stem} ИНВОЙС{suffix}.xlsx",
                "sheets": [{"title": "invoice", "headers": inv_h, "rows": inv_r}],
            },
            {
                "filename": f"{stem} ПАКИНГ{suffix}.xlsx",
                "sheets": [{"title": "packing", "headers": pl_h, "rows": pl_r}],
            },
            {
                "filename": f"{stem}{suffix} СПЕЦИФИКАЦИЯ С ЦВЕТАМИ.xlsx",
                "sheets": [{"title": "specification", "headers": spec_h, "rows": spec_r}],
            },
        ],
    }


def _pad_row(cols: int, *pairs: tuple[int, Any]) -> list[Any]:
    row: list[Any] = [None] * max(cols, 1)
    for idx, value in pairs:
        if 1 <= idx <= len(row):
            row[idx - 1] = value
    return row


def _append_right_pair(ws, left: Any, label: str, value: Any, cols: int) -> None:
    """Left text + label/value in the last two table columns (INV.NO / DATE)."""
    row = [None] * max(cols, 1)
    row[0] = left
    if cols >= 2:
        row[cols - 2] = label
        row[cols - 1] = value
    ws.append(row)
    if cols >= 4:
        merge_row(ws, ws.max_row, 1, cols - 2)
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="center")


def _write_letterhead(ws, kind: str, header: dict[str, Any] | None, cols: int = 10) -> None:
    header = header or {}
    seller = header.get("seller") or ""
    buyer = header.get("buyer") or ""
    contract = header.get("contract_no") or ""
    contract_date = header.get("contract_date") or ""
    invoice_no = header.get("invoice_no") or ""
    date = header.get("invoice_date") or header.get("date") or ""
    delivery = header.get("delivery_terms") or ""
    payment = header.get("payment_terms") or ""
    manufacturer = header.get("manufacturer") or seller
    buyer_address = header.get("buyer_address") or ""
    seller_address = header.get("seller_address") or ""
    warehouse = header.get("warehouse_address") or ""
    delivery_date = header.get("delivery_date") or ""
    title = {"invoice": "INVOICE", "packing": "PACKING LIST", "specification": "SPECIFICATION"}.get(kind, kind.upper())
    if seller:
        ws.append([seller])
        style_letterhead_row(ws, ws.max_row, cols, company=True)
    if seller_address and kind != "packing":
        ws.append([seller_address])
        style_letterhead_row(ws, ws.max_row, cols)
    ws.append([])
    if kind == "specification":
        spec_title = f"SPECIFICATION №{invoice_no}".strip()
        if date:
            spec_title = f"{spec_title} dd {date}"
        ws.append([spec_title])
        style_letterhead_row(ws, ws.max_row, cols, title=True)
    else:
        ws.append([title])
        style_letterhead_row(ws, ws.max_row, cols, title=True)
    ws.append([])
    if kind == "packing":
        _write_packing_letterhead(ws, header, cols)
        return
    split = min(6, cols)
    if contract:
        dated = f" dd {contract_date}" if contract_date else ""
        ws.append([f"Contract №：{contract}{dated}"])
        style_letterhead_row(ws, ws.max_row, cols)
    ws.append(["Buyer:", None, None, None, None, "Seller:"])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split)
    ws.append([buyer, None, None, None, None, seller])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    if buyer_address or seller_address:
        ws.append(["Address:", None, None, None, None, "Address:"])
        style_split_letterhead(ws, ws.max_row, cols, split_at=split)
        ws.append([buyer_address, None, None, None, None, seller_address])
        style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
        ws.row_dimensions[ws.max_row].height = 36
    if warehouse:
        ws.append([None, None, None, None, None, "Warehouse address:"])
        style_split_letterhead(ws, ws.max_row, cols, split_at=split)
        ws.append([None, None, None, None, None, warehouse])
        style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    if payment:
        ws.append([f"Terms of payment: {payment}"])
        style_letterhead_row(ws, ws.max_row, cols)
    if delivery:
        ws.append([f"Terms of delivery: {delivery}"])
        style_letterhead_row(ws, ws.max_row, cols)
    if delivery_date:
        ws.append([f"Delivery date: {delivery_date}"])
        style_letterhead_row(ws, ws.max_row, cols)
    if manufacturer:
        ws.append([f"Manufacturer: {manufacturer}"])
        style_letterhead_row(ws, ws.max_row, cols)
    ws.append([])
    if kind == "specification":
        spec_row = _pad_row(
            cols,
            (1, "инв номер:"),
            (2, invoice_no),
            (cols - 1, "дата:"),
            (cols, date),
        )
        ws.append(spec_row)
    else:
        spec_no = f"SPECIFICATION № {invoice_no}".strip()
        if date:
            spec_no = f"{spec_no} {date}"
        ws.append([spec_no])
        style_letterhead_row(ws, ws.max_row, cols)
        left_incoterm = delivery or ""
        contract_line = f"Contract No {contract}".strip()
        if contract_date:
            contract_line = f"{contract_line} dd {contract_date}".strip()
        _append_right_pair(ws, left_incoterm, "INV.NO.", invoice_no, cols)
        _append_right_pair(ws, contract_line, "DATE:", date, cols)
    ws.append([])


def _write_packing_letterhead(ws, header: dict[str, Any], cols: int) -> None:
    """Packing keeps the compact source block: one Buyer cell, INV.NO in last columns."""
    buyer = header.get("buyer") or ""
    buyer_address = header.get("buyer_address") or ""
    invoice_no = header.get("invoice_no") or ""
    date = header.get("invoice_date") or header.get("date") or ""
    contract = header.get("contract_no") or ""
    contract_date = header.get("contract_date") or ""
    delivery = header.get("delivery_terms") or ""
    block = buyer
    if buyer_address:
        block = f"{buyer}\n{buyer_address}" if buyer else buyer_address
    if block:
        prefix = "" if str(block).lower().startswith("buyer") else "Buyer:"
        ws.append([f"{prefix}{block}"])
        style_letterhead_row(ws, ws.max_row, cols)
        ws.row_dimensions[ws.max_row].height = 48
    contract_line = f"Contract No {contract}".strip() if contract else ""
    if contract and contract_date:
        contract_line = f"Contract No {contract} dd {contract_date}"
    _append_right_pair(ws, delivery, "INV.NO.", invoice_no, cols)
    _append_right_pair(ws, contract_line, "DATE:", date, cols)
    ws.append([])


def _save_simple_book(path: Path, title: str, headers: list[str], rows: list[list[Any]], header: dict[str, Any] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"
    _write_letterhead(ws, title, header, cols=max(len(headers), 1))
    header_row = ws.max_row + 1
    _write_sheet(ws, headers, rows)
    for row in ws.iter_rows(min_row=ws.max_row - len(rows), max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            if type(cell.value) is int:
                cell.number_format = "0"
            elif isinstance(cell.value, float):
                cell.number_format = "0.000" if abs(cell.value) >= 100 or (cell.column in {6, 8, 9} and title == "specification") else "0.00"
            if isinstance(cell.value, str) and cell.value.upper().startswith("TOTAL"):
                cell.font = Font(bold=True)
    style_data_table(
        ws,
        header_row=header_row,
        cols=len(headers),
        n_rows=1 + len(rows),
        headers=headers,
    )
    unfreeze_workbook(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_beijing(items: list[dict[str, Any]], output_path: Path, header: dict[str, Any] | None = None) -> Path:
    products = _product_items(items)
    if is_tsd_layout(products, header):
        stem = output_path.stem
        if "для ЭД" in stem:
            output_path = output_path.with_name(f"ТСД {stem.replace(' для ЭД', '')}.xlsx")
        return export_tsd_book(products, output_path, header)
    return export_beijing_book(products, output_path, header)


def export_18233(
    items: list[dict[str, Any]],
    output_dir: Path,
    header: dict[str, Any] | None = None,
    shipment_title: str = "18233",
) -> list[Path]:
    """Three Excel files with the same product rows. No txt/PDF stubs."""
    output_dir.mkdir(parents=True, exist_ok=True)
    products = _product_items(items)
    kit = _kit_label(products, header)
    layout = _layout_kit(kit, products, header)
    fabric = is_fabric_layout(layout)
    suffix = f" {kit}" if kit else ""
    stem = safe_export_stem(shipment_title)
    writers = {
        "invoice": lambda rows: _aligned_invoice_rows(rows, fabric),
        "packing": lambda rows: _aligned_packing_rows(rows, fabric),
        "specification": lambda rows: _aligned_spec_rows(rows, fabric),
    }
    paths: list[Path] = []
    names = [
        (f"{stem} ИНВОЙС{suffix}.xlsx", "invoice"),
        (f"{stem} ПАКИНГ{suffix}.xlsx", "packing"),
        (f"{stem}{suffix} СПЕЦИФИКАЦИЯ С ЦВЕТАМИ.xlsx", "specification"),
    ]
    for filename, kind in names:
        headers, rows = writers[kind](products)
        paths.append(
            _save_simple_book(
                output_dir / filename,
                kind,
                headers,
                rows,
                header=header,
            )
        )
    return paths
