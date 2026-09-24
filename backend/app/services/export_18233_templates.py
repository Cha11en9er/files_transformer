"""Export 18233 by filling copies of customer эталон templates."""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font

from app.parsing.header_extract import (
    currency_from_sources,
    export_currency_label,
    export_header_fields,
    manufacturer_from_sources,
)
from app.parsing.normalize import normalize_article
from app.services.export_style import safe_export_stem, unfreeze_workbook
from app.services.field_map import article_with_color, classify_header, is_factory_note, parse_number
from app.services.materials_18233 import kit_templates, materials_available
from app.services.profile_18233 import GROUP_PREFIXES, _design_family

_KIT_RE = re.compile(r"(?<![0-9])(\d{2,4}-\d)(?![0-9])")
_INVOICE_TOKEN_RE = re.compile(r"ZFRMB[\w.-]+", re.I)
_PRICE_CCY_RE = re.compile(r"(UNIT PRICE|AMOUNT)\s*\(\s*[A-Z]{3}\s*\)", re.I)

NUM_FMT_INT = "0"
NUM_FMT_2 = "0.00"
NUM_FMT_3 = "0.000"


def _price_headers(ccy: str) -> tuple[str, str]:
    label = export_currency_label(ccy, hangzhou_style=True)
    return f"UNIT PRICE({label})", f"AMOUNT({label})"


def fabric_invoice_headers(ccy: str = "CNY") -> list[str]:
    price, amount = _price_headers(ccy)
    return [
        "NO.",
        "DESIGN",
        "H.S. CODE",
        "ROLLS",
        "WIDTH M",
        "TOTAL M2",
        "METERS",
        "UNIT MT/PIECE",
        price,
        amount,
    ]


def element_invoice_headers(ccy: str = "CNY") -> list[str]:
    price, amount = _price_headers(ccy)
    return [
        "NO.",
        "DESIGN",
        "H.S. CODE",
        "PACKAGES",
        "QUANTITY",
        "UNIT M/PC",
        price,
        amount,
    ]


# Backward-compatible aliases for Hangzhou RMB etalon.
FABRIC_INVOICE_HEADERS = fabric_invoice_headers("CNY")
ELEMENT_INVOICE_HEADERS = element_invoice_headers("CNY")
FABRIC_PACKING_HEADERS = [
    "No.",
    "DESIGN",
    "G/M",
    "ROLLS",
    "TOTAL METERS",
    "NET WEIGHT/KG",
    "GROSS WEIGHT/KG",
    "Total M2",
]
ELEMENT_PACKING_HEADERS = [
    "No.",
    "DESIGN",
    "PACKAGES",
    "QUANTITY",
    "UNIT",
    "NET WEIGHT/KG",
    "GROSS WEIGHT/KG",
]
FABRIC_SPEC_HEADERS = [
    "№",
    "Q-Ty Rolls / Кол-во рулонов",
    "Art./Артикул",
    "Product name / Наименование товара",
    "HS code / Код гармонизированной системы",
    "Customs code / Таможенный код",
    "Q-ty meters / Кол-во погонных метров",
    "Widht, m / Ширина, м",
    "Q-ty m2 / Кол-во кв.м",
    "N.W, kg / Вес Нетто, кг",
    "G.W, kg / Вес Брутто, кг",
    "Price per 1 meter / Цена за 1 пог.метр, Юань",
    "Total price / Цена, Юань",
]
ELEMENT_SPEC_HEADERS = [
    "№",
    "Q-Ty PACKAGES / Кол-во упаковок",
    "Art./Артикул",
    "Product name / Наименование товара",
    "HS code / Код гармонизированной системы",
    "Customs code / Таможенный код",
    "Q-ty UNITS / Кол-во единиц",
    "Unit/Единица измерения",
    "Q-ty m2 / Кол-во кв.м",
    "N.W, kg / Вес Нетто, кг",
    "G.W, kg / Вес Брутто, кг",
    "Price per unit / Цена за единицу, Юань",
    "Total price / Цена, Юань",
]


def _as_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    return parse_number(value)


def _count(value: Any) -> int | float | None:
    number = _as_float(value)
    if number is None:
        return None
    rounded = round(float(number), 2)
    if abs(rounded - round(rounded)) < 1e-9:
        return int(round(rounded))
    return rounded


def _r2(value: Any) -> float | None:
    number = _as_float(value)
    if number is None:
        return None
    return round(float(number), 2)


def _r3(value: Any) -> float | None:
    number = _as_float(value)
    if number is None:
        return None
    return round(float(number), 3)


def _detect_kit(items: list[dict[str, Any]], header: dict[str, Any] | None) -> str | None:
    for item in items:
        sub = item.get("invoice_subkit") or (item.get("commercial_data") or {}).get("invoice_subkit")
        if sub and str(sub) not in {"ALL", "None", ""}:
            return str(sub)
    inv = (header or {}).get("invoice_no") or ""
    match = _KIT_RE.search(str(inv))
    if match:
        return match.group(1)
    arts = {normalize_article(i.get("article") or "") for i in items}
    blob = " ".join(
        str(((i.get("source_traces") or {}).get("packing_list_group") or {}).get("design") or "")
        for i in items
    ).upper()
    if "ARTIFICIAL LEATHER" in blob or any(a.startswith("NAPPA") or a.startswith("MAGIC") for a in arts):
        return "626-2"
    if any(a.startswith("NOBLE") or a.startswith("MELANGE") for a in arts):
        return "626-1"
    return None


def _item_blob(items: list[dict[str, Any]], header: dict[str, Any] | None = None) -> str:
    parts = [
        str((header or {}).get("buyer") or ""),
        str((header or {}).get("invoice_no") or ""),
        str((header or {}).get("contract_no") or ""),
    ]
    for item in items:
        parts.append(str(item.get("article") or ""))
        pl = (item.get("source_traces") or {}).get("packing_list_group") or {}
        parts.append(str(pl.get("design") or ""))
        parts.append(str(pl.get("category") or ""))
    return " ".join(parts).upper()


def _goods_use_fabric_columns(items: list[dict[str, Any]] | None) -> bool | None:
    """True when goods carry meters+width (Hangzhou fabric), False when they do not.

    None means there is not enough to decide. Never keys off a shop name or SKU.
    """
    products = [item for item in (items or []) if item.get("article")]
    if len(products) < 2:
        return None
    meters = sum(1 for item in products if (item.get("packing_data") or {}).get("meters") not in (None, ""))
    width = sum(1 for item in products if (item.get("packing_data") or {}).get("width") not in (None, ""))
    if meters >= max(2, int(len(products) * 0.3)) and width:
        return True
    return False


def _looks_element(items: list[dict[str, Any]], header: dict[str, Any] | None = None) -> bool:
    shaped = _goods_use_fabric_columns(items)
    if shaped is False:
        return True
    if shaped is True:
        return False
    blob = _item_blob(items, header)
    if "SOFA FABRIC" in blob or "ARTIFICIAL LEATHER" in blob:
        return False
    return bool(re.search(r"PACKAGES|UNIT M/PC", blob))


def _layout_kit(kit: str | None, items: list[dict[str, Any]], header: dict[str, Any] | None = None) -> str:
    """Pick a customer layout workbook. 626-1/626-2/18312 are shapes, not job numbers."""
    if kit in {"626-1", "626-2", "18312"}:
        return kit
    blob = _item_blob(items, header)
    if "ARTIFICIAL LEATHER" in blob or "NAPPA" in blob:
        return "626-2"
    if _looks_element(items, header):
        return "18312"
    return "626-1"


def is_fabric_layout(layout: str | None, items: list[dict[str, Any]] | None = None) -> bool:
    shaped = _goods_use_fabric_columns(items)
    if shaped is not None:
        return shaped
    return (layout or "626-1") != "18312"


def _ordered_products(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    products = [
        i
        for i in items
        if (i.get("source_traces") or {}).get("invoice") or (i.get("source_traces") or {}).get("specification")
    ]
    if not products:
        products = [i for i in items if i.get("article")]
    products.sort(key=lambda i: int(((i.get("source_traces") or {}).get("invoice") or {}).get("no") or 999))
    return products


def _format_packing_design(design: str | None, article: str | None = None) -> str:
    text = str(design or "").strip()
    if is_factory_note(text):
        text = ""
    text = re.sub(r"^\([^)]{0,48}\)\s*/\s*", "", text)
    if not text:
        return str(article or "").strip()
    if " / " in text:
        left, right = text.split(" / ", 1)
        if any(left.upper().startswith(prefix) for prefix in GROUP_PREFIXES) or len(left) <= 40:
            return f"{left}\n{right}"
    return text.replace(" / ", "\n")


def _product_description(item: dict[str, Any]) -> str:
    customs = item.get("customs_data") or {}
    for key in ("description", "description_en", "description_ru"):
        value = customs.get(key)
        if value and not is_factory_note(value):
            en = customs.get("description_en")
            ru = customs.get("description_ru")
            if key == "description" and en and ru and not is_factory_note(en) and not is_factory_note(ru):
                return f"{en}/{ru}"
            return str(value)
    return ""


def _group_prefix(item: dict[str, Any], fabric: bool) -> str:
    pl = (item.get("source_traces") or {}).get("packing_list_group") or {}
    design = str(pl.get("design") or pl.get("category") or "")
    if design:
        formatted = _format_packing_design(design, item.get("article"))
        return formatted.split("\n", 1)[0]
    group = (
        (item.get("commercial_data") or {}).get("group")
        or (item.get("source_traces") or {}).get("group")
        or ""
    )
    if group:
        return str(group).split("\n", 1)[0]
    article = str(item.get("article") or "")
    if not fabric:
        return ""
    if article.upper().startswith("NAPPA") or article.upper().startswith("MAGIC"):
        return "ARTIFICIAL LEATHER"
    return ""


def _invoice_group_key(item: dict[str, Any], fabric: bool) -> str:
    if fabric:
        return _design_family(item.get("article") or "")
    return normalize_article(item.get("article") or "") or str(id(item))


def _collect_packing_groups(products: list[dict[str, Any]], fabric: bool) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    seen: set[tuple[Any, ...]] = set()
    for item in products:
        pl = (item.get("source_traces") or {}).get("packing_list_group") or {}
        parts = pl.get("parts") or ([pl] if pl else [])
        for part in parts:
            if not part:
                continue
            sig = (part.get("row_index"), part.get("design"), part.get("rolls"), part.get("meters") or part.get("qty"))
            if sig in seen:
                continue
            seen.add(sig)
            groups.append(part)
    if groups:
        return groups
    # No family packing row in the source: each goods line keeps its own places and weight.
    for item in products:
        commercial = item.get("commercial_data") or {}
        packing = item.get("packing_data") or {}
        label = article_with_color(item.get("article"), commercial.get("color"))
        prefix = _group_prefix(item, fabric)
        design = f"{prefix}\n{label}".strip() if prefix else label
        groups.append(
            {
                "design": design,
                "article": label,
                "gm": packing.get("gm"),
                "rolls": packing.get("rolls") or packing.get("boxes"),
                "meters": packing.get("meters") if packing.get("meters") is not None else commercial.get("qty"),
                "qty": commercial.get("qty") if commercial.get("qty") is not None else packing.get("meters"),
                "unit": commercial.get("unit") or packing.get("unit"),
                "net_weight": packing.get("net_weight"),
                "gross_weight": packing.get("gross_weight"),
                "area": packing.get("area"),
            }
        )
    return groups


def invoice_table_rows(products: list[dict[str, Any]], fabric: bool) -> list[list[Any]]:
    families: list[str] = []
    by_family: dict[str, list[dict[str, Any]]] = {}
    for item in products:
        fam = _invoice_group_key(item, fabric)
        if fam not in by_family:
            families.append(fam)
            by_family[fam] = []
        by_family[fam].append(item)

    rows: list[list[Any]] = []
    no = 1
    for fam in families:
        group_items = by_family[fam]
        # A real family = a printed category (SOFA FABRIC / ARTIFICIAL LEATHER)
        # over colour children, where the source prints the unit price only on the
        # family row. Then children show colour + no price and a summary row carries
        # price and the family total. Distinct SKUs that merely share a design token
        # (MAXWELL 997 / 236, NERGIS 001 / 305) have no category prefix and their own
        # printed price - keep each as its own priced line and emit no summary row,
        # so the export never blanks a real per-position price.
        prefix = _group_prefix(group_items[0], fabric)
        is_family = bool(prefix)
        group_rolls = 0.0
        group_meters = 0.0
        group_qty = 0.0
        group_area = 0.0
        group_amount = 0.0
        price = None
        hs = None
        width = None
        unit = None
        for item in group_items:
            inv = (item.get("source_traces") or {}).get("invoice") or {}
            packing = item.get("packing_data") or {}
            commercial = item.get("commercial_data") or {}
            customs = item.get("customs_data") or {}
            rolls = _count(packing.get("rolls") or inv.get("rolls"))
            meters = _r2(packing.get("meters") or inv.get("meters"))
            qty = _count(commercial.get("qty") if commercial.get("qty") is not None else meters)
            area = _r3(packing.get("area") or inv.get("area"))
            width = packing.get("width") or inv.get("width") or width
            hs = customs.get("hs_code") or inv.get("hs_code") or hs
            unit = commercial.get("unit") or inv.get("unit") or packing.get("unit") or unit
            price = commercial.get("price") or inv.get("price") or price
            amount = _r2(commercial.get("amount"))
            item_price = _as_float(commercial.get("price") or inv.get("price"))
            if fabric:
                rows.append(
                    [
                        int(no),
                        article_with_color(item.get("article"), commercial.get("color")),
                        hs,
                        rolls,
                        _as_float(width) if width is not None else None,
                        area,
                        meters,
                        unit or "meters",
                        None if is_family else item_price,
                        None if is_family else amount,
                    ]
                )
            else:
                rows.append(
                    [
                        int(no),
                        article_with_color(item.get("article"), commercial.get("color")),
                        hs,
                        rolls,
                        qty,
                        unit or "",
                        _as_float(price),
                        amount,
                    ]
                )
            group_rolls += float(rolls or 0)
            group_meters += float(meters or 0)
            group_qty += float(qty or 0)
            group_area += float(area or 0)
            group_amount += float(amount or 0)
            no += 1
        first = group_items[0]
        label_name = (first.get("article") or fam).split()[0] if fabric else (first.get("article") or "")
        group_design = f"{prefix}\n{label_name}".strip() if prefix else str(first.get("article") or "")
        group_amount = _r2(float(price) * group_meters) if fabric and price is not None else _r2(group_amount)
        if fabric and is_family:
            rows.append(
                [
                    None,
                    group_design,
                    hs,
                    _count(group_rolls),
                    _as_float(width) if width is not None else None,
                    _r3(group_area),
                    _r2(group_meters),
                    unit or "meters",
                    _as_float(price),
                    group_amount,
                ]
            )
        elif not fabric and prefix:
            rows.append(
                [
                    None,
                    group_design,
                    None,
                    _count(group_rolls),
                    _count(group_qty),
                    unit or "",
                    None,
                    group_amount,
                ]
            )
    total_rolls = _count(sum(float((i.get("packing_data") or {}).get("rolls") or 0) for i in products))
    total_area = _r3(sum(float((i.get("packing_data") or {}).get("area") or 0) for i in products))
    total_meters = _r2(sum(float((i.get("packing_data") or {}).get("meters") or 0) for i in products))
    total_qty = _count(
        sum(float((i.get("commercial_data") or {}).get("qty") or (i.get("packing_data") or {}).get("meters") or 0) for i in products)
    )
    total_amount = _r2(sum(float((i.get("commercial_data") or {}).get("amount") or 0) for i in products))
    if fabric:
        rows.append([None, "TOTAL:", None, total_rolls, None, total_area, total_meters, None, None, total_amount])
    else:
        rows.append([None, "TOTAL:", None, total_rolls, total_qty, None, None, total_amount])
    return rows


def packing_table_rows(products: list[dict[str, Any]], fabric: bool) -> list[list[Any]]:
    rows: list[list[Any]] = []
    sums = {"rolls": 0.0, "meters": 0.0, "qty": 0.0, "nw": 0.0, "gw": 0.0, "area": 0.0}
    for idx, pl in enumerate(_collect_packing_groups(products, fabric), start=1):
        rolls = _count(pl.get("rolls"))
        meters = _r2(pl.get("meters") if pl.get("meters") is not None else pl.get("qty"))
        qty = _count(pl.get("qty") if pl.get("qty") is not None else meters)
        nw = _r2(pl.get("net_weight"))
        gw = _r2(pl.get("gross_weight"))
        area = _r3(pl.get("area"))
        design = _format_packing_design(pl.get("design"), pl.get("article"))
        if fabric:
            rows.append([int(idx), design, _as_float(pl.get("gm")), rolls, meters, nw, gw, area])
        else:
            rows.append([int(idx), design, rolls, qty, pl.get("unit") or "", nw, gw])
        sums["rolls"] += float(rolls or 0)
        sums["meters"] += float(meters or 0)
        sums["qty"] += float(qty or 0)
        sums["nw"] += float(nw or 0)
        sums["gw"] += float(gw or 0)
        sums["area"] += float(area or 0)
    if fabric:
        rows.append(["TOTAL", None, None, _count(sums["rolls"]), _r2(sums["meters"]), _r2(sums["nw"]), _r2(sums["gw"]), _r3(sums["area"])])
    else:
        rows.append(["TOTAL", None, _count(sums["rolls"]), _count(sums["qty"]), None, _r2(sums["nw"]), _r2(sums["gw"])])
    return rows


def spec_table_rows(products: list[dict[str, Any]], fabric: bool) -> list[list[Any]]:
    rows: list[list[Any]] = []
    sums = {"rolls": 0.0, "meters": 0.0, "qty": 0.0, "area": 0.0, "nw": 0.0, "gw": 0.0, "amount": 0.0}
    for idx, item in enumerate(products, start=1):
        commercial = item.get("commercial_data") or {}
        packing = item.get("packing_data") or {}
        customs = item.get("customs_data") or {}
        rolls = _count(packing.get("rolls"))
        meters = _r2(packing.get("meters"))
        qty = _count(commercial.get("qty") if commercial.get("qty") is not None else meters)
        width = _as_float(packing.get("width"))
        area = _r3(packing.get("area"))
        nw = _r2(packing.get("net_weight"))
        gw = _r2(packing.get("gross_weight"))
        price = _as_float(commercial.get("price"))
        amount = _r2(commercial.get("amount"))
        hs = customs.get("hs_code") or ""
        tnved = customs.get("tnved_code") or hs
        if fabric:
            rows.append(
                [
                    int(idx),
                    rolls,
                    article_with_color(item.get("article"), commercial.get("color")),
                    _product_description(item),
                    hs,
                    tnved,
                    meters,
                    width,
                    area,
                    nw,
                    gw,
                    price,
                    amount,
                ]
            )
        else:
            rows.append(
                [
                    int(idx),
                    rolls,
                    article_with_color(item.get("article"), commercial.get("color")),
                    _product_description(item),
                    hs,
                    tnved,
                    qty,
                    commercial.get("unit") or packing.get("unit") or "",
                    area,
                    nw,
                    gw,
                    price,
                    amount,
                ]
            )
        sums["rolls"] += float(rolls or 0)
        sums["meters"] += float(meters or 0)
        sums["qty"] += float(qty or 0)
        sums["area"] += float(area or 0)
        sums["nw"] += float(nw or 0)
        sums["gw"] += float(gw or 0)
        sums["amount"] += float(amount or 0)
    if fabric:
        rows.append(
            [
                "Total/Итого:",
                _count(sums["rolls"]),
                None,
                None,
                None,
                None,
                _r2(sums["meters"]),
                None,
                _r3(sums["area"]),
                _r2(sums["nw"]),
                _r2(sums["gw"]),
                None,
                _r2(sums["amount"]),
            ]
        )
    else:
        rows.append(
            [
                "Total/Итого:",
                _count(sums["rolls"]),
                None,
                None,
                None,
                None,
                _count(sums["qty"]),
                None,
                _r3(sums["area"]),
                _r2(sums["nw"]),
                _r2(sums["gw"]),
                None,
                _r2(sums["amount"]),
            ]
        )
    return rows


def preview_headers(
    layout: str | None,
    items: list[dict[str, Any]] | None = None,
    header: dict[str, Any] | None = None,
) -> dict[str, list[str]]:
    fabric = is_fabric_layout(layout, items)
    ccy = currency_from_sources(items, header) or "CNY"
    return {
        "invoice": fabric_invoice_headers(ccy) if fabric else element_invoice_headers(ccy),
        "packing": FABRIC_PACKING_HEADERS if fabric else ELEMENT_PACKING_HEADERS,
        "specification": FABRIC_SPEC_HEADERS if fabric else ELEMENT_SPEC_HEADERS,
    }


def _set_cell(ws, row: int, col: int, value: Any, number_format: str | None = None) -> None:
    cell = ws.cell(row, col)
    fmt = number_format
    if type(value) is int:
        fmt = NUM_FMT_INT
    if isinstance(cell, MergedCell):
        for merged in ws.merged_cells.ranges:
            if cell.coordinate in merged:
                target = ws.cell(merged.min_row, merged.min_col)
                target.value = value
                if fmt and isinstance(value, (int, float)):
                    target.number_format = fmt
                return
        return
    cell.value = value
    if fmt and isinstance(value, (int, float)):
        cell.number_format = fmt


def _clear_block(ws, start_row: int, end_row: int, min_col: int, max_col: int) -> None:
    to_unmerge = []
    for merged in ws.merged_cells.ranges:
        if merged.max_row < start_row or merged.min_row > end_row:
            continue
        if merged.max_col < min_col or merged.min_col > max_col:
            continue
        to_unmerge.append(str(merged))
    for ref in to_unmerge:
        try:
            ws.unmerge_cells(ref)
        except KeyError:
            continue
    for r in range(start_row, end_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            try:
                cell.value = None
            except (AttributeError, KeyError):
                continue


def _find_header_row(ws, *needles: str) -> int | None:
    wanted = tuple(n.upper() for n in needles)
    for r in range(1, 40):
        texts = [str(ws.cell(r, c).value or "").strip().upper() for c in range(1, 16)]
        if any(any(needle == text or needle in text for text in texts) for needle in wanted):
            if any("DESIGN" in text or "ART." in text or "АРТИКУЛ" in text or text == "№" for text in texts):
                return r
    return None


def _table_colmap(ws, header_row: int, max_col: int = 22) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, max_col + 1):
        title = str(ws.cell(header_row, col).value or "").strip()
        if not title:
            continue
        low = title.lower()
        if low in {"no.", "no", "№", "nº", "n°"}:
            mapping["no"] = col
            continue
        if low == "design" or low.startswith("design"):
            mapping["design"] = col
            continue
        if "g/m" in low or low in {"g.m", "gsm"}:
            mapping["gm"] = col
            continue
        field = classify_header(title)
        if field and field not in mapping:
            mapping[field] = col
    return mapping


def _prepare_data_block(ws, data_start: int, n_rows: int, total_cols: tuple[int, ...], extra_cols: int = 20) -> int:
    total_row = None
    for r in range(data_start, ws.max_row + 1):
        for col in total_cols:
            if "total" in str(ws.cell(r, col).value or "").lower():
                total_row = r
                break
        if total_row:
            break
    max_col = max(ws.max_column or extra_cols, extra_cols)
    if total_row is None:
        _clear_block(ws, data_start, data_start + max(n_rows, 8) + 2, 1, max_col)
        return data_start + n_rows
    needed_before_total = max(n_rows - 1, 0)
    available = total_row - data_start
    if needed_before_total > available:
        try:
            ws.insert_rows(total_row, needed_before_total - available)
            total_row += needed_before_total - available
        except Exception:
            _clear_block(ws, data_start, data_start + n_rows + 2, 1, max_col)
            return data_start + n_rows - 1
    _clear_block(ws, data_start, total_row, 1, max_col)
    return total_row


def _fmt_for(field: str) -> str | None:
    if field in {"no"}:
        return NUM_FMT_INT
    if field in {"area"}:
        return NUM_FMT_3
    if field in {
        "rolls",
        "meters",
        "qty",
        "width",
        "price",
        "amount",
        "net_weight",
        "gross_weight",
        "gm",
    }:
        return NUM_FMT_2
    return None


def _write_mapped_row(ws, row: int, colmap: dict[str, int], values: dict[str, Any]) -> None:
    for field, value in values.items():
        col = colmap.get(field)
        if not col:
            continue
        _set_cell(ws, row, col, value, _fmt_for(field))


def _header_value(header: dict[str, Any] | None, *keys: str) -> str:
    for key in keys:
        value = (header or {}).get(key)
        if value not in (None, ""):
            return str(value)
    return ""


def _apply_letterhead(ws, header: dict[str, Any] | None, kit: str, header_row: int) -> None:
    inv = _header_value(header, "invoice_no") or f"ZFRMB-{kit}"
    date = _header_value(header, "invoice_date", "date")
    contract = _header_value(header, "contract_no")
    buyer = _header_value(header, "buyer")
    buyer_addr = _header_value(header, "buyer_address")
    seller = _header_value(header, "seller")
    delivery = _header_value(header, "delivery_terms")
    payment = _header_value(header, "payment_terms")
    delivery_date = _header_value(header, "delivery_date")
    manufacturer = _header_value(header, "manufacturer")
    ccy_label = export_currency_label(_header_value(header, "currency") or "CNY", hangzhou_style=True)
    for r in range(1, max(header_row, 22)):
        for c in range(1, 16):
            val = ws.cell(r, c).value
            if val is None:
                continue
            text = str(val)
            stripped = text.strip()
            upper = stripped.upper()
            if isinstance(val, str) and val.lstrip().startswith("=") and "SPECIFICATION" in val.upper() and inv:
                ws.cell(r, c).value = f"SPECIFICATION №{inv} dd {date}".strip() if date else f"SPECIFICATION №{inv}"
                continue
            if _INVOICE_TOKEN_RE.search(text) and inv:
                ws.cell(r, c).value = _INVOICE_TOKEN_RE.sub(inv, text)
                text = str(ws.cell(r, c).value)
            if text.upper().startswith("SPECIFICATION") and inv:
                if date:
                    ws.cell(r, c).value = f"SPECIFICATION №{inv} dd {date}"
                else:
                    ws.cell(r, c).value = f"SPECIFICATION №{inv}"
            if upper in {"INV.NO.", "INV.NO"}:
                ws.cell(r, c + 1).value = inv
            if upper in {"DATE:", "ДАТА:"} and date:
                ws.cell(r, c + 1).value = date
            if "инв номер" in text.lower() and inv:
                ws.cell(r, c + 1).value = inv
            if contract and upper.startswith("CONTRACT") and "№" in text:
                if "dated" in text.lower() or "dd" in text.lower() or "контракт" in text.lower():
                    keep_date = ""
                    dated = re.search(r"(dated|dd|от)\s*.+$", text, re.I)
                    if dated:
                        keep_date = " " + dated.group(0)
                    ws.cell(r, c).value = f"Contract/Контракт №：{contract}{keep_date}" if "контракт" in text.lower() else f"Contract №：{contract}{keep_date}"
            if buyer and upper in {"BUYER:", "BUYER/ПОКУПАТЕЛЬ:"}:
                ws.cell(r + 1, c).value = buyer
            if buyer_addr and upper in {"ADDRESS:", "ADDRESS/АДРЕС:"} and c <= 4:
                ws.cell(r + 1, c).value = buyer_addr
            if seller and upper in {"SELLER:", "SELLER / ПРОДАВЕЦ:", "SELLER/ПРОДАВЕЦ:"}:
                ws.cell(r + 1, c).value = seller
            if payment and text.lower().startswith("terms of payment"):
                ws.cell(r, c).value = re.sub(r"(?i)terms of payment\s*:?\s*.*", f"Terms of payment: {payment}", text)
            if delivery and text.lower().startswith("terms of delivery"):
                ws.cell(r, c).value = re.sub(r"(?i)terms of delivery\s*:?\s*.*", f"Terms of delivery: {delivery}", text)
            if delivery_date and text.lower().startswith("delivery date"):
                ws.cell(r, c).value = re.sub(r"(?i)delivery date\s*:?\s*.*", f"Delivery date: {delivery_date}", text)
            if manufacturer and text.lower().startswith("manufacturer"):
                ws.cell(r, c).value = f"Manufacturer: {manufacturer}"
            # Rewrite hardcoded etalon currency labels to the shipment currency.
            if _PRICE_CCY_RE.search(stripped):
                ws.cell(r, c).value = _PRICE_CCY_RE.sub(
                    lambda m: f"{m.group(1).upper()}({ccy_label})",
                    stripped,
                )


def fill_specification_template(
    template: Path,
    output: Path,
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    kit: str,
) -> Path:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb[wb.sheetnames[0]]
    products = _ordered_products(items)
    fabric = is_fabric_layout(_layout_kit(kit, products, header), products)
    rows = spec_table_rows(products, fabric)
    header_row = _find_header_row(ws, "№", "ART", "PRODUCT") or 22
    data_start = header_row + 1
    colmap = _table_colmap(ws, header_row)
    total_row = _prepare_data_block(ws, data_start, len(rows), (2, 1, colmap.get("no", 2)))
    field_order = [
        "no",
        "rolls",
        "article",
        "description",
        "hs_code",
        "tnved_code",
        "meters" if fabric else "qty",
        "width" if fabric else "unit",
        "area",
        "net_weight",
        "gross_weight",
        "price",
        "amount",
    ]
    for offset, row in enumerate(rows):
        values = {field: row[idx] for idx, field in enumerate(field_order) if idx < len(row)}
        if str(row[0]).lower().startswith("total"):
            values["no"] = row[0]
            _write_mapped_row(ws, data_start + offset, colmap, values)
            if colmap.get("article"):
                _set_cell(ws, data_start + offset, colmap["article"], None)
        else:
            _write_mapped_row(ws, data_start + offset, colmap, values)
        extra = {
            "country": (header or {}).get("country") or ((products[offset].get("customs_data") if offset < len(products) else {}) or {}).get("country") or "КИТАЙ",
            "manufacturer": ((products[offset].get("customs_data") if offset < len(products) else {}) or {}).get("manufacturer")
            or (header or {}).get("manufacturer")
            or "",
        }
        if offset < len(products):
            _write_mapped_row(ws, data_start + offset, colmap, extra)
            if colmap.get("article") and colmap.get("description"):
                # keep TM / repeat columns used by some broker files
                for col in range(max(colmap.values(), default=14) + 1, min((ws.max_column or 20) + 1, 21)):
                    title = str(ws.cell(header_row, col).value or "").lower()
                    if "country" in title or "страна" in title:
                        _set_cell(ws, data_start + offset, col, extra["country"])
                    elif "manufact" in title or "производ" in title:
                        _set_cell(ws, data_start + offset, col, extra["manufacturer"])
                    elif "наименов" in title or "product" in title:
                        _set_cell(ws, data_start + offset, col, _product_description(products[offset]))
                    elif "артикул" in title or title.startswith("art"):
                        _set_cell(ws, data_start + offset, col, products[offset].get("article"))
                    elif "hs" in title:
                        _set_cell(ws, data_start + offset, col, (products[offset].get("customs_data") or {}).get("hs_code"))
    if data_start + len(rows) - 1 != total_row:
        # total already written as last data row
        pass
    _apply_letterhead(ws, header, kit, header_row)

    if len(wb.sheetnames) > 1:
        desc_ws = wb[wb.sheetnames[1]]
        d_start = 13
        for r in range(1, 20):
            if str(desc_ws.cell(r, 4).value or "").lower().startswith("art"):
                d_start = r + 1
                break
        _clear_block(desc_ws, d_start, d_start + max(40, len(products) + 2), 2, 15)
        for idx, item in enumerate(products, start=1):
            row = d_start + idx - 1
            commercial = item.get("commercial_data") or {}
            packing = item.get("packing_data") or {}
            customs = item.get("customs_data") or {}
            _set_cell(desc_ws, row, 2, int(idx), NUM_FMT_INT)
            _set_cell(desc_ws, row, 3, _count(packing.get("rolls")), NUM_FMT_INT)
            _set_cell(desc_ws, row, 4, item.get("article"))
            _set_cell(desc_ws, row, 5, _product_description(item))
            _set_cell(desc_ws, row, 6, customs.get("hs_code"))
            _set_cell(desc_ws, row, 7, customs.get("tnved_code") or customs.get("hs_code"))
            if fabric:
                _set_cell(desc_ws, row, 8, _r2(packing.get("meters")), NUM_FMT_2)
            else:
                _set_cell(desc_ws, row, 8, _count(commercial.get("qty")), NUM_FMT_INT)
            _set_cell(desc_ws, row, 9, packing.get("width") if fabric else commercial.get("unit"), NUM_FMT_2 if fabric else None)
            _set_cell(desc_ws, row, 10, _r3(packing.get("area")), NUM_FMT_3)
            _set_cell(desc_ws, row, 11, _r2(packing.get("net_weight")), NUM_FMT_2)
            _set_cell(desc_ws, row, 12, _r2(packing.get("gross_weight")), NUM_FMT_2)
            _set_cell(desc_ws, row, 13, _as_float(commercial.get("price")), NUM_FMT_2)
            _set_cell(desc_ws, row, 14, _r2(commercial.get("amount")), NUM_FMT_2)

    unfreeze_workbook(wb)
    wb.save(output)
    return output


def fill_invoice_template(
    template: Path,
    output: Path,
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    kit: str,
    packing_groups: list[dict[str, Any]] | None = None,
) -> Path:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb.active
    products = _ordered_products(items)
    fabric = is_fabric_layout(_layout_kit(kit, products, header), products)
    rows = invoice_table_rows(products, fabric)
    header_row = _find_header_row(ws, "NO.", "DESIGN") or 24
    data_start = header_row + 1
    colmap = _table_colmap(ws, header_row)
    _prepare_data_block(ws, data_start, len(rows), (1, 2, colmap.get("design", 2), colmap.get("no", 1)))
    field_order = (
        ["no", "design", "hs_code", "rolls", "width", "area", "meters", "unit", "price", "amount"]
        if fabric
        else ["no", "design", "hs_code", "rolls", "qty", "unit", "price", "amount"]
    )
    for offset, row in enumerate(rows):
        values = {field: row[idx] for idx, field in enumerate(field_order) if idx < len(row)}
        _write_mapped_row(ws, data_start + offset, colmap, values)
        if str(values.get("design") or "").upper().startswith("TOTAL"):
            cell = ws.cell(data_start + offset, colmap.get("design", 2))
            cell.font = Font(bold=True)
    _apply_letterhead(ws, header, kit, header_row)
    unfreeze_workbook(wb)
    wb.save(output)
    return output


def fill_packing_template(
    template: Path,
    output: Path,
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    kit: str,
) -> Path:
    shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb.active
    products = _ordered_products(items)
    fabric = is_fabric_layout(_layout_kit(kit, products, header), products)
    rows = packing_table_rows(products, fabric)
    header_row = _find_header_row(ws, "NO.", "NO", "DESIGN") or 9
    data_start = header_row + 1
    colmap = _table_colmap(ws, header_row)
    _prepare_data_block(ws, data_start, len(rows), (1, 2, colmap.get("no", 1), colmap.get("design", 2)))
    field_order = (
        ["no", "design", "gm", "rolls", "meters", "net_weight", "gross_weight", "area"]
        if fabric
        else ["no", "design", "rolls", "qty", "unit", "net_weight", "gross_weight"]
    )
    for offset, row in enumerate(rows):
        values = {field: row[idx] for idx, field in enumerate(field_order) if idx < len(row)}
        if str(row[0] or "").upper() == "TOTAL":
            values["no"] = "TOTAL"
            values["design"] = None
        _write_mapped_row(ws, data_start + offset, colmap, values)
    _apply_letterhead(ws, header, kit, header_row)
    unfreeze_workbook(wb)
    wb.save(output)
    return output


def export_18233_from_templates(
    items: list[dict[str, Any]],
    output_dir: Path,
    header: dict[str, Any] | None = None,
    shipment_title: str = "18233",
) -> list[Path]:
    if not materials_available():
        raise FileNotFoundError("MVP_18233 materials not found; cannot use эталон templates")

    header = export_header_fields(header or {}, items)
    kit = _detect_kit(items, header)
    layout = _layout_kit(kit, items, header)
    templates = kit_templates(layout)
    output_dir.mkdir(parents=True, exist_ok=True)
    label = kit or layout
    stem = safe_export_stem(shipment_title)

    inv_out = output_dir / f"{stem} ИНВОЙС {label}.xlsx"
    pl_out = output_dir / f"{stem} ПАКИНГ {label}.xlsx"
    spec_out = output_dir / f"{stem} {label} СПЕЦИФИКАЦИЯ С ЦВЕТАМИ.xlsx"

    paths = [
        fill_invoice_template(templates["invoice"], inv_out, items, header, label),
        fill_packing_template(templates["packing"], pl_out, items, header, label),
        fill_specification_template(templates["specification"], spec_out, items, header, label),
    ]
    return paths
