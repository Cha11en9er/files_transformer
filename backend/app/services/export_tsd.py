"""Commercial TSD book: invoice/packing/spec as printed, not the Goldluck ED layout.

Used when goods already carry HS + manufacturer + packages and almost no color
column — the same shape as a finished-goods invoice/packing/specification PDF.
18233 fabric kits and Beijing Goldluck (color/volume/CNY) keep their own writers.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.parsing.header_extract import currency_from_sources, export_header_fields, split_party_address
from app.services.bilingual import bilingual, label
from app.services.column_layout import apply_sheet_layout
from app.services.export_style import (
    style_data_table,
    style_letterhead_row,
    style_split_letterhead,
    unfreeze_workbook,
)
from app.services.field_map import is_factory_note, parse_number

_ROOT = Path(__file__).resolve().parents[3]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))
from prepare_transform_code.fields import is_factory_list

def _party_address(text: str | None, *, seller: bool) -> str:
    return split_party_address(text, seller=seller) or ""


NUM_FMT = "0.00"
NUM_FMT_4 = "0.0000"
NUM_FMT_INT = "0"


def is_tsd_layout(items: list[dict[str, Any]]) -> bool:
    """Штучный бланк ТСД: CODE + страна/шт, без метров ткани и без Goldluck-цвета/объёма.

    Примеры: Bestway, шары, матрасы, гайки Necargo, светильники. Не Hangzhou-ткань
    (метры+ширина) и не Beijing Goldluck (цвет+CBM на строке).
    """
    products = [item for item in items if item.get("article")]
    if len(products) < 2:
        return False
    n = len(products)
    hs = 0
    mfr = 0
    country = 0
    color = 0
    volume = 0
    fabric = 0
    for item in products:
        customs = item.get("customs_data") or {}
        commercial = item.get("commercial_data") or {}
        packing = item.get("packing_data") or {}
        if customs.get("tnved_code") or customs.get("hs_code"):
            hs += 1
        if customs.get("manufacturer"):
            mfr += 1
        if customs.get("country"):
            country += 1
        if commercial.get("color"):
            color += 1
        if packing.get("volume") not in (None, ""):
            volume += 1
        if packing.get("meters") not in (None, "") or packing.get("width") not in (None, ""):
            fabric += 1
    if fabric >= max(2, int(n * 0.3)):
        return False
    if color > n * 0.25 or volume > n * 0.25:
        return False
    if hs < n * 0.6:
        return False
    # Завод на строке или страна на строке — достаточно, чтобы не уйти в тканевые столбцы.
    return mfr >= n * 0.4 or country >= n * 0.4


def _sum_present(body: list[list[Any]], index: int) -> float | None:
    nums = [float(row[index]) for row in body if isinstance(row[index], (int, float))]
    if not nums:
        return None
    return _fnum(sum(nums))


def _fnum(value: Any, digits: int = 2) -> float | None:
    if value in (None, ""):
        return None
    number = parse_number(value)
    if number is None:
        return None
    return round(float(number), digits)


def _count(value: Any) -> int | float | None:
    number = _fnum(value, 2)
    if number is None:
        return None
    if abs(number - round(number)) < 1e-9:
        return int(round(number))
    return number


def _qty(item: dict[str, Any]) -> Any:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    if commercial.get("qty") not in (None, ""):
        return commercial.get("qty")
    return packing.get("meters")


def _unit(item: dict[str, Any]) -> str:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    # Единицу не выдумывать: пустая клетка лучше ложного «ШТ».
    return str(commercial.get("unit") or packing.get("unit") or "")


def _packages(item: dict[str, Any]) -> Any:
    packing = item.get("packing_data") or {}
    return packing.get("rolls") or packing.get("boxes")


def _currency(items: list[dict[str, Any]], header: dict[str, Any] | None) -> str:
    return currency_from_sources(items, header) or "USD"


def _hs(item: dict[str, Any]) -> str:
    customs = item.get("customs_data") or {}
    return str(customs.get("tnved_code") or customs.get("hs_code") or "")


def _desc(item: dict[str, Any]) -> str:
    customs = item.get("customs_data") or {}
    en = customs.get("description_en")
    ru = customs.get("description_ru")
    if en and ru and not is_factory_note(en) and not is_factory_note(ru):
        return f"{en} / {ru}"
    for key in ("description", "description_en", "description_ru"):
        value = customs.get(key)
        if value and not is_factory_note(value):
            return str(value)
    return ""


def _ru_desc(item: dict[str, Any]) -> str:
    customs = item.get("customs_data") or {}
    ru = customs.get("description_ru")
    if ru and not is_factory_note(ru):
        return str(ru)
    blob = _desc(item)
    if " / " in blob:
        return blob.split(" / ", 1)[-1]
    return blob


def _mfr(item: dict[str, Any], header: dict[str, Any] | None) -> str:
    """Завод строки. Список всех заводов поставки в клетку не подставляется.
    Колонка MANUFACTURER / BRAND: до слэша завод, после слэша марка."""
    customs = item.get("customs_data") or {}
    row = str(customs.get("manufacturer") or "").strip()
    head = str((header or {}).get("manufacturer") or "").strip()
    if is_factory_list(row):
        row = ""
    if is_factory_list(head) or " / " in head:
        head = ""
    name = row or head
    brand = str(customs.get("brand") or "").strip()
    if brand and brand.casefold() not in name.casefold():
        return f"{name} / {brand}".strip(" /")
    return name


def _country(item: dict[str, Any], header: dict[str, Any] | None = None) -> str:
    customs = item.get("customs_data") or {}
    row = str(customs.get("country") or "").strip()
    if row:
        return row
    return str((header or {}).get("country") or "").strip()


def _net(item: dict[str, Any]) -> Any:
    return (item.get("packing_data") or {}).get("net_weight")


def _net_primary(item: dict[str, Any]) -> Any:
    return (item.get("packing_data") or {}).get("net_primary")


def _gross(item: dict[str, Any]) -> Any:
    return (item.get("packing_data") or {}).get("gross_weight")


def _price(item: dict[str, Any]) -> Any:
    return (item.get("commercial_data") or {}).get("price")


def _amount(item: dict[str, Any]) -> Any:
    return (item.get("commercial_data") or {}).get("amount")


def _append_table(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    header_row = ws.max_row
    start = ws.max_row + 1
    price_cols = {
        i + 1
        for i, title in enumerate(headers)
        if "price" in title.lower() or "цена" in title.lower() or title.upper().startswith("PRICE")
    }
    for row in rows:
        ws.append(row)
    for r in range(start, ws.max_row + 1):
        for cell in ws[r]:
            if type(cell.value) is int:
                cell.number_format = NUM_FMT_INT
            elif isinstance(cell.value, float):
                cell.number_format = NUM_FMT_4 if cell.column in price_cols else NUM_FMT
            if isinstance(cell.value, str) and cell.value.upper().startswith(("TOTAL", "ИТОГО")):
                cell.font = Font(bold=True)
    style_data_table(
        ws,
        header_row=header_row,
        cols=len(headers),
        n_rows=1 + len(rows),
        headers=headers,
    )


def _header_parties(header: dict[str, Any]) -> tuple[str, str, str, str]:
    header = header or {}
    mixed = header.get("seller_address") or ""
    seller_address = _party_address(mixed, seller=True)
    buyer_address = _party_address(header.get("buyer_address") or mixed, seller=False)
    return (
        header.get("seller") or "",
        seller_address,
        header.get("buyer") or "",
        buyer_address,
    )


def _write_invoice_letterhead(ws, header: dict[str, Any], cols: int) -> None:
    header = header or {}
    seller, seller_address, buyer, buyer_address = _header_parties(header)
    invoice_no = header.get("invoice_no") or ""
    date = header.get("invoice_date") or header.get("date") or ""
    contract = header.get("contract_no") or ""
    contract_date = header.get("contract_date") or ""
    delivery = header.get("delivery_terms") or ""
    if seller:
        ws.append([seller])
        style_letterhead_row(ws, ws.max_row, cols, company=True)
    ws.append([])
    split = min(7, cols)
    pad = [None] * max(split - 2, 0)
    ws.append([label("invoice", colon=True), *pad, invoice_no])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split)
    ws.append([label("date", colon=True), *pad, date])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split)
    ws.append([])
    ws.append([label("the_seller", colon=True)])
    style_letterhead_row(ws, ws.max_row, cols)
    ws.append([seller, *pad, label("delivery_basis")])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split)
    contract_line = f"{label('contract_word', colon=True)} {contract}".strip()
    if contract_date:
        contract_line = f"{contract_line} dd {contract_date}"
    ws.append([f"{label('address', colon=True)} {seller_address}".strip(), *pad, contract_line])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    ws.row_dimensions[ws.max_row].height = 36
    spec_line = f"{bilingual('Specifications from', 'Спецификация к')} {invoice_no}".strip()
    if date:
        spec_line = f"{spec_line} dd {date}"
    ws.append([None, *pad, spec_line])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    ws.append([])
    recipient = header.get("consignee") or header.get("recipient") or ""
    recipient_address = header.get("consignee_address") or header.get("recipient_address") or ""
    ws.append([label("the_buyer", colon=True), *pad, label("recipient", colon=True)])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split)
    ws.append([buyer, *pad, recipient])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    ws.append([f"{label('address', colon=True)} {buyer_address}".strip(), *pad, f"{label('address', colon=True)} {recipient_address}".strip()])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    ws.row_dimensions[ws.max_row].height = 36
    if delivery:
        ws.append([f"{label('delivery', colon=True)} {delivery}"])
        style_letterhead_row(ws, ws.max_row, cols)
    payment = header.get("payment_terms") or ""
    if payment:
        ws.append([f"{label('payment', colon=True)} {payment}"])
        style_letterhead_row(ws, ws.max_row, cols)
    ws.append([])


def _write_packing_letterhead(ws, header: dict[str, Any], cols: int) -> None:
    header = header or {}
    seller, seller_address, buyer, buyer_address = _header_parties(header)
    invoice_no = header.get("invoice_no") or ""
    date = header.get("invoice_date") or header.get("date") or ""
    contract = header.get("contract_no") or ""
    contract_date = header.get("contract_date") or ""
    delivery = header.get("delivery_terms") or ""
    container = header.get("container_no") or header.get("container") or ""
    if seller:
        ws.append([seller])
        style_letterhead_row(ws, ws.max_row, cols, company=True)
    ws.append([label("packing")])
    style_letterhead_row(ws, ws.max_row, cols, title=True)
    ws.append([f"{bilingual('To invoice', 'К инвойсу')}: {invoice_no}".strip()])
    style_letterhead_row(ws, ws.max_row, cols)
    ws.append([f"{label('date', colon=True)} {date}".strip()])
    style_letterhead_row(ws, ws.max_row, cols)
    ws.append([])
    split = min(5, cols)
    recipient = header.get("consignee") or header.get("recipient") or ""
    recipient_address = header.get("consignee_address") or header.get("recipient_address") or ""
    ws.append([label("the_seller", colon=True), None, label("the_buyer", colon=True), None, label("recipient", colon=True)])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split)
    ws.append([seller, None, buyer, None, recipient])
    style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    ws.append(
        [
            f"{label('address', colon=True)} {seller_address}".strip(),
            None,
            f"{label('address', colon=True)} {buyer_address}".strip(),
            None,
            f"{label('address', colon=True)} {recipient_address}".strip(),
        ]
    )
    style_split_letterhead(ws, ws.max_row, cols, split_at=split, bold=False)
    ws.row_dimensions[ws.max_row].height = 48
    contract_line = f"{label('contract_word', colon=True)} {contract}".strip()
    if contract_date:
        contract_line = f"{contract_line} dd {contract_date}"
    ws.append([contract_line])
    style_letterhead_row(ws, ws.max_row, cols)
    if container:
        ws.append([f"{label('container', colon=True)} {container}"])
        style_letterhead_row(ws, ws.max_row, cols)
    if delivery:
        ws.append([f"{label('delivery', colon=True)} {delivery}"])
        style_letterhead_row(ws, ws.max_row, cols)
    payment = header.get("payment_terms") or ""
    if payment:
        ws.append([f"{label('payment', colon=True)} {payment}"])
        style_letterhead_row(ws, ws.max_row, cols)
    ws.append([])


def _write_spec_letterhead(ws, header: dict[str, Any], cols: int) -> None:
    header = header or {}
    invoice_no = header.get("invoice_no") or ""
    contract = header.get("contract_no") or ""
    contract_date = header.get("contract_date") or ""
    date = header.get("invoice_date") or header.get("date") or ""
    buyer = header.get("buyer") or ""
    seller = header.get("seller") or ""
    ws.append([f"{bilingual('Annex No.', 'Приложение №')} {invoice_no}".strip()])
    style_letterhead_row(ws, ws.max_row, cols, title=True)
    contract_line = f"{bilingual('To the contract No.', 'к контракту №')} {contract}".strip()
    if contract_date:
        contract_line = f"{contract_line} {bilingual('dated', 'от')} {contract_date}"
    ws.append([contract_line])
    style_letterhead_row(ws, ws.max_row, cols)
    if date:
        ws.append([date])
        style_letterhead_row(ws, ws.max_row, cols)
    if buyer or seller:
        ws.append([f"{buyer} / {seller}".strip(" /")])
        style_letterhead_row(ws, ws.max_row, cols)
        ws.row_dimensions[ws.max_row].height = 36
    ws.append(
        [
            bilingual(
                "1. The Seller sells, the Buyer buys, and the Recipient accepts the goods according to the following specification:",
                "1. Продавец продает, Покупатель покупает, а Получатель принимает товар согласно следующей спецификации:",
            )
        ]
    )
    style_letterhead_row(ws, ws.max_row, cols)
    ws.append([])


def invoice_headers(ccy: str) -> list[str]:
    return [
        bilingual("No.", "№"),
        "Customs code / Таможенный код",
        "DESCRIPTION / Описание",
        "COUNTRY OF ORIGIN / Страна происхождения",
        "MODEL / SERIES / ART. / Модель, серия, арт.",
        "MANUFACTURER / BRAND / Производитель",
        "WEIGHT NETTO, kg / Вес нетто, кг",
        "NETTO WITH PRIMARY PACKAGING, kg / Нетто с первичной упаковкой, кг",
        "QTY / Кол-во",
        f"PRICE PER {ccy} / Цена за ед., {ccy}",
        "pcs, pcg, set / шт, уп, компл",
        f"AMOUNT, {ccy} / Сумма, {ccy}",
    ]


def packing_headers() -> list[str]:
    return [
        bilingual("No.", "№"),
        "Customs code / Таможенный код",
        "DESCRIPTION / Описание",
        "MANUFACTURER / Производитель",
        "MODEL / SERIES / ART. / Модель, серия, арт.",
        "PACKAGE / Места",
        "QTY / Кол-во",
        "WEIGHT NETTO, kg / Вес нетто, кг",
        "WEIGHT NETTO WITH PRIMARY PACKAGING, kg / Нетто с первичной упаковкой, кг",
        "WEIGHT BRUTTO, kg / Вес брутто, кг",
        "COUNTRY OF ORIGIN / Страна происхождения",
    ]


def spec_headers(ccy: str) -> list[str]:
    return [
        "НАИМЕНОВАНИЕ ТОВАРА / Description of goods",
        "МОДЕЛЬ, СЕРИЯ, АРТ. / Model, series, art.",
        "ФИРМА ПРОИЗВ-ЛЬ / Manufacturer",
        "СТРАНА ПРОИСХ. / Country of origin",
        "КОЛ-ВО / Quantity",
        "ЕД.ИЗМ. / Unit",
        "ВЕС БРУТТО, КГ / Gross weight, kg",
        "ВЕС НЕТТО, КГ / Net weight, kg",
        "ВЕС НЕТТО С ПЕРВИЧНОЙ УПАКОВКОЙ, КГ / Net with primary packing, kg",
        f"СТОИМОСТЬ, {ccy} / Amount, {ccy}",
        "ЕД.ИЗМ. / Unit",
        f"СТ-СТЬ {ccy} / Unit price, {ccy}",
        "РАЗМЕР / Size",
    ]


def invoice_rows(items: list[dict[str, Any]], header: dict[str, Any] | None) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for idx, item in enumerate(items, start=1):
        net = _fnum(_net(item))
        primary = _fnum(_net_primary(item))
        rows.append(
            [
                idx,
                _hs(item),
                _desc(item),
                _country(item, header),
                item.get("article") or "",
                _mfr(item, header),
                net,
                primary,
                _count(_qty(item)),
                _fnum(_price(item), 4),
                _unit(item),
                _fnum(_amount(item)),
            ]
        )
    if rows:
        body = rows[:]
        rows.append(
            [
                "TOTAL :",
                None,
                None,
                None,
                None,
                None,
                _fnum(sum(float(r[6] or 0) for r in body)),
                _sum_present(body, 7),
                _count(sum(float(r[8] or 0) for r in body)),
                None,
                None,
                _fnum(sum(float(r[11] or 0) for r in body)),
            ]
        )
    return rows


def packing_rows(items: list[dict[str, Any]], header: dict[str, Any] | None) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for idx, item in enumerate(items, start=1):
        net = _fnum(_net(item))
        primary = _fnum(_net_primary(item))
        rows.append(
            [
                idx,
                _hs(item),
                _desc(item),
                _mfr(item, header),
                item.get("article") or "",
                _count(_packages(item)),
                _count(_qty(item)),
                net,
                primary,
                _fnum(_gross(item)),
                _country(item, header),
            ]
        )
    if rows:
        body = rows[:]
        rows.append(
            [
                "TOTAL",
                None,
                None,
                None,
                None,
                _count(sum(float(r[5] or 0) for r in body)),
                _count(sum(float(r[6] or 0) for r in body)),
                _fnum(sum(float(r[7] or 0) for r in body)),
                _sum_present(body, 8),
                _fnum(sum(float(r[9] or 0) for r in body)),
                None,
            ]
        )
    return rows


def spec_rows(items: list[dict[str, Any]], header: dict[str, Any] | None) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for item in items:
        net = _fnum(_net(item))
        primary = _fnum(_net_primary(item))
        rows.append(
            [
                _desc(item),
                item.get("article") or "",
                _mfr(item, header),
                _country(item, header),
                _count(_qty(item)),
                _unit(item),
                _fnum(_gross(item)),
                net,
                primary,
                _fnum(_amount(item)),
                _unit(item),
                _fnum(_price(item), 4),
                (item.get("commercial_data") or {}).get("size") or None,
            ]
        )
    if rows:
        body = rows[:]
        rows.append(
            [
                "ИТОГО:",
                None,
                None,
                None,
                _count(sum(float(r[4] or 0) for r in body)),
                None,
                _fnum(sum(float(r[6] or 0) for r in body)),
                _fnum(sum(float(r[7] or 0) for r in body)),
                _sum_present(body, 8),
                _fnum(sum(float(r[9] or 0) for r in body)),
                None,
                None,
                None,
            ]
        )
    return rows


def description_headers() -> list[str]:
    return [
        bilingual("No.", "№"),
        bilingual("HS Code", "Код ТН ВЭД"),
        bilingual("Manufacturer / Brand", "Изготовитель / Торговая марка"),
        bilingual("Model / Article", "Модель / артикул"),
        bilingual("Description", "Описание"),
    ]


def description_rows(items: list[dict[str, Any]], header: dict[str, Any] | None) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for idx, item in enumerate(items, start=1):
        rows.append(
            [
                idx,
                _hs(item),
                _mfr(item, header),
                item.get("article") or "",
                _ru_desc(item) or _desc(item),
            ]
        )
    return rows


# Лист ДЛЯ ДТ — шапка как в documents_final/Документы_OI-451.xlsx.
# Безымянные клетки (Y1/Z1) не выводим. Пустые по значениям столбцы режем при записи.
# Исключение: «цена со скидкой» остаётся даже пустой.
_DT_COUNT = {"кол-во товара", "кол-во мест", "Кол-во упак"}
_DT_MONEY = {"брутто", "нетто", "нетто без упаковки", "ст-ть товара", "цена за ед. товара", "цена со скидкой"}
_DT_KEEP_ALWAYS = {"no", "price_discounted"}

# Порядок и имена как на бланке OI-451 «ДЛЯ ДТ» (без пустых заголовков).
_DT_COLUMN_SPECS: list[tuple[str, str]] = [
    ("no", "№"),
    ("hs_code", "Код ТН ВЭД"),
    ("description", "Описание"),
    ("group_description", "Описание в группе"),
    ("manufacturer", "Изготовитель"),
    ("trade_mark", "ТЗ"),
    ("article", "артикул"),
    ("brand", "марка"),
    ("model", "модель"),
    ("serial", "серийный номер"),
    ("size", "размер"),
    ("qty", "кол-во товара"),
    ("unit", "Ед.изм."),
    ("price", "цена за ед. товара"),
    ("packages", "кол-во мест"),
    ("extra_unit", "доп.единица"),
    ("packing_presence", "Код наличия упаковки"),
    ("packing_code", "Код упаковки"),
    ("packing_count", "Кол-во упак"),
    ("container", "контейнер"),
    ("pallets", "поддоны кол-во"),
    ("container_numbers", "номера контейнеры"),
    ("fill_mark", "признак заполнения"),
    ("country", "страна происхождения"),
    ("gross", "брутто"),
    ("net", "нетто"),
    ("net_bare", "нетто без упаковки"),
    ("amount", "ст-ть товара"),
    ("invoice_no", "номер инвойса"),
    ("invoice_date", "дата инвойса"),
    ("invoice_currency", "Валюта инвойса"),
    ("description_short", "Описание"),
    ("its_request", "ИТС по запросу"),
    ("its_price_net", "ИТС цена/нетто"),
    ("declaration", "Декларация соответствия"),
    ("price_discounted", "цена со скидкой"),
]


def dt_headers() -> list[str]:
    return [title for _key, title in _DT_COLUMN_SPECS]


def dt_keys() -> list[str]:
    return [key for key, _title in _DT_COLUMN_SPECS]


def _plain(value: Any) -> str:
    return str(value or "").strip()


def _dt_producer(item: dict[str, Any], header: dict[str, Any] | None) -> str:
    """Завод своей строки. Список заводов и марка в эту клетку не клеятся."""
    row = _plain((item.get("customs_data") or {}).get("manufacturer"))
    if row and not is_factory_list(row):
        return row
    head = _plain((header or {}).get("manufacturer"))
    if not head or is_factory_list(head) or "/" in head:
        return ""
    return head


def _dt_article(item: dict[str, Any]) -> str:
    text = _plain(item.get("article"))
    return "" if text == "-" else text


def _dt_unit(item: dict[str, Any]) -> str:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    return _plain(commercial.get("unit") or packing.get("unit"))


def _dt_price_discounted(item: dict[str, Any]) -> float | None:
    """Цена/сумма со скидкой, если скидка напечатана. Иначе пусто — столбец всё равно оставляем."""
    commercial = item.get("commercial_data") or {}
    if commercial.get("price_discounted") not in (None, ""):
        return _fnum(commercial.get("price_discounted"), 4)
    discount = commercial.get("discount")
    amount = commercial.get("amount")
    if discount not in (None, "") and amount not in (None, ""):
        try:
            return _fnum(float(amount) - float(discount))
        except (TypeError, ValueError):
            return None
    return None


def _dt_values(item: dict[str, Any], header: dict[str, Any] | None, index: int) -> dict[str, Any]:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    customs = item.get("customs_data") or {}
    header = header or {}
    places = _count(_packages(item))
    pallets = packing.get("pallets")
    net = _fnum(_net(item))
    primary = _fnum(_net_primary(item))
    desc = _desc(item) or None
    short = _plain(customs.get("description_ru") or customs.get("description")) or None
    if short and desc and short == desc:
        short = None
    return {
        "no": index,
        "hs_code": _hs(item) or None,
        "description": desc,
        "group_description": _plain(customs.get("group_description")) or None,
        "manufacturer": _dt_producer(item, header) or None,
        "trade_mark": _plain(customs.get("brand")) or None,
        "article": _dt_article(item) or None,
        # «марка» отдельным столбцом только если есть своё поле; ТЗ уже держит brand.
        "brand": _plain(customs.get("mark") or customs.get("trade_mark_alt")) or None,
        "model": _plain(commercial.get("model")) or None,
        "serial": _plain(commercial.get("serial") or packing.get("serial")) or None,
        "size": _plain(commercial.get("size")) or None,
        "qty": _count(_qty(item)),
        "unit": _dt_unit(item) or None,
        "price": _fnum(_price(item), 4),
        "packages": places,
        "extra_unit": _plain(commercial.get("unit_alt") or packing.get("extra_unit")) or None,
        "packing_presence": _plain(packing.get("packing_presence")) or None,
        "packing_code": _plain(packing.get("package_type")) or None,
        "packing_count": places,
        "container": _plain(header.get("container") or packing.get("container")) or None,
        "pallets": _count(pallets) if pallets not in (None, "") else None,
        "container_numbers": _plain(header.get("container") or packing.get("container")) or None,
        "fill_mark": None,
        "country": _plain(customs.get("country") or header.get("country")) or None,
        "gross": _fnum(_gross(item)),
        "net": net,
        "net_bare": primary if primary not in (None, net) else None,
        "amount": _fnum(_amount(item)),
        "invoice_no": _plain(header.get("invoice_no")) or None,
        "invoice_date": _plain(header.get("invoice_date") or header.get("date")) or None,
        "invoice_currency": _plain(header.get("currency")) or None,
        "description_short": short,
        "its_request": None,
        "its_price_net": None,
        "declaration": None,
        "price_discounted": _dt_price_discounted(item),
    }


def dt_table(items: list[dict[str, Any]], header: dict[str, Any] | None) -> tuple[list[str], list[list[Any]]]:
    """Шапка и строки ДЛЯ ДТ: только именованные столбцы бланка, пустые по данным выкинуты."""
    from app.services.column_layout import prune_empty_columns

    headers = dt_headers()
    keys = dt_keys()
    rows: list[list[Any]] = []
    for index, item in enumerate(items, start=1):
        values = _dt_values(item, header, index)
        rows.append([values.get(key) for key in keys])
    if not rows:
        return headers, rows
    body = rows[:]
    total: list[Any] = []
    title_by_key = {key: title for key, title in _DT_COLUMN_SPECS}
    for pos, key in enumerate(keys):
        title = title_by_key.get(key, "")
        if title not in _DT_COUNT and title not in _DT_MONEY:
            total.append(None)
            continue
        nums = [float(row[pos]) for row in body if isinstance(row[pos], (int, float))]
        if not nums:
            total.append(None)
            continue
        total.append(_count(sum(nums)) if title in _DT_COUNT else _fnum(sum(nums)))
    rows.append(total)
    next_h, _keys, next_r = prune_empty_columns(headers, keys, rows, keep_keys=_DT_KEEP_ALWAYS)
    return next_h, next_r


def dt_rows(items: list[dict[str, Any]], header: dict[str, Any] | None) -> list[list[Any]]:
    _headers, rows = dt_table(items, header)
    return rows


TSD_INV_KEYS = [
    "no",
    "hs_code",
    "description",
    "country",
    "article",
    "manufacturer",
    "net_weight",
    "net_primary",
    "qty",
    "price",
    "unit",
    "amount",
]
TSD_PL_KEYS = [
    "no",
    "hs_code",
    "description",
    "manufacturer",
    "article",
    "packages",
    "qty",
    "net_weight",
    "net_primary",
    "gross_weight",
    "country",
]
TSD_SPEC_KEYS = [
    "description",
    "article",
    "manufacturer",
    "country",
    "qty",
    "unit",
    "gross_weight",
    "net_weight",
    "net_primary",
    "amount",
    "unit_2",
    "price",
    "size",
]


def _tsd_role_tables(
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    column_layout: dict[str, Any] | None,
) -> tuple[tuple[list[str], list[list[Any]]], tuple[list[str], list[list[Any]]], tuple[list[str], list[list[Any]]], dict[str, Any]]:
    header = export_header_fields(header or {}, items)
    ccy = _currency(items, header)
    inv_h, _, inv_r = apply_sheet_layout(
        invoice_headers(ccy), TSD_INV_KEYS, invoice_rows(items, header), items, column_layout, "invoice"
    )
    pl_h, _, pl_r = apply_sheet_layout(
        packing_headers(), TSD_PL_KEYS, packing_rows(items, header), items, column_layout, "packing"
    )
    spec_h, _, spec_r = apply_sheet_layout(
        spec_headers(ccy), TSD_SPEC_KEYS, spec_rows(items, header), items, column_layout, "specification"
    )
    return (inv_h, inv_r), (pl_h, pl_r), (spec_h, spec_r), header


def export_tsd_book(
    items: list[dict[str, Any]],
    output_path,
    header: dict[str, Any] | None = None,
    column_layout: dict[str, Any] | None = None,
):
    (inv_h, inv_r), (pl_h, pl_r), (spec_h, spec_r), header = _tsd_role_tables(items, header, column_layout)
    wb = Workbook()
    ws = wb.active
    ws.title = "INV"
    _write_invoice_letterhead(ws, header, len(inv_h))
    _append_table(ws, inv_h, inv_r)

    ws_pl = wb.create_sheet("PAK")
    _write_packing_letterhead(ws_pl, header, len(pl_h))
    _append_table(ws_pl, pl_h, pl_r)

    ws_spec = wb.create_sheet("Specification")
    _write_spec_letterhead(ws_spec, header, len(spec_h))
    _append_table(ws_spec, spec_h, spec_r)

    ws_desc = wb.create_sheet("ОПИСАНИЕ")
    desc_h, desc_r = _description_table(items, header)
    _append_table(ws_desc, desc_h, desc_r)

    ws_dt = wb.create_sheet("ДЛЯ ДТ")
    dt_h, dt_r = dt_table(items, header)
    _append_table(ws_dt, dt_h, dt_r)

    unfreeze_workbook(wb)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _description_table(
    items: list[dict[str, Any]], header: dict[str, Any] | None
) -> tuple[list[str], list[list[Any]]]:
    from app.services.column_layout import prune_empty_columns

    headers = description_headers()
    keys = ["no", "hs_code", "manufacturer", "article", "description"]
    rows = description_rows(items, header)
    next_h, _keys, next_r = prune_empty_columns(headers, keys, rows, keep_keys={"no", "article"})
    return next_h, next_r


def tsd_preview(
    items: list[dict[str, Any]],
    header: dict[str, Any] | None,
    filename: str,
    column_layout: dict[str, Any] | None = None,
) -> dict[str, Any]:
    (inv_h, inv_r), (pl_h, pl_r), (spec_h, spec_r), header = _tsd_role_tables(items, header, column_layout)
    desc_h, desc_r = _description_table(items, header)
    dt_h, dt_r = dt_table(items, header)
    return {
        "filename": filename,
        "sheets": [
            {"title": "INV", "headers": inv_h, "rows": inv_r},
            {"title": "PAK", "headers": pl_h, "rows": pl_r},
            {"title": "Specification", "headers": spec_h, "rows": spec_r},
            {"title": "ОПИСАНИЕ", "headers": desc_h, "rows": desc_r},
            {"title": "ДЛЯ ДТ", "headers": dt_h, "rows": dt_r},
        ],
    }
