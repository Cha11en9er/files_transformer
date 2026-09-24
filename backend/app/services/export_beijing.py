"""Beijing Goldluck export: one workbook, four sheets matching the ED etalon."""

from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.parsing.header_extract import currency_from_sources
from app.services.field_map import article_with_color
from app.services.export_style import (
    merge_column,
    style_data_table,
    style_letterhead_row,
    style_split_letterhead,
    unfreeze_workbook,
)
from app.services.field_map import is_factory_note, parse_number

NUM_FMT = "0.00"
NUM_FMT_4 = "0.0000"
NUM_FMT_6 = "0.000000"
NUM_FMT_INT = "0"


def _fnum(value: Any, digits: int = 2) -> float | None:
    if value in (None, ""):
        return None
    number = parse_number(value)
    if number is None:
        return None
    return round(float(number), digits)


def _count(value: Any) -> int | float | None:
    number = _fnum(value, 2)
    if number is None:
        return None
    if abs(number - round(number)) < 1e-9:
        return int(round(number))
    return number


def _qty(item: dict[str, Any]) -> Any:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    if commercial.get("qty") not in (None, ""):
        return commercial.get("qty")
    return packing.get("meters")


def _unit(item: dict[str, Any]) -> str:
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    return str(commercial.get("unit") or packing.get("unit") or "")


def _cartons(item: dict[str, Any]) -> Any:
    packing = item.get("packing_data") or {}
    traces = item.get("source_traces") or {}
    inv = traces.get("invoice") or {}
    return packing.get("boxes") or packing.get("rolls") or inv.get("boxes") or inv.get("cartons")


def _desc(item: dict[str, Any]) -> str:
    customs = item.get("customs_data") or {}
    en = customs.get("description_en")
    ru = customs.get("description_ru")
    if en and ru and not is_factory_note(en) and not is_factory_note(ru):
        return f"{en}/{ru}"
    for key in ("description", "description_en", "description_ru"):
        value = customs.get(key)
        if value and not is_factory_note(value):
            return str(value)
    return ""


def _qty_lots(item: dict[str, Any]) -> list[dict[str, Any]]:
    commercial = item.get("commercial_data") or {}
    return [
        lot
        for lot in (commercial.get("lots") or [])
        if isinstance(lot, dict) and lot.get("qty") not in (None, "")
    ]


def _visual_lines(item: dict[str, Any], *, mode: str) -> list[dict[str, Any]]:
    """Invoice follows commercial lots; packing follows packing lines; spec is one totals row."""
    commercial = item.get("commercial_data") or {}
    packing = item.get("packing_data") or {}
    lots = _qty_lots(item)
    plines = [line for line in (packing.get("lines") or []) if isinstance(line, dict)]
    if mode == "invoice":
        n = max(len(lots), 1)
    elif mode == "packing":
        n = max(len(plines), 1)
    else:
        n = 1
    lines: list[dict[str, Any]] = []
    for i in range(n):
        lot = lots[i] if mode != "spec" and i < len(lots) else {}
        pline = plines[i] if mode == "packing" and i < len(plines) else {}
        line = {**pline, **lot}
        if i == 0:
            line.setdefault("qty", _qty(item))
            line.setdefault("unit", _unit(item))
            line.setdefault("price", commercial.get("price"))
            line.setdefault("amount", commercial.get("amount"))
            line.setdefault("color", commercial.get("color"))
            line.setdefault("measurement", packing.get("measurement"))
            line.setdefault("gross_weight", packing.get("gross_weight"))
            line.setdefault("net_weight", packing.get("net_weight"))
            line.setdefault("boxes", packing.get("boxes") or packing.get("rolls"))
            line.setdefault("volume", packing.get("volume"))
            line.setdefault("pcs_per_carton", packing.get("pcs_per_carton"))
        lines.append(line)
    return lines


def invoice_headers(ccy: str = "CNY") -> list[str]:
    return [
        "No",
        "Customs Code",
        "Description of the goods",
        "Art No.",
        "Color",
        "Quantity",
        "Unit",
        f"Price ({ccy})",
        f"Amount ({ccy})",
    ]


def packing_headers() -> list[str]:
    return [
        "No",
        "Description of the goods",
        "Art No.",
        "Color",
        "Quantity",
        "Unit",
        "Measurement",
        "Gross Wt. (kg)",
        "Net Wt. (kg)",
        "Cartons",
        "Volume (m3)",
        "unit per carton",
    ]


def spec_headers() -> list[str]:
    return [
        "No",
        "Item/ Артикул",
        "Description/ Наименование",
        "Quantity, ctns/ Количество, коробок",
        "CT",
        "Quantity, unit/ Количество, единиц",
        "Unit/Единица измерения",
        "Netto weight, kg/ Вес нетто, кг",
        "GROSS weight, kg/ Вес брутто, кг",
        "Unit Price / Цена за единицу",
        "Amount / Cтоимость",
        "Customs code / Таможенный код",
    ]


def description_headers() -> list[str]:
    return ["Item/ Артикул", "Description/ Наименование", "Manufacturer", "Country"]


def _invoice_table(items: list[dict[str, Any]]) -> tuple[list[list[Any]], list[tuple[int, int]], list[tuple[str, int, int]]]:
    rows: list[list[Any]] = []
    item_spans: list[tuple[int, int]] = []
    pack_spans: list[tuple[str, int, int]] = []
    for idx, item in enumerate(items, start=1):
        customs = item.get("customs_data") or {}
        commercial = item.get("commercial_data") or {}
        lines = _visual_lines(item, mode="invoice")
        start = len(rows)
        for line_i, line in enumerate(lines):
            first = line_i == 0
            rows.append(
                [
                    idx if first else None,
                    (customs.get("tnved_code") or customs.get("hs_code")) if first else None,
                    _desc(item) if first else None,
                    (article_with_color(item.get("article") or item.get("normalized_article"), line.get("color") or commercial.get("color")) if first else None),
                    line.get("color"),
                    _count(line.get("qty")),
                    line.get("unit") or (_unit(item) if first else None),
                    _fnum(line.get("price"), 4),
                    _fnum(line.get("amount")),
                ]
            )
        item_spans.append((start, len(rows) - 1))
        group = str((item.get("packing_data") or {}).get("pack_group") or "")
        if group:
            pack_spans.append((group, start, len(rows) - 1))
    if rows:
        body = rows[:]
        rows.append(
            [
                None,
                None,
                "TOTAL:",
                None,
                None,
                _count(sum(float(r[5] or 0) for r in body)),
                None,
                None,
                _fnum(sum(float(r[8] or 0) for r in body)),
            ]
        )
    return rows, item_spans, pack_spans


def _packing_table(items: list[dict[str, Any]]) -> tuple[list[list[Any]], list[tuple[int, int]], list[tuple[str, int, int]]]:
    rows: list[list[Any]] = []
    item_spans: list[tuple[int, int]] = []
    pack_spans: list[tuple[str, int, int]] = []
    for idx, item in enumerate(items, start=1):
        packing = item.get("packing_data") or {}
        traces = item.get("source_traces") or {}
        inv = traces.get("invoice") or {}
        lines = _visual_lines(item, mode="packing")
        start = len(rows)
        for line_i, line in enumerate(lines):
            first = line_i == 0
            rows.append(
                [
                    idx if first else None,
                    _desc(item) if first else None,
                    (article_with_color(item.get("article"), line.get("color") or (item.get("commercial_data") or {}).get("color")) if first else None),
                    line.get("color"),
                    _count(line.get("qty")),
                    line.get("unit") or (_unit(item) if first else None),
                    line.get("measurement") or (packing.get("measurement") if first else None) or inv.get("measurement"),
                    _fnum(line.get("gross_weight") if line.get("gross_weight") not in (None, "") else (packing.get("gross_weight") if first else None)),
                    _fnum(line.get("net_weight") if line.get("net_weight") not in (None, "") else (packing.get("net_weight") if first else None)),
                    _count(line.get("boxes") if line.get("boxes") not in (None, "") else (_cartons(item) if first else None)),
                    _fnum(line.get("volume") if line.get("volume") not in (None, "") else (packing.get("volume") if first else None), 6),
                    _count(line.get("pcs_per_carton") or (packing.get("pcs_per_carton") if first else None) or inv.get("pcs_per_carton")),
                ]
            )
        item_spans.append((start, len(rows) - 1))
        group = str(packing.get("pack_group") or "")
        if group:
            pack_spans.append((group, start, len(rows) - 1))
    if rows:
        body = rows[:]
        rows.append(
            [
                None,
                "TOTAL:",
                None,
                None,
                _count(sum(float(r[4] or 0) for r in body)),
                None,
                None,
                _fnum(sum(float(r[7] or 0) for r in body)),
                _fnum(sum(float(r[8] or 0) for r in body)),
                _count(sum(float(r[9] or 0) for r in body)),
                _fnum(sum(float(r[10] or 0) for r in body), 6),
                None,
            ]
        )
    return rows, item_spans, pack_spans


def _spec_table(items: list[dict[str, Any]]) -> tuple[list[list[Any]], list[tuple[int, int]], list[tuple[str, int, int]]]:
    rows: list[list[Any]] = []
    item_spans: list[tuple[int, int]] = []
    pack_spans: list[tuple[str, int, int]] = []
    for idx, item in enumerate(items, start=1):
        commercial = item.get("commercial_data") or {}
        packing = item.get("packing_data") or {}
        customs = item.get("customs_data") or {}
        lines = _visual_lines(item, mode="spec")
        start = len(rows)
        for line_i, line in enumerate(lines):
            first = line_i == 0
            rows.append(
                [
                    idx if first else None,
                    (article_with_color(item.get("article"), line.get("color") or (item.get("commercial_data") or {}).get("color")) if first else None),
                    _desc(item) if first else None,
                    _count(line.get("boxes") if line.get("boxes") not in (None, "") else (_cartons(item) if first else None)),
                    "CT" if first or line.get("boxes") not in (None, "") else None,
                    _count(line.get("qty")),
                    line.get("unit") or (_unit(item) if first else None),
                    _fnum(line.get("net_weight") if line.get("net_weight") not in (None, "") else (packing.get("net_weight") if first else None)),
                    _fnum(line.get("gross_weight") if line.get("gross_weight") not in (None, "") else (packing.get("gross_weight") if first else None)),
                    _fnum(line.get("price") if line.get("price") not in (None, "") else (commercial.get("price") if first else None), 4),
                    _fnum(line.get("amount") if line.get("amount") not in (None, "") else (commercial.get("amount") if first else None)),
                    (customs.get("tnved_code") or customs.get("hs_code")) if first else None,
                ]
            )
        item_spans.append((start, len(rows) - 1))
        group = str(packing.get("pack_group") or "")
        if group:
            pack_spans.append((group, start, len(rows) - 1))
    if rows:
        body = rows[:]
        rows.append(
            [
                "Total/Итого:",
                None,
                None,
                _count(sum(float(r[3] or 0) for r in body)),
                None,
                _count(sum(float(r[5] or 0) for r in body)),
                None,
                _fnum(sum(float(r[7] or 0) for r in body)),
                _fnum(sum(float(r[8] or 0) for r in body)),
                None,
                _fnum(sum(float(r[10] or 0) for r in body)),
                None,
            ]
        )
    return rows, item_spans, pack_spans


def invoice_rows(items: list[dict[str, Any]]) -> list[list[Any]]:
    return _invoice_table(items)[0]


def packing_rows(items: list[dict[str, Any]]) -> list[list[Any]]:
    return _packing_table(items)[0]


def spec_rows(items: list[dict[str, Any]]) -> list[list[Any]]:
    return _spec_table(items)[0]


def description_rows(items: list[dict[str, Any]], header: dict[str, Any] | None) -> list[list[Any]]:
    # Operator header overrides row manufacturer when set.
    manufacturer = (header or {}).get("manufacturer") or "BEIJING GOLDLUCK CO., LTD"
    country = (header or {}).get("country") or "CN"
    rows: list[list[Any]] = []
    for item in items:
        customs = item.get("customs_data") or {}
        rows.append(
            [
                article_with_color(item.get("article"), (item.get("commercial_data") or {}).get("color")),
                _desc(item),
                (header or {}).get("manufacturer") or customs.get("manufacturer") or manufacturer,
                customs.get("country") or country,
            ]
        )
    return rows


def _write_letterhead(ws, kind: str, header: dict[str, Any] | None, cols: int) -> None:
    header = header or {}
    seller = header.get("seller") or "BEIJING GOLDLUCK CO., LTD"
    address = header.get("seller_address") or ""
    if not header.get("seller") and not address:
        address = "RM.317, NO.33 DENGSHIKOU STREET, DONGCHENG DISTRICT, BEIJING, CHINA"
    buyer = header.get("buyer") or ""
    buyer_address = header.get("buyer_address") or ""
    contract = header.get("contract_no") or ""
    invoice_no = header.get("invoice_no") or ""
    date = header.get("invoice_date") or header.get("date") or ""
    container = header.get("container_no") or header.get("container") or ""
    ws.append([seller])
    style_letterhead_row(ws, ws.max_row, cols, company=True)
    ws.append([address])
    style_letterhead_row(ws, ws.max_row, cols)
    ws.append([])
    if kind == "specification":
        ws.append([f"Specification / Спецификация №: {invoice_no} dated {date}".strip()])
        style_letterhead_row(ws, ws.max_row, cols, title=True)
        ws.append([f"To the contract / К контракту:{contract}".strip()])
        style_letterhead_row(ws, ws.max_row, cols)
    elif kind == "description":
        ws.append([f"DESCRIPTION / Описание №: {invoice_no} dated {date}".strip()])
        style_letterhead_row(ws, ws.max_row, cols, title=True)
        ws.append([f"To the contract / К контракту:{contract}".strip()])
        style_letterhead_row(ws, ws.max_row, cols)
    else:
        title = "Commercial invoice" if kind == "invoice" else "Packing list"
        ws.append([title])
        style_letterhead_row(ws, ws.max_row, cols, title=True)
        split = min(6, cols)
        pad = [None] * max(split - 2, 0)
        ws.append([f"Buyer: {buyer}", *pad, f"Contract No.: {contract}"])
        style_split_letterhead(ws, ws.max_row, cols, split_at=split)
        ws.append([f"Add: {buyer_address}", *pad, f"Invoice No.: {invoice_no}"])
        style_split_letterhead(ws, ws.max_row, cols, split_at=split)
        ws.append([None, *pad, f"Invoice date: {date}"])
        style_split_letterhead(ws, ws.max_row, cols, split_at=split)
        if container:
            ws.append([None, *pad, f"Container No.: {container}"])
            style_split_letterhead(ws, ws.max_row, cols, split_at=split)
    ws.append([])


def _append_table(
    ws,
    headers: list[str],
    rows: list[list[Any]],
    *,
    item_spans: list[tuple[int, int]] | None = None,
    pack_spans: list[tuple[str, int, int]] | None = None,
    identity_cols: tuple[int, ...] = (),
    pack_cols: tuple[int, ...] = (),
) -> None:
    ws.append(headers)
    header_row = ws.max_row
    start = ws.max_row + 1
    price_cols = {
        i + 1
        for i, title in enumerate(headers)
        if "price" in title.lower() or "цена" in title.lower()
    }
    vol_cols = {
        i + 1
        for i, title in enumerate(headers)
        if "volume" in title.lower() or "m3" in title.lower()
    }
    for row in rows:
        ws.append(row)
    for r in range(start, ws.max_row + 1):
        for cell in ws[r]:
            if type(cell.value) is int:
                cell.number_format = NUM_FMT_INT
            elif isinstance(cell.value, float):
                if cell.column in vol_cols:
                    cell.number_format = NUM_FMT_6
                elif cell.column in price_cols:
                    cell.number_format = NUM_FMT_4
                else:
                    cell.number_format = NUM_FMT
            if isinstance(cell.value, str) and cell.value.upper().startswith("TOTAL"):
                cell.font = Font(bold=True)
    for first, last in item_spans or []:
        if last <= first:
            continue
        for col in identity_cols:
            merge_column(ws, col, start + first, start + last)
    grouped: dict[str, list[tuple[int, int]]] = {}
    for group, first, last in pack_spans or []:
        grouped.setdefault(group, []).append((first, last))
    for spans in grouped.values():
        block_first = min(span[0] for span in spans)
        block_last = max(span[1] for span in spans)
        if block_last <= block_first:
            continue
        for col in pack_cols:
            merge_column(ws, col, start + block_first, start + block_last)
    style_data_table(
        ws,
        header_row=header_row,
        cols=len(headers),
        n_rows=1 + len(rows),
        headers=headers,
    )


def export_beijing_book(items: list[dict[str, Any]], output_path, header: dict[str, Any] | None = None):
    ccy = currency_from_sources(items, header) or "CNY"
    inv_h = invoice_headers(ccy)
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    inv_rows, inv_spans, inv_pack = _invoice_table(items)
    _write_letterhead(ws, "invoice", header, len(inv_h))
    _append_table(ws, inv_h, inv_rows, item_spans=inv_spans, identity_cols=(1, 2, 3, 4))

    ws_pl = wb.create_sheet("Packing list")
    pl_rows, pl_spans, pl_pack = _packing_table(items)
    _write_letterhead(ws_pl, "packing", header, len(packing_headers()))
    _append_table(
        ws_pl,
        packing_headers(),
        pl_rows,
        item_spans=pl_spans,
        pack_spans=pl_pack,
        identity_cols=(1, 2, 3),
        pack_cols=(7, 8, 9, 10, 11),
    )

    ws_spec = wb.create_sheet("Specification")
    spec_body, spec_spans, spec_pack = _spec_table(items)
    _write_letterhead(ws_spec, "specification", header, len(spec_headers()))
    _append_table(
        ws_spec,
        spec_headers(),
        spec_body,
        item_spans=spec_spans,
        pack_spans=spec_pack,
        identity_cols=(1, 2, 3, 12),
        pack_cols=(4, 8, 9),
    )

    ws_desc = wb.create_sheet("описание")
    _write_letterhead(ws_desc, "description", header, len(description_headers()))
    _append_table(ws_desc, description_headers(), description_rows(items, header))

    unfreeze_workbook(wb)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
