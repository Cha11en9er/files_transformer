from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import pandas as pd

from app.parsing.languages import decode_bytes
from app.parsing.normalize import normalize_text
from app.parsing.table_rows import lines_from_matrix

ARTICLE_HEADERS = (
    "articule",
    "articul",
    "article",
    "art.",
    "art",
    "артикул",
    "арт",
    "model",
    "модель",
    "design",
    "дизайн",
    "product name",
    "product",
    "наименование",
    "item no",
    "item no.",
    "item number",
    "art no",
    "art no.",
    "style",
    "sku",
    "код",
    "货号",
    "品号",
    "pattern",
    "ürün kodu",
    "urun kodu",
    "müşteri kodu",
    "musteri kodu",
)

HEADER_SCAN_ROWS = 30


def _cell_str(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return normalize_text(str(value))


def _header_key(value: Any) -> str:
    return _cell_str(value).lower()


def detect_header_row(rows: list[list[Any]]) -> int | None:
    best_idx: int | None = None
    best_hits = 0
    limit = min(len(rows), HEADER_SCAN_ROWS)
    for idx in range(limit):
        keys = [_header_key(cell) for cell in rows[idx]]
        hits = 0
        joined = " | ".join(keys)
        for header in ARTICLE_HEADERS:
            if any(header == key or header in key for key in keys) or header in joined:
                hits += 1
        extra = (
            "qty",
            "quantity",
            "price",
            "amount",
            "weight",
            "hs",
            "tn",
            "packages",
            "meters",
            "hides",
            "количество",
            "数量",
            "fiyat",
        )
        hits += sum(1 for token in extra if any(token in key for key in keys))
        if hits > best_hits:
            best_hits = hits
            best_idx = idx
    if best_hits < 1:
        return None
    return best_idx


def _unique_columns(headers: Iterable[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for raw in headers:
        base = _header_key(raw) or "column"
        n = seen.get(base, 0)
        seen[base] = n + 1
        result.append(base if n == 0 else f"{base}_{n+1}")
    return result


def _article_rank(key: str) -> int | None:
    lowered = key.lower()
    if "price" in lowered or "cart" in lowered:
        return None
    if "артикул" in lowered or "articul" in lowered:
        return 0
    if re.search(r"(^|[^a-zа-я])article([^a-zа-я]|$)", lowered) and "name" not in lowered:
        return 1
    if "art no" in lowered or lowered.strip() in {"art no", "art no.", "art. no"}:
        return 1
    if any(token in lowered for token in ("sku", "item no", "style")) or lowered.strip() in {
        "art",
        "art.",
        "арт",
    }:
        return 2
    if "design" in lowered or "дизайн" in lowered:
        return 3
    if "product name" in lowered or "наименование" in lowered or lowered.strip() == "product":
        return 5
    return None


def _pick_article(row: dict[str, Any]) -> tuple[str | None, str | None]:
    article = None
    article_rank = 99
    model = None
    for key, value in row.items():
        text = _cell_str(value)
        if not text:
            continue
        lowered = key.lower()
        rank = _article_rank(key)
        if rank is not None and rank < article_rank:
            article = text
            article_rank = rank
        if "model" in lowered or "модель" in lowered:
            if model is None:
                model = text
    if article is None and model is not None:
        article = model
    return article, model


@dataclass(frozen=True)
class ExcelSheetRead:
    name: str
    lines: list[dict[str, Any]]
    preview: str


def _sheet_preview(sheet_name: str, frame: pd.DataFrame) -> str:
    sample = frame.head(80).fillna("").astype(str)
    body = sample.to_string(index=False, max_cols=24)
    return f"{sheet_name}\n{body}"


def dataframe_to_lines(
    frame: pd.DataFrame,
    *,
    sheet_name: str,
    header_row: int | None = None,
    inherited_cells: set[tuple[int, int]] | None = None,
) -> list[dict[str, Any]]:
    if frame.empty:
        return []
    raw_rows = frame.where(pd.notna(frame), None).values.tolist()
    return lines_from_matrix(
        raw_rows,
        sheet_name=sheet_name,
        header_row=header_row,
        inherited_cells=inherited_cells,
    )


def _worksheet_matrix(ws) -> tuple[list[list[Any]], set[tuple[int, int]]]:
    """Copy merged origin values into every cell of the range.

    inherited marks cells that were empty and filled from a merge (not the origin).
    """
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    grid = [[ws.cell(row, col).value for col in range(1, max_col + 1)] for row in range(1, max_row + 1)]
    inherited: set[tuple[int, int]] = set()
    for merged in ws.merged_cells.ranges:
        origin = ws.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                r0, c0 = row - 1, col - 1
                if (row, col) == (merged.min_row, merged.min_col):
                    continue
                if grid[r0][c0] in (None, ""):
                    grid[r0][c0] = origin
                    inherited.add((r0, c0))
    return grid, inherited


def _sniff_excel_engines(path: str) -> list[str]:
    suffix = Path(path).suffix.lower()
    try:
        head = Path(path).read_bytes()[:256]
    except Exception:
        head = b""
    if head.startswith(b"PK"):
        return ["openpyxl"]
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return ["xlrd"]
    low = head.lstrip().lower()
    if low.startswith((b"<html", b"<!doctype", b"<?xml")) or b"<table" in low:
        return ["html"]
    if suffix in {".xlsx", ".xlsm"}:
        return ["openpyxl", "xlrd"]
    if suffix == ".xls":
        return ["xlrd", "openpyxl"]
    return ["openpyxl", "xlrd"]


def load_excel_book(path: str | Path) -> dict[str, pd.DataFrame]:
    """Read every sheet. Always pass an engine so pandas does not guess blindly."""
    path_s = str(path)
    last_exc: Exception | None = None
    for engine in _sniff_excel_engines(path_s):
        try:
            if engine == "html":
                raw = Path(path_s).read_bytes()
                text, _encoding = decode_bytes(raw)
                frames = pd.read_html(io.StringIO(text), header=None)
                return {f"Sheet{index + 1}": frame for index, frame in enumerate(frames)}
            book = pd.read_excel(
                path_s,
                sheet_name=None,
                header=None,
                dtype=object,
                engine=engine,
            )
            if isinstance(book, dict):
                return book
            return {"Sheet1": book}
        except Exception as exc:
            last_exc = exc
            continue
    raise last_exc or ValueError("excel_unreadable")


def read_excel_sheets(path: str) -> list[ExcelSheetRead]:
    suffix = str(path).lower()
    if suffix.endswith(".xlsx") or suffix.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=True, read_only=False)
            sheets: list[ExcelSheetRead] = []
            try:
                for worksheet in workbook.worksheets:
                    name = str(worksheet.title)
                    matrix, inherited = _worksheet_matrix(worksheet)
                    if not matrix:
                        continue
                    frame = pd.DataFrame(matrix)
                    sheets.append(
                        ExcelSheetRead(
                            name=name,
                            lines=lines_from_matrix(
                                matrix,
                                sheet_name=name,
                                inherited_cells=inherited,
                            ),
                            preview=_sheet_preview(name, frame),
                        )
                    )
            finally:
                workbook.close()
            if sheets:
                return sheets
        except Exception:
            pass
    book = load_excel_book(path)
    sheets = []
    for sheet_name, frame in book.items():
        name = str(sheet_name)
        sheets.append(
            ExcelSheetRead(
                name=name,
                lines=dataframe_to_lines(frame, sheet_name=name),
                preview=_sheet_preview(name, frame),
            )
        )
    return sheets


def read_excel(path: str) -> tuple[list[str], list[dict[str, Any]], str]:
    """Read all sheets. Returns sheet names, parsed line dicts, concatenated preview text."""
    sheets = read_excel_sheets(path)
    lines: list[dict[str, Any]] = []
    preview_chunks: list[str] = []
    for sheet in sheets:
        preview_chunks.append(sheet.preview)
        lines.extend(sheet.lines)
    return [sheet.name for sheet in sheets], lines, "\n".join(preview_chunks)
