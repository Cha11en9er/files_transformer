"""Agentic UI audit against the live Excel Transformer site.

Opens the site in Chromium, uploads one shipment kit at a time, waits for
parse (+ model review), downloads the export ZIP, dumps input / etalon /
export for human+LLM comparison, and appends verdicts to findings.md.

Does NOT patch the transformer. Discrepancies are recorded only.

Run from anywhere:
  python backend/scripts/agent_site_audit/run_site_audit.py
  python backend/scripts/agent_site_audit/run_site_audit.py --limit 2
  python backend/scripts/agent_site_audit/run_site_audit.py --base-url http://127.0.0.1:8012
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[3]
KITS_ROOT = ROOT / "documents" / "я_тестирую"
OUT_ROOT = Path(__file__).resolve().parents[1] / "debug_docs" / "out" / "agent_site_audit"
VENV_PY_HINT = ROOT / "backend" / ".venv" / "Scripts" / "python.exe"

ACCEPT_RE = re.compile(r"\.(xlsx|xls|xlsm|pdf|png|jpe?g|tif{1,2}|bmp|webp)$", re.I)
CATALOG_RE = re.compile(r"(справочник|сводная|catalog|catalogue|описание)", re.I)
SKIP_NAME_RE = re.compile(r"^(thumbs\.db|\.ds_store)$", re.I)

# Ten kits: both profiles, prefer PDF / scans, cover control + foreign shapes.
KIT_PLAN: list[dict[str, Any]] = [
    {
        "id": "01_ханчжоу_18233_626-1",
        "profile": "18233",
        "why": "Контроль ТЗ: Noble 110, дети vs семейство, веса с пакинга",
        "checks": ["Noble 110 отдельно", "15 rolls / 611.5 m", "нетто≈580.93"],
    },
    {
        "id": "11_beijing_goldluck_003",
        "profile": "BEIJING",
        "why": "Контроль Beijing + подписанный скан без текста + справочник",
        "checks": ["MD 812 200100 pcs", "пустой Art No. = лот", "после TOTAL не новые товары"],
    },
    {
        "id": "10_ханчжоу_18312_элемент",
        "profile": "18233",
        "why": "Hangzhou не ткань: PACKAGES ≠ QUANTITY + большой справочник",
        "checks": ["packages=места", "quantity=метры/шт", "профиль О-30"],
    },
    {
        "id": "12_beijing_goldluck_002",
        "profile": "BEIJING",
        "why": "Второй Beijing без скана, другая отгрузка",
        "checks": ["пустой артикул после I2388", "4 листа на выходе"],
    },
    {
        "id": "13_bestway_way04",
        "profile": "18233",
        "why": "Только PDF с текстом: PACKAGE ≠ QTY (критично)",
        "checks": ["модель 31021", "PACKAGE=10 places", "QTY=360 pcs"],
    },
    {
        "id": "14_bestway_way05",
        "profile": "BEIJING",
        "why": "PDF INV+PL+SPEC, чужой профиль выгрузки",
        "checks": ["PACKAGE≠QTY", "брутто на месте", "не пустой items"],
    },
    {
        "id": "15_матрац_intex_sc50226",
        "profile": "18233",
        "why": "PDF + скан ТСД без текста",
        "checks": ["PACKAGE≠QTY", "многостраничный INV", "скан не как таблица товаров"],
    },
    {
        "id": "16_шары_nh26001002",
        "profile": "BEIJING",
        "why": "PDF раньше «не отработал»; 2 позиции шаров",
        "checks": ["2 позиции", "PACKAGE 21+39=60", "сумма≈5831.54"],
    },
    {
        "id": "18_мора_18049",
        "profile": "18233",
        "why": "Скан инвойса без текста + пакинг xls (OCR)",
        "checks": ["OCR не пустой", "пакинг связан", "не_вход не грузим"],
    },
    {
        "id": "19_чжэцзян_18252",
        "profile": "BEIJING",
        "why": "PDF текст + Excel пакинг + подписанный скан",
        "checks": ["KASHEMIR метры", "rtf можно пропустить", "скан сверка"],
    },
]

# Round 2: retry 19 + 10 new + 3 self-judge (без опоры на эталон).
KIT_PLAN_ROUND2: list[dict[str, Any]] = [
    {
        "id": "19_чжэцзян_18252",
        "profile": "BEIJING",
        "why": "Retry: ранее network error / модель не отвечает",
        "checks": ["не network error", "KASHEMIR метры", "позиции > 0"],
    },
    {
        "id": "02_ханчжоу_18233_626-2",
        "profile": "18233",
        "why": "Hangzhou вторая половина 18233, NAPPA",
        "checks": ["отдельно от 626-1", "NAPPA дети", "3 книги"],
    },
    {
        "id": "03_ханчжоу_18080_621-1",
        "profile": "18233",
        "why": "Hangzhou 621-1 семейства в пакинге",
        "checks": ["дети с инвойса", "веса", "не пустой export"],
    },
    {
        "id": "05_ханчжоу_18206_624-1",
        "profile": "18233",
        "why": "Hangzhou 624-1 Mengma в имени",
        "checks": ["DESIGN/ROLLS/METERS", "итоги"],
    },
    {
        "id": "07_ханчжоу_18297_627-1",
        "profile": "18233",
        "why": "Hangzhou 627-1 ENZO/Preston…",
        "checks": ["дети отдельно", "3 книги"],
    },
    {
        "id": "08_ханчжоу_18297_627-2",
        "profile": "BEIJING",
        "why": "Hangzhou вход, профиль Beijing (чужой формат выгрузки)",
        "checks": ["позиции собрались", "одна книга"],
    },
    {
        "id": "09_ханчжоу_18297_627-3",
        "profile": "18233",
        "why": "Hangzhou 627-3, в эталоне бывает Sample",
        "checks": ["позиции", "веса"],
    },
    {
        "id": "20_виверс_18259",
        "profile": "18233",
        "why": "PDF Турция Weavers, метры/рулоны",
        "checks": ["DYER/дизайн", "net/gross", "не пустой"],
    },
    {
        "id": "21_ипекис_18259",
        "profile": "BEIJING",
        "why": "PDF инвойс + скан пакинга + лишний скан бухгалтерии",
        "checks": ["инвойс прочитан", "скан бухгалтерии не товары", "позиции"],
    },
    {
        "id": "22_кдф_18259",
        "profile": "18233",
        "why": "PDF Kadifeteks INV + packing/weight",
        "checks": ["метры", "rolls", "net/gross≈3047/3122"],
    },
    {
        "id": "23_тосун_18259",
        "profile": "BEIJING",
        "why": "PDF Tosunoglu + турецкий packing по рулонам",
        "checks": ["ZIMMY+SINDRI", "сумма≈10311", "рулоны"],
    },
    {
        "id": "17_logdv_sh_ldv01_6",
        "profile": "BEIJING",
        "why": "Self-judge: уже готовый ТСД ~30МБ, эталона нет",
        "checks": ["не упал", "позиции > 0 или явный skip", "не смешать CMR как товар"],
        "self_judge": True,
        "no_etalon": True,
    },
    {
        "id": "24_эскада_18259",
        "profile": "18233",
        "why": "Self-judge: .doc инвойс (не обещан) + пакинг xlsx",
        "checks": [".doc отклонён или понятная ошибка", "пакинг если один"],
        "self_judge": True,
    },
    {
        "id": "04_ханчжоу_18080_621-2",
        "profile": "18233",
        "why": "Self-judge: смотрим только вход vs вывод, эталон не сверяем",
        "checks": ["124 rolls в пакинге как контроль из README", "позиции > 0"],
        "self_judge": True,
    },
]


@dataclass
class KitResult:
    kit_id: str
    profile: str
    status: str
    why: str
    input_files: list[str] = field(default_factory=list)
    catalog_files: list[str] = field(default_factory=list)
    item_count: int | None = None
    warning_count: int | None = None
    upload_status: str = ""
    model_review: dict[str, Any] = field(default_factory=dict)
    workspace_summary: dict[str, Any] = field(default_factory=dict)
    export_zip: str | None = None
    export_files: list[str] = field(default_factory=list)
    totals_got: dict[str, float] = field(default_factory=dict)
    totals_etalon_hint: dict[str, Any] = field(default_factory=dict)
    findings: list[str] = field(default_factory=list)
    verdict: str = "unknown"
    elapsed_sec: float = 0.0
    error: str | None = None
    run_dir: str = ""
    attempt: int = 1
    attempts_total: int = 1
    error_kind: str = ""
    self_judge: bool = False


def _venv_python() -> Path | None:
    if VENV_PY_HINT.is_file():
        return VENV_PY_HINT
    return None


def ensure_dump_deps() -> None:
    """Prefer backend venv for openpyxl/pdf tools when system Python lacks them."""
    try:
        import openpyxl  # noqa: F401
        import xlrd  # noqa: F401
    except ImportError:
        vpy = _venv_python()
        if vpy:
            print(
                f"[warn] dump libs missing in current interpreter; "
                f"prefer: {vpy} {Path(__file__).name}",
                file=sys.stderr,
            )


def list_input_files(kit_dir: Path) -> tuple[list[Path], list[Path]]:
    ingress = kit_dir / "вход"
    if not ingress.is_dir():
        raise FileNotFoundError(f"нет папки вход/: {kit_dir}")
    goods: list[Path] = []
    catalog: list[Path] = []
    for path in sorted(ingress.iterdir()):
        if not path.is_file():
            continue
        if SKIP_NAME_RE.match(path.name):
            continue
        if not ACCEPT_RE.search(path.name):
            continue
        if CATALOG_RE.search(path.name):
            catalog.append(path)
        else:
            goods.append(path)
    if not goods and catalog:
        # rare: only catalog — treat as goods so site still gets something
        goods, catalog = catalog, []
    return goods, catalog


def safe_name(text: str) -> str:
    return re.sub(r"[^\w\-]+", "_", text, flags=re.U).strip("_")[:80]


def dump_excel_text(path: Path, max_rows: int = 60, max_cols: int = 16) -> str:
    lines = [f"FILE: {path.name}", f"SIZE: {path.stat().st_size}"]
    suffix = path.suffix.lower()
    try:
        if suffix == ".xls":
            import pandas as pd

            book = pd.read_excel(path, sheet_name=None, header=None, dtype=object, engine="xlrd")
            lines.append(f"SHEETS: {list(book)}")
            for name, frame in book.items():
                lines.append(f"-- sheet '{name}' dims={frame.shape}")
                rows = min(len(frame), max_rows)
                cols = min(frame.shape[1], max_cols)
                for r in range(rows):
                    cells = []
                    for c in range(cols):
                        v = frame.iat[r, c]
                        if v is None or (isinstance(v, float) and v != v):
                            cells.append("")
                        else:
                            cells.append(str(v).replace("\n", " ")[:50])
                    if any(cells):
                        lines.append(f"  r{r}| " + " | ".join(cells))
        else:
            from openpyxl import load_workbook

            wb = load_workbook(path, data_only=True, read_only=True)
            lines.append(f"SHEETS: {wb.sheetnames}")
            for ws in wb.worksheets:
                lines.append(f"-- sheet '{ws.title}'")
                for i, row in enumerate(ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True), 1):
                    cells = [("" if v is None else str(v).replace("\n", " ")[:50]) for v in row]
                    if any(cells):
                        lines.append(f"  r{i}| " + " | ".join(cells))
            wb.close()
    except Exception as exc:
        lines.append(f"[dump failed: {exc}]")
    return "\n".join(lines)


def dump_pdf_text(path: Path, max_pages: int = 4) -> str:
    lines = [f"FILE: {path.name}", f"SIZE: {path.stat().st_size}"]
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            lines.append(f"PAGES: {len(pdf.pages)}")
            for i, page in enumerate(pdf.pages[:max_pages], 1):
                text = (page.extract_text() or "").strip()
                lines.append(f"-- page {i} chars={len(text)}")
                if text:
                    lines.append(text[:4000])
                else:
                    lines.append("[no text layer]")
    except Exception as exc:
        lines.append(f"[pdfplumber failed: {exc}]")
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            lines.append(f"PAGES(pypdf): {len(reader.pages)}")
            for i, page in enumerate(reader.pages[:max_pages], 1):
                text = (page.extract_text() or "").strip()
                lines.append(f"-- page {i} chars={len(text)}")
                lines.append(text[:4000] if text else "[no text layer]")
        except Exception as exc2:
            lines.append(f"[pypdf failed: {exc2}]")
    return "\n".join(lines)


def render_pdf_previews(path: Path, out_dir: Path, max_pages: int = 2) -> list[str]:
    out_dir.mkdir(parents=True, exist_ok=True)
    saved: list[str] = []
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(path))
        for i in range(min(len(pdf), max_pages)):
            page = pdf[i]
            bitmap = page.render(scale=1.5)
            pil = bitmap.to_pil()
            target = out_dir / f"{safe_name(path.stem)}_p{i + 1}.png"
            pil.save(target)
            saved.append(str(target))
    except Exception as exc:
        (out_dir / f"{safe_name(path.stem)}_render_error.txt").write_text(str(exc), encoding="utf-8")
    return saved


def summarize_numbers_from_text(blob: str) -> dict[str, Any]:
    """Loose hints for the findings file — not a second parser."""
    hints: dict[str, Any] = {}
    for key, pattern in (
        ("total_rolls", r"(?:TOTAL|ИТОГО)[^\n]{0,40}?(\d+[.,]?\d*)\s*(?:rolls?|рулон)"),
        ("total_meters", r"(?:TOTAL|ИТОГО)[^\n]{0,60}?(\d+[.,]\d+)\s*m"),
        ("amount", r"(?:TOTAL|AMOUNT|СУММА)[^\n]{0,40}?(\d+[.,]\d+)"),
    ):
        m = re.search(pattern, blob, re.I)
        if m:
            hints[key] = m.group(1)
    return hints


def dump_kit_sources(
    kit_dir: Path,
    run_dir: Path,
    goods: list[Path],
    catalog: list[Path],
    *,
    skip_etalon: bool = False,
) -> None:
    src_dir = run_dir / "00_input_dump"
    et_dir = run_dir / "01_etalon_dump"
    img_dir = run_dir / "02_previews"
    src_dir.mkdir(parents=True, exist_ok=True)
    et_dir.mkdir(parents=True, exist_ok=True)
    img_dir.mkdir(parents=True, exist_ok=True)

    for path in goods + catalog:
        if path.suffix.lower() in {".xlsx", ".xls", ".xlsm"}:
            (src_dir / f"{safe_name(path.name)}.txt").write_text(dump_excel_text(path), encoding="utf-8")
        elif path.suffix.lower() == ".pdf":
            (src_dir / f"{safe_name(path.name)}.txt").write_text(dump_pdf_text(path), encoding="utf-8")
            render_pdf_previews(path, img_dir / "input", max_pages=2)

    etalon = kit_dir / "эталон_заказчика"
    if etalon.is_dir() and not skip_etalon:
        for path in sorted(etalon.iterdir()):
            if not path.is_file():
                continue
            if path.suffix.lower() in {".xlsx", ".xls", ".xlsm"}:
                (et_dir / f"{safe_name(path.name)}.txt").write_text(dump_excel_text(path), encoding="utf-8")
            elif path.suffix.lower() == ".pdf":
                (et_dir / f"{safe_name(path.name)}.txt").write_text(dump_pdf_text(path), encoding="utf-8")
                render_pdf_previews(path, img_dir / "etalon", max_pages=2)
    elif skip_etalon or not etalon.is_dir():
        (et_dir / "SKIPPED.txt").write_text(
            "эталон не сверяем (self_judge / no_etalon)\n", encoding="utf-8"
        )

    notes = kit_dir / "замечания"
    if notes.is_dir():
        for path in sorted(notes.iterdir()):
            if path.suffix.lower() == ".docx":
                try:
                    import zipfile as zf
                    import xml.etree.ElementTree as ET

                    with zf.ZipFile(path) as archive:
                        xml = archive.read("word/document.xml")
                    root = ET.fromstring(xml)
                    paras = []
                    for para in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                        texts = [
                            n.text or ""
                            for n in para.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
                        ]
                        blob = "".join(texts).strip()
                        if blob:
                            paras.append(blob)
                    (run_dir / f"notes_{safe_name(path.stem)}.txt").write_text(
                        "\n".join(paras) or "[empty]", encoding="utf-8"
                    )
                except Exception as exc:
                    (run_dir / f"notes_{safe_name(path.stem)}.txt").write_text(str(exc), encoding="utf-8")


def classify_error(message: str | None) -> str:
    text = (message or "").lower()
    if not text:
        return ""
    if text.startswith("готово") or ("позиций" in text and "ошибка" not in text):
        return ""
    if "network error" in text or "net::" in text or "failed to fetch" in text:
        return "network"
    if "timeout" in text:
        return "timeout"
    if "список файлов пуст" in text or "set_input_files" in text:
        return "upload_empty"
    if "не поймали done" in text or "no_api" in text:
        return "api_capture"
    if "ошибка:" in text:
        return "server_ui"
    return "other"


def needs_retry(result: KitResult) -> bool:
    if result.verdict == "fail":
        return True
    if result.error:
        return True
    status = (result.upload_status or "").lower()
    if status.startswith("ошибка") or "network error" in status or status == "timeout":
        return True
    if result.status in {"failed", "error", "no_api_payload", "no_workspace"}:
        return True
    mr = result.model_review or {}
    if str(mr.get("status") or "").lower() == "error":
        return True
    err = str(mr.get("error") or "").lower()
    if "не ответила" in err or "timeout" in err or "timed out" in err:
        return True
    return False


def dump_export_zip(zip_path: Path, run_dir: Path) -> list[str]:
    out = run_dir / "03_export"
    out.mkdir(parents=True, exist_ok=True)
    names: list[str] = []
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            target = out / Path(info.filename).name
            target.write_bytes(zf.read(info))
            names.append(target.name)
            if target.suffix.lower() in {".xlsx", ".xls", ".xlsm"}:
                (out / f"{safe_name(target.name)}.txt").write_text(
                    dump_excel_text(target), encoding="utf-8"
                )
    return names


def sum_workspace_totals(items: list[dict[str, Any]]) -> dict[str, float]:
    totals = {
        "qty": 0.0,
        "meters": 0.0,
        "amount": 0.0,
        "rolls": 0.0,
        "boxes": 0.0,
        "net_weight": 0.0,
        "gross_weight": 0.0,
    }
    for item in items:
        comm = item.get("commercial_data") or {}
        pack = item.get("packing_data") or {}
        for key in totals:
            for blob in (comm, pack):
                val = blob.get(key)
                if isinstance(val, (int, float)):
                    totals[key] += float(val)
                    break
    return totals


def workspace_rows_brief(items: list[dict[str, Any]], limit: int = 40) -> list[dict[str, Any]]:
    rows = []
    for item in items[:limit]:
        comm = item.get("commercial_data") or {}
        pack = item.get("packing_data") or {}
        cust = item.get("customs_data") or {}
        rows.append(
            {
                "article": item.get("article"),
                "qty": comm.get("qty") or pack.get("meters"),
                "unit": comm.get("unit"),
                "price": comm.get("price"),
                "amount": comm.get("amount"),
                "rolls": pack.get("rolls") or pack.get("boxes"),
                "net": pack.get("net_weight"),
                "gross": pack.get("gross_weight"),
                "hs": cust.get("hs_code"),
                "tnved": cust.get("tnved_code"),
                "desc": str(cust.get("description") or cust.get("description_en") or "")[:80],
                "flags": [
                    e.get("error_type")
                    for e in (item.get("validation_errors") or [])
                    if isinstance(e, dict)
                ][:5],
            }
        )
    return rows


def heuristic_verdict(kit: dict[str, Any], result: KitResult) -> tuple[str, list[str]]:
    findings = list(result.findings)
    if result.error:
        return "fail", findings + [f"ошибка прогона: {result.error}"]
    if result.item_count is None:
        return "fail", findings + ["нет item_count после обработки"]
    if result.item_count == 0:
        findings.append("0 позиций после обработки — для товарного комплекта это провал")
        return "fail", findings
    if not result.export_zip:
        findings.append("экспорт ZIP не скачан")
        return "fail", findings
    if not result.export_files:
        findings.append("ZIP пустой или без xlsx")
        return "fail", findings

    profile = result.profile
    if profile == "18233" and len(result.export_files) < 3:
        findings.append(f"профиль 18233: ожидали 3 книги, получили {result.export_files}")
    if profile == "BEIJING" and len(result.export_files) < 1:
        findings.append("профиль BEIJING: нет книги в ZIP")

    mr = result.model_review or {}
    if mr.get("status") == "ok" and mr.get("totals_mismatch"):
        findings.append("модель: totals_mismatch=true (скан/черновик не сошлись)")
    if mr.get("status") not in {None, "", "skipped", "ok"}:
        findings.append(f"модель status={mr.get('status')} error={mr.get('error')}")

    # Kit-specific soft checks — record, do not auto-fix except empty.
    article_blob = " ".join(
        str(r.get("article") or "") for r in (result.workspace_summary.get("rows") or [])
    ).upper()
    if kit["id"].startswith("01_") and "NOBLE" not in article_blob and "110" not in article_blob:
        findings.append("01: в позициях не видно Noble/110 — сверь с эталоном")
    if kit["id"].startswith("11_") and "MD" not in article_blob and "812" not in article_blob:
        findings.append("11: в позициях не видно MD 812 — сверь с эталоном")
    if "bestway" in kit["id"] or "матрац" in kit["id"] or "шары" in kit["id"]:
        totals = result.totals_got
        if totals.get("qty") and totals.get("rolls") and abs(totals["qty"] - totals["rolls"]) < 1e-6:
            findings.append(
                "PACKAGE/QTY: qty≈rolls — подозрение что места ушли в quantity "
                f"(qty={totals['qty']}, rolls={totals['rolls']})"
            )

    if findings:
        return "review", findings
    return "ok_candidate", findings + ["автоматически явных красных флагов нет — нужна ручная сверка с эталоном"]


def append_findings_md(path: Path, result: KitResult, kit: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    block = [
        f"## {result.kit_id}",
        f"- profile: `{result.profile}`",
        f"- why: {kit.get('why')}",
        f"- self_judge: {bool(kit.get('self_judge') or result.self_judge)}",
        f"- status/verdict: **{result.verdict}** ({result.status})",
        f"- attempt: {result.attempt}/{result.attempts_total}",
        f"- error_kind: {result.error_kind or '-'}",
        f"- elapsed: {result.elapsed_sec:.1f}s",
        f"- items: {result.item_count}, warnings: {result.warning_count}",
        f"- upload_status: {result.upload_status}",
        f"- model: {json.dumps(result.model_review, ensure_ascii=False)[:500]}",
        f"- totals_got: {result.totals_got}",
        f"- export: {result.export_files}",
        f"- run_dir: `{result.run_dir}`",
        "- checks:",
    ]
    for c in kit.get("checks") or []:
        block.append(f"  - {c}")
    block.append("- findings:")
    for f in result.findings or ["(empty)"]:
        block.append(f"  - {f}")
    if result.error:
        block.append(f"- error: `{result.error}`")
    block.append("")
    with path.open("a", encoding="utf-8") as fh:
        fh.write("\n".join(block) + "\n")


def wait_processing(page, timeout_sec: float = 600.0, poll_sec: float = 10.0) -> str:
    """Poll UI every poll_sec until workspace appears or hard failure / timeout."""
    deadline = time.time() + timeout_sec
    last = ""
    while time.time() < deadline:
        status = ""
        try:
            status = page.locator("#upload-status").inner_text(timeout=2000).strip()
        except Exception:
            status = ""
        last = status or last
        btn = ""
        try:
            btn = page.locator("#btn-process").inner_text(timeout=1000).strip()
        except Exception:
            pass
        ws_visible = False
        try:
            ws_visible = page.locator("#workspace").is_visible()
        except Exception:
            ws_visible = False
        if status.startswith("Ошибка"):
            return status
        if ws_visible and (btn == "Обработано" or status.startswith("Готово")):
            return status or "Готово"
        if "Готово:" in status:
            return status
        time.sleep(poll_sec)
    return last or "timeout"


def parse_shipment_done_body(body: str) -> dict[str, Any] | None:
    """Parse JSON or NDJSON stream from POST /api/v1/shipments/."""
    text = (body or "").strip()
    if not text:
        return None
    if text.startswith("{"):
        try:
            data = json.loads(text)
        except json.JSONDecodeError:
            data = None
        if isinstance(data, dict) and "items" in data:
            return data
    done: dict[str, Any] | None = None
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            continue
        if not isinstance(event, dict):
            continue
        if event.get("event") == "done":
            done = {k: v for k, v in event.items() if k != "event"}
        elif event.get("event") == "error":
            raise RuntimeError(event.get("detail") or "server error event")
    return done


def run_one_kit(
    page,
    base_url: str,
    kit: dict[str, Any],
    session_dir: Path,
    *,
    attempt: int = 1,
) -> KitResult:
    kit_id = kit["id"]
    kit_dir = KITS_ROOT / kit_id
    suffix = "" if attempt == 1 else f"_try{attempt}"
    run_dir = session_dir / f"{safe_name(kit_id)}_{kit['profile']}{suffix}"
    run_dir.mkdir(parents=True, exist_ok=True)
    result = KitResult(
        kit_id=kit_id,
        profile=kit["profile"],
        status="started",
        why=kit.get("why", ""),
        run_dir=str(run_dir),
        attempt=attempt,
        self_judge=bool(kit.get("self_judge")),
    )
    t0 = time.time()
    captured: dict[str, Any] = {"done": None, "raw": ""}

    def on_response(response) -> None:
        try:
            url = response.url or ""
            if response.request.method != "POST":
                return
            if "/api/v1/shipments/" not in url:
                return
            if "/export" in url:
                return
            body = response.text()
            captured["raw"] = body
            captured["done"] = parse_shipment_done_body(body)
        except Exception as exc:
            captured["error"] = f"{type(exc).__name__}: {exc}"

    page.on("response", on_response)
    try:
        goods, catalog = list_input_files(kit_dir)
        result.input_files = [p.name for p in goods]
        result.catalog_files = [p.name for p in catalog]
        dump_kit_sources(
            kit_dir,
            run_dir,
            goods,
            catalog,
            skip_etalon=bool(kit.get("self_judge") or kit.get("no_etalon")),
        )

        page.goto(base_url, wait_until="domcontentloaded", timeout=60_000)
        page.wait_for_selector("#create-form", timeout=30_000)
        page.select_option('select[name="profile_type"]', kit["profile"])
        title = f"audit-{kit_id}"
        page.fill('input[name="title"]', title)

        if goods:
            page.set_input_files("#files", [str(p) for p in goods])
        if catalog:
            page.set_input_files("#extra-files", [str(p) for p in catalog])

        page.wait_for_timeout(800)
        chips = page.locator("#file-list .file-chip").count()
        (run_dir / "ui_before_process.txt").write_text(
            f"profile={kit['profile']}\ngoods={result.input_files}\n"
            f"catalog={result.catalog_files}\nchips={chips}\n",
            encoding="utf-8",
        )
        if chips == 0:
            raise RuntimeError("после set_input_files список файлов пуст")

        page.click("#btn-process")
        status = wait_processing(page, timeout_sec=600.0, poll_sec=10.0)
        result.upload_status = status
        page.screenshot(path=str(run_dir / "ui_after_process.png"), full_page=True)

        if status.startswith("Ошибка") or status == "timeout" or status.startswith("timeout"):
            result.status = "failed"
            result.error = status
            result.verdict = "fail"
            if captured.get("raw"):
                (run_dir / "api_raw.txt").write_text(str(captured["raw"])[:200_000], encoding="utf-8")
            return result

        # Wait a bit more if stream body still buffering into captured.
        deadline = time.time() + 30
        while captured.get("done") is None and time.time() < deadline:
            time.sleep(1)

        ws = captured.get("done")
        if isinstance(ws, dict):
            (run_dir / "api_done.json").write_text(
                json.dumps(ws, ensure_ascii=False, indent=2, default=str)[:2_000_000],
                encoding="utf-8",
            )
        else:
            if captured.get("raw"):
                (run_dir / "api_raw.txt").write_text(str(captured["raw"])[:200_000], encoding="utf-8")
            # Fallback: scrape visible table (no full commercial_data).
            rows_ui = page.evaluate(
                """() => {
                  const out = [];
                  document.querySelectorAll('#items-table tbody tr').forEach((tr) => {
                    const tds = [...tr.querySelectorAll('td')].map(td => (td.dataset.full || td.innerText || '').trim());
                    if (tds.length) out.push(tds);
                  });
                  return out;
                }"""
            )
            (run_dir / "table_scrape.json").write_text(
                json.dumps(rows_ui, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            result.status = "no_api_payload"
            result.error = "не поймали done из POST /api/v1/shipments/"
            if captured.get("error"):
                result.error += f" ({captured['error']})"
            result.item_count = len(rows_ui or [])
            result.verdict = "fail"
            # Still try UI export so we have files to inspect.
            try:
                with page.expect_download(timeout=120_000) as dl_info:
                    page.click("#btn-export")
                download = dl_info.value
                zip_path = run_dir / (download.suggested_filename or f"{safe_name(kit_id)}.zip")
                download.save_as(str(zip_path))
                result.export_zip = str(zip_path)
                result.export_files = dump_export_zip(zip_path, run_dir)
            except Exception as export_exc:
                result.findings.append(f"export fallback failed: {export_exc}")
            result.verdict, result.findings = heuristic_verdict(kit, result)
            return result

        items = ws.get("items") or []
        result.item_count = int(ws.get("item_count") or len(items))
        result.warning_count = ws.get("warning_count")
        result.totals_got = sum_workspace_totals(items)
        mr = ws.get("model_review") or {}
        result.model_review = {
            "status": mr.get("status"),
            "model": mr.get("model") or mr.get("model_label"),
            "totals_mismatch": mr.get("totals_mismatch"),
            "error": mr.get("error"),
            "meaning": (mr.get("meaning") or "")[:400],
            "item_count_model": len(mr.get("items") or []) if isinstance(mr.get("items"), list) else None,
            "excel_attached": mr.get("excel_attached"),
            "image_count": mr.get("image_count"),
        }
        result.workspace_summary = {
            "title": ws.get("title"),
            "header": ws.get("header_fields"),
            "files": [
                {
                    "filename": f.get("filename"),
                    "parse_status": f.get("parse_status"),
                    "parse_message": f.get("parse_message"),
                    "doc_type": f.get("doc_type"),
                }
                for f in (ws.get("files") or [])
                if isinstance(f, dict)
            ],
            "rows": workspace_rows_brief(items),
        }
        (run_dir / "workspace.json").write_text(
            json.dumps(
                {
                    "meta": {
                        "id": ws.get("id"),
                        "title": ws.get("title"),
                        "profile_type": ws.get("profile_type"),
                        "status": ws.get("status"),
                        "item_count": result.item_count,
                        "warning_count": result.warning_count,
                    },
                    "files": result.workspace_summary["files"],
                    "header": ws.get("header_fields"),
                    "totals": result.totals_got,
                    "model_review": result.model_review,
                    "rows": result.workspace_summary["rows"],
                    "items_full": items,
                },
                ensure_ascii=False,
                indent=2,
                default=str,
            ),
            encoding="utf-8",
        )

        # Export via UI button (real user path).
        with page.expect_download(timeout=120_000) as dl_info:
            page.click("#btn-export")
        download = dl_info.value
        zip_path = run_dir / (download.suggested_filename or f"{safe_name(kit_id)}.zip")
        download.save_as(str(zip_path))
        result.export_zip = str(zip_path)
        result.export_files = dump_export_zip(zip_path, run_dir)

        export_status = ""
        try:
            export_status = page.locator("#export-result").inner_text(timeout=3000)
        except Exception:
            pass
        (run_dir / "export_status.txt").write_text(export_status, encoding="utf-8")

        result.status = "done"
        result.verdict, result.findings = heuristic_verdict(kit, result)
    except Exception as exc:
        result.status = "error"
        result.error = f"{type(exc).__name__}: {exc}"
        result.verdict = "fail"
        result.findings.append(result.error)
        try:
            page.screenshot(path=str(run_dir / "ui_error.png"), full_page=True)
        except Exception:
            pass
    finally:
        try:
            page.remove_listener("response", on_response)
        except Exception:
            pass
        result.elapsed_sec = time.time() - t0
        (run_dir / "result.json").write_text(
            json.dumps(asdict(result), ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
    return result


def main() -> int:
    ensure_dump_deps()
    parser = argparse.ArgumentParser(description="UI agent audit of live transformer site")
    parser.add_argument("--base-url", default="http://87.251.86.53:8010/")
    parser.add_argument("--limit", type=int, default=0, help="process only first N kits (0=all)")
    parser.add_argument("--only", default="", help="comma-separated kit id prefixes")
    parser.add_argument("--plan", choices=("round1", "round2"), default="round1")
    parser.add_argument("--retries", type=int, default=2, help="extra attempts after error (default 2)")
    parser.add_argument("--headed", action="store_true", help="show browser window")
    parser.add_argument("--out", default="", help="override output session dir")
    args = parser.parse_args()

    if not KITS_ROOT.is_dir():
        print(f"нет комплектов: {KITS_ROOT}", file=sys.stderr)
        return 2

    source_plan = KIT_PLAN_ROUND2 if args.plan == "round2" else KIT_PLAN
    only = {x.strip() for x in args.only.split(",") if x.strip()}
    plan = []
    for kit in source_plan:
        if only and not any(kit["id"].startswith(x) or x in kit["id"] for x in only):
            continue
        if not (KITS_ROOT / kit["id"]).is_dir():
            print(f"[skip missing] {kit['id']}")
            continue
        plan.append(kit)
    if args.limit:
        plan = plan[: args.limit]
    if not plan:
        print("пустой план", file=sys.stderr)
        return 2

    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    session_dir = Path(args.out) if args.out else OUT_ROOT / f"run_{stamp}"
    session_dir.mkdir(parents=True, exist_ok=True)
    findings_path = session_dir / "findings.md"
    findings_path.write_text(
        f"# Agent site audit\n\n"
        f"- base_url: {args.base_url}\n"
        f"- started: {stamp}\n"
        f"- plan: {args.plan}\n"
        f"- kits: {len(plan)}\n"
        f"- retries_on_error: {args.retries}\n"
        f"- note: расхождения только фиксируем, код парсера не правим в этом прогоне\n\n",
        encoding="utf-8",
    )
    (session_dir / "plan.json").write_text(
        json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    from playwright.sync_api import sync_playwright

    results: list[KitResult] = []
    max_attempts = 1 + max(0, args.retries)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not args.headed)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()
        page.set_default_timeout(60_000)
        for i, kit in enumerate(plan, 1):
            print(
                f"\n=== [{i}/{len(plan)}] {kit['id']} profile={kit['profile']} "
                f"self_judge={bool(kit.get('self_judge'))} ===",
                flush=True,
            )
            result: KitResult | None = None
            for attempt in range(1, max_attempts + 1):
                if attempt > 1:
                    print(f"  retry {attempt - 1}/{args.retries} after error…", flush=True)
                    time.sleep(3)
                    try:
                        page.goto(args.base_url.rstrip("/") + "/", wait_until="domcontentloaded", timeout=60_000)
                    except Exception as nav_exc:
                        print(f"  reload failed: {nav_exc}", flush=True)
                result = run_one_kit(
                    page,
                    args.base_url.rstrip("/") + "/",
                    kit,
                    session_dir,
                    attempt=attempt,
                )
                result.attempt = attempt
                result.attempts_total = max_attempts
                result.self_judge = bool(kit.get("self_judge"))
                result.error_kind = classify_error(result.error or result.upload_status)
                (Path(result.run_dir) / "result.json").write_text(
                    json.dumps(asdict(result), ensure_ascii=False, indent=2, default=str),
                    encoding="utf-8",
                )
                if not needs_retry(result):
                    break
                print(
                    f"  attempt {attempt} fail kind={result.error_kind or 'unknown'} "
                    f"err={result.error or result.upload_status}",
                    flush=True,
                )
            assert result is not None
            results.append(result)
            append_findings_md(findings_path, result, kit)
            print(
                f"→ verdict={result.verdict} items={result.item_count} "
                f"attempt={result.attempt}/{result.attempts_total} "
                f"kind={result.error_kind or '-'} elapsed={result.elapsed_sec:.0f}s "
                f"error={result.error}",
                flush=True,
            )
        browser.close()

    summary = {
        "base_url": args.base_url,
        "session_dir": str(session_dir),
        "plan": args.plan,
        "results": [asdict(r) for r in results],
        "counts": {
            "ok_candidate": sum(1 for r in results if r.verdict == "ok_candidate"),
            "review": sum(1 for r in results if r.verdict == "review"),
            "fail": sum(1 for r in results if r.verdict == "fail"),
        },
    }
    (session_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    with findings_path.open("a", encoding="utf-8") as fh:
        fh.write("\n## Summary\n")
        fh.write(json.dumps(summary["counts"], ensure_ascii=False) + "\n")
        fh.write("\n## Error kinds\n")
        for r in results:
            if r.verdict == "fail" or r.error:
                fh.write(
                    f"- {r.kit_id}: kind={r.error_kind or 'n/a'} "
                    f"attempts={r.attempt}/{r.attempts_total} "
                    f"err={r.error or r.upload_status}\n"
                )
    print(f"\nDONE session={session_dir}")
    print(json.dumps(summary["counts"], ensure_ascii=False))
    return 0 if summary["counts"]["fail"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
