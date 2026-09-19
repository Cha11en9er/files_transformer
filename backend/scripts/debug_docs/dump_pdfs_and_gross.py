"""Dump PDFs with pypdf + extract docx images + compare net/gross totals."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[3] / "documents" / "4_pravka"
OUT = Path(__file__).resolve().parent / "out"
IMG = OUT / "docx_images"


def dump_pdfs() -> None:
    lines: list[str] = []
    for path in sorted(ROOT.rglob("*.pdf")):
        rel = path.relative_to(ROOT)
        lines.append("=" * 90)
        lines.append(f"PDF: {rel}  ({path.stat().st_size} bytes)")
        try:
            reader = PdfReader(str(path))
            lines.append(f"pages={len(reader.pages)} encrypted={reader.is_encrypted}")
            for i, page in enumerate(reader.pages[:4]):
                text = (page.extract_text() or "").strip()
                lines.append(f"-- page {i+1} ({len(text)} chars)")
                if not text:
                    lines.append("   [NO TEXT LAYER]")
                    continue
                for row in text.splitlines()[:90]:
                    lines.append("   | " + row[:160])
        except Exception as exc:
            lines.append(f"[failed: {exc}]")
    (OUT / "93_pdfs.txt").write_text("\n".join(lines), encoding="utf-8")
    print("wrote 93_pdfs.txt", len(lines))


def dump_docx_images() -> None:
    IMG.mkdir(parents=True, exist_ok=True)
    index: list[str] = []
    n = 0
    for path in sorted(ROOT.rglob("*.docx")):
        rel = path.relative_to(ROOT)
        with zipfile.ZipFile(path) as zf:
            media = [n for n in zf.namelist() if n.startswith("word/media/")]
            index.append(f"{rel}: {len(media)} images")
            for mi, name in enumerate(media, start=1):
                data = zf.read(name)
                ext = Path(name).suffix or ".bin"
                slug = re.sub(r"[^\w.\-]+", "_", str(rel), flags=re.UNICODE)[:80]
                dest = IMG / f"{slug}__{mi}{ext}"
                dest.write_bytes(data)
                n += 1
                index.append(f"  {dest.name} {len(data)} bytes")
    (OUT / "94_docx_images.txt").write_text("\n".join(index), encoding="utf-8")
    print("extracted", n, "images")


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace(" ", ""))
    except Exception:
        return None


def sheet_totals(path: Path, sheet: str, header_needles: dict[str, tuple[str, ...]]) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row = None
    colmap: dict[str, int] = {}
    for r in range(1, min(ws.max_row or 1, 40) + 1):
        vals = [str(ws.cell(r, c).value or "").lower() for c in range(1, min(ws.max_column or 1, 20) + 1)]
        blob = " | ".join(vals)
        hits = 0
        cmap = {}
        for key, needles in header_needles.items():
            for i, v in enumerate(vals, start=1):
                if any(n in v for n in needles):
                    cmap[key] = i
                    hits += 1
                    break
        if hits >= 3:
            header_row = r
            colmap = cmap
            break
    body = []
    total = None
    if header_row:
        for r in range(header_row + 1, (ws.max_row or 0) + 1):
            first = str(ws.cell(r, 1).value or "")
            row = {k: _num(ws.cell(r, c).value) for k, c in colmap.items()}
            if first.upper().startswith("TOTAL") or "итого" in first.lower():
                total = row
                break
            if any(v is not None for v in row.values()):
                body.append((first, row))
    summed = {k: round(sum(r[k] or 0 for _, r in body), 4) for k in colmap}
    wb.close()
    return {
        "file": path.name,
        "sheet": ws.title,
        "cols": colmap,
        "n": len(body),
        "sum": summed,
        "total_row": total,
    }


def compare_beijing() -> None:
    src = next(ROOT.rglob("*003-26E*00.xlsx"))
    export = next((ROOT / "для тест" / "BEIJING GOLDLUCK CO., LTD").glob("*beijing_export*"))
    needles_inv = {
        "qty": ("quantity",),
        "amount": ("amount",),
        "gross": ("gross", "брутто"),
        "net": ("net wt", "net wt.", "нетто", "netto"),
        "cartons": ("carton",),
    }
    needles_pl = {
        "qty": ("quantity",),
        "gross": ("gross",),
        "net": ("net wt", "net wt.", "нетто"),
        "cartons": ("carton",),
    }
    blocks = [
        sheet_totals(src, "Invoice + Packing list", needles_inv),
        sheet_totals(src, "Packing list", needles_pl),
        sheet_totals(export, "Invoice", needles_inv),
        sheet_totals(export, "Packing list", needles_pl),
        sheet_totals(export, "Specification", {
            "qty": ("quantity, unit", "количество, единиц"),
            "amount": ("amount", "стоимость"),
            "gross": ("gross", "брутто"),
            "net": ("netto", "нетто"),
            "cartons": ("ctns", "коробок"),
        }),
    ]
    text = []
    for b in blocks:
        text.append(f"{b['file']} / {b['sheet']} rows={b['n']} cols={b['cols']}")
        text.append(f"  SUM   {b['sum']}")
        text.append(f"  TOTAL {b['total_row']}")
        text.append("")
    (OUT / "95_beijing_net_gross.txt").write_text("\n".join(text), encoding="utf-8")
    print("wrote 95_beijing_net_gross.txt")


def tsd_spec_gross() -> None:
    files = [
        next((ROOT / "для тест" / "bestway").glob("*.xlsx")),
        next((ROOT / "для тест" / "besway 2").glob("*.xlsx")),
        next((ROOT / "для тест" / "Матрац").glob("*.xlsx")),
        next((ROOT / "для тест" / "шары").glob("*.xlsx")),
    ]
    needles = {
        "qty": ("кол-во", "qty"),
        "gross": ("брутто", "gross", "brutto"),
        "net": ("нетто", "netto", "net"),
        "amount": ("ст-сть", "amount", "стоимость"),
        "pkg": ("package", "мест"),
    }
    out = []
    for path in files:
        wb = load_workbook(path, data_only=True)
        out.append(f"FILE {path.name} sheets={wb.sheetnames}")
        for name in wb.sheetnames:
            ws = wb[name]
            out.append(f"  sheet {name} {ws.max_row}x{ws.max_column}")
        wb.close()
        if "Specification" in load_workbook(path, read_only=True).sheetnames:
            out.append(str(sheet_totals(path, "Specification", needles)))
        if "PAK" in load_workbook(path, read_only=True).sheetnames:
            out.append("PAK " + str(sheet_totals(path, "PAK", {
                "qty": ("qty",),
                "gross": ("brutto", "брутто", "gross"),
                "net": ("netto", "нетто", "net"),
                "pkg": ("package",),
            })))
        out.append("")
    (OUT / "96_tsd_sheets_gross.txt").write_text("\n".join(out), encoding="utf-8")
    print("wrote 96_tsd_sheets_gross.txt")


if __name__ == "__main__":
    dump_pdfs()
    dump_docx_images()
    compare_beijing()
    tsd_spec_gross()
