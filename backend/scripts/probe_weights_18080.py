"""Compare PL vs sender-spec vs etalon vs our export weights for 18080 kits."""
from __future__ import annotations

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from app.transform.extract import extract_sheet
from app.transform.merge import _aggregate_detail, _family_model, match_key
from app.transform.reader import read_workbook
from app.transform.service import canonical_to_rows, transform_paths

ROOT = Path(__file__).resolve().parents[2] / "documents" / "3_pravka" / "18080"
SRC = ROOT / "Исходные"
ETALON = ROOT / "Готовые"
MINE = ROOT / "Готовые мои"


def _load_spec_weights(path: Path) -> dict[str, tuple[float | None, float | None]]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = None
    cols: dict[str, int] = {}
    for r in range(1, 40):
        headers = [str(ws.cell(r, c).value or "") for c in range(1, 20)]
        blob = " ".join(headers).lower()
        if "артикул" in blob and ("нетто" in blob or "n.w" in blob):
            header_row = r
            for c, h in enumerate(headers, start=1):
                hl = h.lower()
                if "артикул" in hl or hl.startswith("art"):
                    cols["a"] = c
                if "нетто" in hl or "n.w" in hl:
                    cols["n"] = c
                if "брутто" in hl or "g.w" in hl:
                    cols["g"] = c
            break
    out: dict[str, tuple[float | None, float | None]] = {}
    if header_row is None or "a" not in cols:
        return out
    for r in range(header_row + 1, (ws.max_row or 0) + 1):
        a = ws.cell(r, cols["a"]).value
        if not a:
            continue
        text = str(a)
        if text.lower().startswith("total") or text.lower().startswith("итого"):
            continue
        n = ws.cell(r, cols["n"]).value if "n" in cols else None
        g = ws.cell(r, cols["g"]).value if "g" in cols else None
        out[match_key(text)] = (
            float(n) if isinstance(n, (int, float)) else None,
            float(g) if isinstance(g, (int, float)) else None,
        )
    return out


def analyze(prefix: str) -> None:
    inv = next(SRC.glob(f"{prefix}*INVOICE*"))
    pl = next(SRC.glob(f"{prefix}*PL*"))
    sp = next(SRC.glob(f"{prefix}*Specification*"))

    pl_fam: dict[str, tuple[str, float | None, float | None]] = {}
    for sheet in read_workbook(str(pl)):
        ex = extract_sheet(sheet)
        if not ex:
            continue
        for row in ex.rows:
            fam = _family_model(row.article) or row.article
            pl_fam[match_key(fam)] = (
                row.article,
                row.fields.get("net_weight") if isinstance(row.fields.get("net_weight"), (int, float)) else None,
                row.fields.get("gross_weight") if isinstance(row.fields.get("gross_weight"), (int, float)) else None,
            )

    agg: dict[str, dict] = {}
    for sheet in read_workbook(str(sp)):
        ex = extract_sheet(sheet)
        if not ex:
            continue
        print(prefix, "spec role=", ex.role, "detail=", ex.detail, "rows=", len(ex.rows))
        agg = _aggregate_detail(ex.rows) if ex.detail else {
            match_key(r.article): {"article": r.article, **r.fields} for r in ex.rows
        }

    res = transform_paths([(str(inv), inv.name), (str(pl), pl.name), (str(sp), sp.name)])
    rows = canonical_to_rows(res.items)

    etalon_files = list(ETALON.glob(f"*{prefix}*СПЕЦ*"))
    et = _load_spec_weights(etalon_files[0]) if etalon_files else {}
    mine_files = list((MINE / prefix).rglob("*СПЕЦ*")) if (MINE / prefix).exists() else []
    mine = _load_spec_weights(mine_files[0]) if mine_files else {}

    print(f"=== {prefix}: transform={len(rows)} etalon={len(et)} mine={len(mine)} pl_families={len(pl_fam)}")

    # Family-level: sum of children vs PL family vs sum of spec rolls
    for fam_key, (fam_name, pn, pg) in sorted(pl_fam.items(), key=lambda x: x[0]):
        children = [r for r in rows if match_key(r["article"]).startswith(fam_key)]
        if not children:
            continue
        our_n = sum(float(r["packing_data"].get("net_weight") or 0) for r in children)
        our_g = sum(float(r["packing_data"].get("gross_weight") or 0) for r in children)
        spec_n = sum(float(agg.get(match_key(r["article"]), {}).get("net_weight") or 0) for r in children)
        spec_g = sum(float(agg.get(match_key(r["article"]), {}).get("gross_weight") or 0) for r in children)
        et_n = sum(float(et.get(match_key(r["article"]), (0, 0))[0] or 0) for r in children)
        et_g = sum(float(et.get(match_key(r["article"]), (0, 0))[1] or 0) for r in children)
        dn_ps = abs((pn or 0) - spec_n)
        dn_oe = abs(our_n - et_n) if et else None
        if dn_ps > 0.05 or (dn_oe is not None and dn_oe > 0.05) or abs(our_n - (pn or 0)) > 0.05:
            print(
                f"  FAM {fam_name}: PL={pn}/{pg} "
                f"spec_sum={round(spec_n,2)}/{round(spec_g,2)} "
                f"ours_sum={round(our_n,2)}/{round(our_g,2)} "
                f"etalon_sum={round(et_n,2)}/{round(et_g,2)} "
                f"|PL-spec|={round(dn_ps,2)}"
            )

    print("  per-article ours vs etalon (delta>0.05):")
    for r in rows:
        k = match_key(r["article"])
        pn = r["packing_data"].get("net_weight")
        pg = r["packing_data"].get("gross_weight")
        en, eg = et.get(k, (None, None))
        mn, mg = mine.get(k, (None, None))
        sn = agg.get(k, {}).get("net_weight")
        if not isinstance(en, (int, float)) or not isinstance(pn, (int, float)):
            continue
        if abs(pn - en) > 0.05 or (isinstance(pg, (int, float)) and isinstance(eg, (int, float)) and abs(pg - eg) > 0.05):
            print(
                f"    {r['article']}: ours={pn}/{pg} etalon={en}/{eg} mine={mn}/{mg} "
                f"spec_agg={sn} dN={round(pn-en,2)} dG={round((pg or 0)-(eg or 0),2)}"
            )


if __name__ == "__main__":
    for prefix in ("621-1", "621-2", "624-1", "624-2"):
        analyze(prefix)
        print()
