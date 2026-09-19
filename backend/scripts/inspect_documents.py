"""Ad-hoc inspection: dump every supply file (Excel + PDF) so we can design the
transformer from the real data, not from assumptions. Output is UTF-8 text.

Run:
    uv run --with pandas --with openpyxl --with xlrd --with pdfplumber \
        python scripts/inspect_documents.py
"""

from __future__ import annotations

import sys
from pathlib import Path

DOCS = Path(__file__).resolve().parents[2] / "documents"
OUT = Path(__file__).resolve().parents[1] / "_inspect_out.txt"

MAX_ROWS = 30
MAX_COLS = 22


def _p(fh, *args):
    print(*args, file=fh)


def dump_xlsx(fh, path: Path) -> None:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        _p(fh, f"    [openpyxl failed: {exc}]")
        return
    for ws in wb.worksheets:
        merges = list(ws.merged_cells.ranges)
        _p(fh, f"  -- sheet '{ws.title}'  dims={ws.max_row}x{ws.max_column}  merges={len(merges)}")
        if merges:
            _p(fh, "     merged: " + ", ".join(str(m) for m in merges[:20]))
        for r in range(1, min(ws.max_row, MAX_ROWS) + 1):
            cells = []
            for c in range(1, min(ws.max_column, MAX_COLS) + 1):
                v = ws.cell(r, c).value
                cells.append("" if v is None else str(v).replace("\n", " ")[:22])
            if any(cells):
                _p(fh, f"     r{r:>3}| " + " | ".join(cells))
    wb.close()


def dump_xls(fh, path: Path) -> None:
    import pandas as pd

    try:
        book = pd.read_excel(path, sheet_name=None, header=None, dtype=object, engine="xlrd")
    except Exception as exc:
        try:
            book = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        except Exception as exc2:
            _p(fh, f"    [xls failed: {exc} / {exc2}]")
            return
    for name, frame in book.items():
        _p(fh, f"  -- sheet '{name}'  dims={frame.shape}")
        for r in range(min(len(frame), MAX_ROWS)):
            cells = []
            for c in range(min(frame.shape[1], MAX_COLS)):
                v = frame.iat[r, c]
                cells.append("" if v is None or (isinstance(v, float) and v != v) else str(v).replace("\n", " ")[:22])
            if any(cells):
                _p(fh, f"     r{r:>3}| " + " | ".join(cells))


def dump_pdf(fh, path: Path) -> None:
    try:
        import pdfplumber
    except Exception as exc:
        _p(fh, f"    [pdfplumber missing: {exc}]")
        return
    try:
        with pdfplumber.open(path) as pdf:
            _p(fh, f"  -- pages={len(pdf.pages)}")
            for i, page in enumerate(pdf.pages[:3]):
                text = (page.extract_text() or "").strip()
                _p(fh, f"  -- page {i+1} text ({len(text)} chars):")
                for line in text.splitlines()[:60]:
                    _p(fh, "     | " + line[:120])
                tables = page.extract_tables() or []
                _p(fh, f"  -- page {i+1} tables={len(tables)}")
                for ti, t in enumerate(tables[:2]):
                    _p(fh, f"     table {ti+1} ({len(t)} rows):")
                    for row in t[:12]:
                        _p(fh, "       " + " | ".join("" if x is None else str(x).replace("\n", " ")[:18] for x in row[:MAX_COLS]))
    except Exception as exc:
        _p(fh, f"    [pdf failed: {exc}]")


def main() -> None:
    files = sorted(
        [p for p in DOCS.rglob("*") if p.is_file() and p.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".pdf"}],
        key=lambda p: str(p).lower(),
    )
    with OUT.open("w", encoding="utf-8") as fh:
        _p(fh, f"documents root: {DOCS}")
        _p(fh, f"total files: {len(files)}\n")
        for path in files:
            rel = path.relative_to(DOCS)
            _p(fh, "=" * 100)
            _p(fh, f"FILE: {rel}  ({path.stat().st_size} bytes)")
            suffix = path.suffix.lower()
            try:
                if suffix in {".xlsx", ".xlsm"}:
                    dump_xlsx(fh, path)
                elif suffix == ".xls":
                    dump_xls(fh, path)
                elif suffix == ".pdf":
                    dump_pdf(fh, path)
            except Exception as exc:
                _p(fh, f"    [ERROR: {exc}]")
    print(f"wrote {OUT}")


if __name__ == "__main__":
    sys.exit(main())
