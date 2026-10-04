import unittest
from pathlib import Path

from prepare_transform_code.fields import LOT_FIELDS, column_of, is_header_row, is_stop_label, party_after
from prepare_transform_code.numbers import currencies_of, parse_number, row_is_doubled, undouble_row
from prepare_transform_code.shipment import _INVOICE_NO, _one_label, analyze

ROOT = Path(__file__).resolve().parents[2] / "documents_final"


def load(name):
    return analyze(ROOT / name)


class HeaderContinuationTest(unittest.TestCase):
    def test_unit_word_next_to_a_name_is_a_goods_row(self):
        from prepare_transform_code.pdfdoc import _header_continuation, stated_from_text

        row = [
            "110.60",
            "METERS",
            "BLOOM (BLOOM) QUALITY UPHOLSTERY FABRIC",
            "54.07.73.00.90.11",
            "ROLLS",
            "3",
            "12.38",
            "1,369.23",
        ]
        self.assertFalse(_header_continuation(row))
        stated = stated_from_text(
            "3,654.50 METERS ROLLS 101 41,431.38\n"
            "TOTAL WEIGHT : 3,047.000 KGS NET / 3,122.500 KGS GROSS\n"
        )
        self.assertAlmostEqual(stated["pieces"], 3654.50)
        self.assertAlmostEqual(stated["packages"], 101)
        self.assertAlmostEqual(stated["amount"], 41431.38)
        self.assertAlmostEqual(stated["net"], 3047)
        self.assertAlmostEqual(stated["gross"], 3122.5)


class NumbersTest(unittest.TestCase):
    def test_money_formats(self):
        self.assertEqual(parse_number("US$55,000.00"), 55000)
        self.assertEqual(parse_number("¥1 267,00"), 1267)
        self.assertEqual(parse_number("¥291 557,00"), 291557)
        self.assertAlmostEqual(parse_number("6809.6"), 6809.6)
        self.assertAlmostEqual(parse_number("21.03 Cbm"), 21.03)
        self.assertEqual(parse_number("3.900"), 3900)
        self.assertEqual(parse_number("22.000"), 22000)
        self.assertEqual(parse_number("1.000"), 1000)
        self.assertAlmostEqual(parse_number("5.952,40"), 5952.40)
        self.assertAlmostEqual(parse_number("5.952,\n40"), 5952.40)
        self.assertAlmostEqual(parse_number("31.467\n,80"), 31467.80)
        self.assertAlmostEqual(parse_number("51,76"), 51.76)
        self.assertEqual(parse_number("80,00000"), 80)
        self.assertEqual(parse_number("20 000,0"), 20000)
        from prepare_transform_code.numbers import _COMMA_IS_DECIMAL

        token = _COMMA_IS_DECIMAL.set(True)
        try:
            self.assertAlmostEqual(parse_number("722,820"), 722.82)
            self.assertAlmostEqual(parse_number("1251,040"), 1251.04)
            self.assertAlmostEqual(parse_number("516,30"), 516.30)
        finally:
            _COMMA_IS_DECIMAL.reset(token)
        self.assertEqual(parse_number("722,820"), 722820)


class LanguageTest(unittest.TestCase):
    def test_labels_are_not_tied_to_two_languages(self):
        self.assertEqual(column_of("Bezeichnung"), "description")
        self.assertEqual(column_of("Menge"), "qty")
        self.assertEqual(column_of("Gesamtpreis"), "amount")
        self.assertEqual(column_of("Item number"), "vendor")
        self.assertEqual(column_of("Customs code"), "hs")
        self.assertEqual(column_of("Country of origi"), "origin")
        self.assertEqual(column_of("Q\u2010ty"), "qty")
        self.assertTrue(is_stop_label("Gesamt"))
        self.assertTrue(is_stop_label("Suma"))
        self.assertTrue(is_header_row(["Item number", "Contents", "Quantity", "Price"]))
        self.assertFalse(
            is_header_row(
                ["1", "M-G(2100) Automatic nonwoven fabrics folding machine", "5 Packages", "2980.00 Kgs"]
            )
        )
        self.assertEqual(currencies_of("Price EUR, przeliczenie PLN 4.63"), ["EUR", "PLN"])
        doubled = ["JJAALLAASS 11661155-4433", "IInnssoolleess", "22", "33,4455"]
        self.assertTrue(row_is_doubled(doubled))
        self.assertFalse(row_is_doubled(["Safetyy shoes", "1", "41,61"]))
        self.assertEqual(undouble_row(doubled)[0], "JALAS 1615-43")
        parties = party_after("Verkäufer\nEjendals AB\nBuyer adress\nEjendals LLC Ejendals LLC")
        self.assertEqual(parties["seller"], "Ejendals AB")
        self.assertEqual(parties["buyer"], "Ejendals LLC")
        same_line = party_after("Покупатель ООО «ПО ОКТЯБРЬ» Продавец YANTAI FOREVER TRADE CO., LTD.")
        self.assertIn("ОКТЯБРЬ", same_line["buyer"])
        self.assertIn("YANTAI", same_line["seller"])
        self.assertNotIn("Продавец", same_line["buyer"])
        columns = party_after(
            "Buyer: Seller:\n"
            "“SM REGIONTEKSTIL'” LLC HANGZHOU ZHONGFANG TEXTILE IMP./EXP. CO., LTD."
        )
        self.assertIn("REGIONTEKSTIL", columns["buyer"])
        self.assertNotIn("HANGZHOU", columns["buyer"])
        self.assertIn("HANGZHOU", columns["seller"])
        self.assertNotIn("REGIONTEKSTIL", columns["seller"])
        bilingual = party_after(
            "Buyer/Покупатель: Seller / Продавец:\n"
            "“SM REGIONTEKSTIL'” LLC / “СМ Регионтекстиль” ООО HANGZHOU ZHONGFANG TEXTILE IMP./EXP. CO., LTD."
        )
        self.assertIn("REGIONTEKSTIL", bilingual["buyer"])
        self.assertNotIn("HANGZHOU", bilingual["buyer"])
        self.assertIn("HANGZHOU", bilingual["seller"])
        self.assertNotIn("REGIONTEKSTIL", bilingual["seller"])
        self.assertNotIn("Регион", bilingual["seller"])
        self.assertTrue(bilingual["seller"].startswith("HANGZHOU"))
        self.assertEqual(column_of("Код изделия"), "vendor")
        self.assertEqual(column_of("Item Code"), "vendor")
        self.assertEqual(column_of("Цвет /Тип покрытия"), "finish")
        self.assertEqual(column_of("Количество мест"), "packages")
        self.assertEqual(column_of("Количество упаковок"), "packages")
        self.assertEqual(column_of("Number of packages"), "packages")
        self.assertEqual(column_of("Name of product"), "description")
        self.assertEqual(column_of("Количество, шт."), "qty")
        named = party_after(
            'Компания HUANAN INT\'L HARDWARE CO., LIMITED в лице директора, именуемая в дальнейшем "Продавец", '
            'с одной стороны, и Общество с ограниченной ответственностью «Респект», в лице директора, '
            'именуемая в дальнейшем "Покупатель", с другой стороны'
        )
        self.assertIn("HUANAN", named["seller"])
        self.assertIn("Респект", named["buyer"])
        self.assertNotIn("HUANAN", named["buyer"])
        self.assertEqual(column_of("Наименование модели// Style"), "model")
        self.assertEqual(column_of("Производитель, наименование и характеристики товара"), "description")
        self.assertEqual(column_of("Shipper's Custom Tariff"), "hs")
        self.assertEqual(column_of("Таможенный код поставщика"), "hs")
        sides = party_after(
            "Покупатель//Buyer:\nПродавец//Seller:\n"
            'ООО "Обувь XXI века"\nAIRWAIR INTERNATIONAL LTD'
        )
        self.assertIn("Обувь", sides["buyer"])
        self.assertIn("AIRWAIR", sides["seller"])
        self.assertNotIn("AIRWAIR", sides["buyer"])
        self.assertEqual(currencies_of("Felix Rub Off, Price for pair, USD"), ["USD"])
        self.assertEqual(currencies_of("Vessel EURO MAX, price $28.48"), ["USD"])
        self.assertIn("EUR", currencies_of("amount 100 EURO"))
        self.assertEqual(column_of("H.S. Code"), "hs")
        self.assertIsNone(column_of("CBM/CTN"))
        self.assertEqual(column_of("MEASUREMENT (CBM)"), "volume")
        self.assertEqual(column_of("Measurement"), "volume")
        self.assertEqual(column_of("Volume (m3)"), "volume")
        self.assertEqual(column_of("Art No."), "vendor")
        self.assertIsNone(column_of("Part No."))
        self.assertEqual(column_of("Gross Wt. (kg)"), "gross")
        self.assertEqual(column_of("Net Wt. (kg)"), "net")
        self.assertIsNone(column_of("UNIT G.W.(KGS)"))
        self.assertEqual(column_of("UNIT N.W.(KGS)"), "unit_net")
        self.assertEqual(column_of("G.W.(KGS)"), "gross")
        self.assertEqual(column_of("N.W.(KGS)"), "net")
        bank = party_after("Seller's bank information\nBeneficiary: Shinegarden Shoes Limited")
        self.assertNotIn("bank", bank.get("seller", "").lower())
        self.assertEqual(column_of("CODE"), "vendor")
        self.assertEqual(column_of("HS CODE"), "hs")
        self.assertEqual(column_of("BOX QTY KOLI ADEDI"), "packages")
        self.assertEqual(column_of("GOODS PCS"), "qty")
        self.assertEqual(column_of("TOTAL NETT KG"), "net")
        self.assertEqual(column_of("TOTAL KG"), "gross")
        self.assertIsNone(column_of("EACH BOX KG"))
        bilingual = party_after('Buyer:\nПокупатель:\nLLC "RESPECT"\nООО «РЕСПЕКТ»')
        self.assertIn("RESPECT", bilingual["buyer"])
        self.assertNotEqual(bilingual["buyer"], "Покупатель")
        prose = party_after(
            "In the event that the Seller fails to fulfill the order\nSeller:\nACME MACHINERY CO., LTD"
        )
        self.assertIn("ACME", prose["seller"])
        self.assertNotIn("fails", prose["seller"].lower())
        self.assertEqual(column_of("BRUT KG"), "gross")
        self.assertEqual(column_of("PALLET"), "pallet")
        self.assertEqual(column_of("ТМ"), "brand")
        self.assertEqual(column_of("with pallet"), "gross_with_pallet")
        self.assertIn("VEK", party_after('BYUER: "TD 21 VEK" LLC')["buyer"])
        self.assertEqual(
            _one_label(
                "ORIGIN: TURKEY / СТРАНА ПРОИСХОЖДЕНИЯ ТОВАРА: ТУРЦИЯ",
                r"(?:country\s+of\s+origin|origin)\s*:",
            ),
            "TURKEY / ТУРЦИЯ",
        )
        self.assertIsNone(column_of("размер коробки"))
        self.assertIsNone(column_of("G.W./CTN"))
        self.assertEqual(column_of("N.W./CTN"), "unit_net")
        self.assertEqual(column_of("нетто коробки"), "unit_net")
        self.assertIsNone(column_of("Qty/Ctn/кол-во в коробке"))
        self.assertEqual(column_of("CTNS"), "packages")
        self.assertEqual(column_of("Box/коробка"), "packages")
        self.assertEqual(column_of("Specifications"), "description")
        self.assertEqual(column_of("Net weight kg 1 pair"), "unit_net")
        self.assertEqual(currencies_of("SAY TOTAL: YUAN ONLY"), ["CNY"])
        buyer = party_after('BUYER: LIMITED LIABILITY COMPANY "TD 21 VEK" DATE: 22.12.2025')
        self.assertIn("TD 21 VEK", buyer["buyer"])
        self.assertNotIn("DATE", buyer["buyer"])
        self.assertEqual(_INVOICE_NO.search("No INVOICE RU30006").group(1), "RU30006")
        self.assertEqual(_INVOICE_NO.search("INVOICE NO: SG1251").group(1), "SG1251")
        self.assertEqual(_INVOICE_NO.search("INVOICE No. CI240301RU").group(1), "CI240301RU")


class NoTemplateBranchTest(unittest.TestCase):
    def test_reader_does_not_name_shipments(self):
        package = Path(__file__).resolve().parents[1]
        for path in package.rglob("*.py"):
            if path.name.startswith("test_"):
                continue
            text = path.read_text(encoding="utf-8")
            self.assertNotIn("Шаблон", text, path.name)
            self.assertNotIn("TIG200", text, path.name)
            self.assertNotIn("JH21024", text, path.name)


class Shipment1Test(unittest.TestCase):
    def test_one_machine(self):
        result = load("Шаблон 1")
        roles = {doc["role"] for doc in result["documents"]}
        self.assertIn("invoice", roles)
        self.assertIn("packing", roles)
        self.assertNotIn("duplicate", roles)
        self.assertEqual(len(result["lots"]), 1)
        lot = result["lots"][0]
        self.assertEqual(lot["pieces"], 1)
        self.assertEqual(lot["packages"], 5)
        self.assertEqual(lot["net"], 2980)
        self.assertEqual(lot["gross"], 3270)
        self.assertEqual(lot["volume"], 21.03)
        self.assertEqual(lot["amount"], 55000)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["invoice_no"], "JH21024")


class Shipment2Test(unittest.TestCase):
    def test_seven_lots_and_weight_conflict(self):
        result = load("Шаблон 2")
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 7)
        pieces = sorted(lot["pieces"] for lot in goods)
        self.assertEqual(pieces, [60, 60, 200, 240, 650, 1000, 2840])
        self.assertIn("weight_conflict", result["flags"])
        self.assertTrue(all(lot["gross"] is None for lot in goods))


class Shipment3Test(unittest.TestCase):
    def test_fabric(self):
        result = load("Шаблон 3")
        self.assertEqual(len(result["lots"]), 1)
        lot = result["lots"][0]
        self.assertAlmostEqual(lot["pieces"], 6809.6)
        self.assertEqual(lot["packages"], 178)
        self.assertAlmostEqual(lot["net"], 5154.9)
        self.assertAlmostEqual(lot["gross"], 5332.9)
        self.assertAlmostEqual(lot["volume"], 12.5)
        self.assertAlmostEqual(lot["price"], 6.25)
        self.assertEqual(lot["amount"], 42560)
        self.assertEqual(result["invoice_no"], "HF20230106")
        self.assertIn("POLYESTER", lot["description"].upper())


class Shipment4Test(unittest.TestCase):
    def test_duplicate_scan_and_spec_lots(self):
        result = load("Шаблон 4")
        duplicates = [doc for doc in result["documents"] if doc["role"] == "duplicate"]
        scans = [doc for doc in result["documents"] if doc["role"] == "scan"]
        self.assertEqual(len(duplicates), 1)
        self.assertEqual(len(scans), 1)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        articles = [lot["vendor"] for lot in goods]
        self.assertEqual(sum(1 for article in articles if str(article).startswith("1971")), 2)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 2261)
        self.assertEqual(result["currency"], "EUR")
        self.assertTrue(result["freights"])


class Shipment5Test(unittest.TestCase):
    def test_join_by_model_not_row_order(self):
        result = load("Шаблон 5")
        self.assertEqual(result["currency"], "CNY")
        self.assertEqual(result["invoice_no"], "CI240301RU")
        by_model = {lot["model"]: lot for lot in result["lots"]}
        self.assertEqual(set(by_model), {"TIG200Pro", "TIG1000", "TIG630", "ASAW1600II", "Axis"})
        tig1000 = by_model["TIG1000"]
        self.assertEqual(tig1000["pieces"], 2)
        self.assertEqual(tig1000["packages"], 2)
        self.assertEqual(tig1000["gross"], 234)
        self.assertEqual(tig1000["net"], 182)
        self.assertEqual(tig1000["price"], 11200)
        axis = by_model["Axis"]
        self.assertEqual(axis["pieces"], 2)
        self.assertEqual(axis["price"], 540)
        shared = by_model["TIG200Pro"]["measure_group"]
        self.assertIsNotNone(shared)
        self.assertEqual(shared["gross"], 404)
        self.assertEqual(shared["net"], 357)
        self.assertIs(by_model["TIG200Pro"]["measure_group"], by_model["Axis"]["measure_group"])
        self.assertEqual(by_model["TIG200Pro"]["pieces"], 35)
        self.assertIsNone(by_model["TIG200Pro"]["gross"])
        groups = {}
        own = 0
        for lot in result["lots"]:
            if lot["measure_group"]:
                groups[id(lot["measure_group"])] = lot["measure_group"]["gross"]
            else:
                own += lot["gross"] or 0
        self.assertEqual(own + sum(groups.values()), 3079)


class Shipment6Test(unittest.TestCase):
    def test_split_article_and_pallet_gross(self):
        result = load("Шаблон 6")
        roles = {}
        for doc in result["documents"]:
            roles.setdefault(doc["role"], []).append(doc)
        self.assertEqual(len(roles["proforma"]), 1)
        self.assertEqual(roles["proforma"][0]["line_count"], 9)
        self.assertEqual(roles["packing"][0]["line_count"], 10)
        self.assertEqual(roles["specification"][0]["line_count"], 10)
        self.assertEqual(roles["scan"][0]["line_count"], 0)
        self.assertIsNone(result["invoice_no"])
        self.assertEqual(result["order_no"], "ZW/PRO_FORM/21/00724")
        self.assertEqual(result["contract"], "STL-1EL")
        self.assertEqual(result["currency"], "EUR")
        self.assertIn("unit_conflict", result["flags"])
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 10)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 27500)
        self.assertAlmostEqual(sum(lot["amount"] or 0 for lot in goods), 31467.80, places=2)
        self.assertAlmostEqual(sum(lot["net"] or 0 for lot in goods), 5028.58, places=2)
        self.assertAlmostEqual(sum(lot["gross"] or 0 for lot in goods), 5112.68, places=2)
        self.assertAlmostEqual(sum(lot["gross_with_pallet"] or 0 for lot in goods), 5216, places=2)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 1216)
        dish = next(lot for lot in goods if lot["vendor"] == "SP 085-28")
        self.assertEqual(dish["pieces"], 115)
        self.assertEqual(dish["packages"], 58)
        self.assertAlmostEqual(dish["price"], 51.76)
        self.assertAlmostEqual(dish["amount"], 5952.40, places=2)
        self.assertEqual(dish["net"], 1564)
        self.assertEqual(dish["hs"], "9401908009")
        self.assertEqual(dish["origin"], "PL")
        self.assertIn("7", dish["pallet"])
        self.assertIn("Опора", dish["description"])
        self.assertFalse(dish["unit_conflict"])
        covers = [lot for lot in goods if lot["vendor"] == "TO 002-4"]
        self.assertEqual(sorted(lot["pieces"] for lot in covers), [2000, 20000])
        self.assertEqual(sorted(lot["packages"] for lot in covers), [10, 800])
        hinge = next(lot for lot in goods if lot["vendor"] == "CH 031-1")
        self.assertEqual(hinge["pieces"], 3900)
        self.assertEqual(hinge["packages"], 50)
        self.assertTrue(hinge["unit_conflict"])
        self.assertAlmostEqual(hinge["price"], 1.05)
        self.assertAlmostEqual(hinge["net"], 1107.60, places=2)


class Shipment7Test(unittest.TestCase):
    def test_repeated_header_and_two_codes(self):
        result = load("Шаблон 7")
        roles = {doc["role"] for doc in result["documents"]}
        self.assertIn("invoice", roles)
        self.assertIn("packing", roles)
        specs = [doc for doc in result["documents"] if doc["role"] == "specification"]
        self.assertEqual(len(specs), 1)
        self.assertEqual(specs[0]["line_count"], 115)
        self.assertTrue(any(doc["role"] == "unknown" for doc in result["documents"]))
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 115)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 681)
        self.assertAlmostEqual(sum(lot["amount"] or 0 for lot in goods), 11341.61, places=2)
        self.assertEqual(result["currency"], "EUR")
        self.assertEqual(result["invoice_nos"], ["10587886", "10587887"])
        self.assertEqual(result["order_nos"], ["510211", "516361"])
        self.assertEqual(result["contract"], "001/02/12")
        self.assertEqual(result["buyer"], "Ejendals LLC")
        self.assertIn("hs_conflict", result["flags"])
        self.assertEqual(sum(1 for lot in goods if lot["vendor"] == "JALAS 1615-44" and lot["pieces"] == 1), 2)
        insole = next(lot for lot in goods if lot["vendor"] == "JALAS 8102-36")
        self.assertEqual(insole["pieces"], 2)
        self.assertAlmostEqual(insole["price"], 3.45)
        self.assertEqual(insole["origin"], "Vietnam")
        boot = next(lot for lot in goods if "6468" in lot["vendor"])
        self.assertEqual(boot["hs"], "6404199000")
        self.assertEqual(boot["hs_alt"], "6403999300")
        self.assertEqual(boot["unit"], "PAR")
        self.assertTrue(all(lot["vendor"] != "Purchase contract" for lot in goods))


class PalletsAndAddressesTest(unittest.TestCase):
    def test_pallet_row_is_not_places_and_addresses_are_clean(self):
        result = load("Шаблон 19")
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 1)
        lot = goods[0]
        self.assertEqual(lot["packages"], 2150)
        self.assertEqual(lot["pallet_count"], 20)
        self.assertEqual(lot["pallet_weight"], 250)
        self.assertEqual(lot["gross"], 21750)
        self.assertEqual(lot["gross_with_pallet"], 22000)
        self.assertIn("Maxim Gorky", result["buyer_address"])
        self.assertIn("room 26", result["buyer_address"])
        self.assertNotIn("PAKISTAN", result["buyer_address"])
        self.assertIn("PAKISTAN", result["seller_address"])
        self.assertEqual(result["currency"], "USD")

    def test_shipper_code_second_date_and_fees(self):
        result = load("Шаблон 13")
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertTrue(goods[0]["hs_alt_shipper"])
        self.assertEqual(result["invoice_date"], "06.02.2018")
        self.assertNotIn("Container", result["buyer_address"])
        self.assertNotIn("BL No", result["buyer_address"])
        self.assertTrue(result["seller_address"].endswith("CHINA"))
        self.assertEqual(len(result["freights"]), 2)
        self.assertIn("weight_suspect", result["flags"])


class DatesAndPromptModesTest(unittest.TestCase):
    def test_date_with_month_name_and_year_first(self):
        from prepare_transform_code.shipment import _named_date, _spaced_date

        self.assertEqual(_named_date("Mar. 31, 2014"), "31.03.2014")
        self.assertEqual(_named_date("5 September 2020"), "05.09.2020")
        self.assertEqual(_spaced_date("Invoice No: 7\nDate\n2018-2-6"), "06.02.2018")
        self.assertEqual(_spaced_date("Date: Mar. 31, 2014"), "31.03.2014")

    def test_full_rules_stay_available_for_comparison(self):
        import os
        from prepare_transform_code.verdict import FULL_RULES, RULES, build_prompt

        self.assertLess(len(RULES), len(FULL_RULES))
        os.environ["VERDICT_PROMPT"] = "full"
        try:
            self.assertTrue(build_prompt({"lots": []}).startswith(FULL_RULES))
        finally:
            del os.environ["VERDICT_PROMPT"]


class PhotoTileTest(unittest.TestCase):
    def test_left_to_right_then_down_skip_empty_corner(self):
        from prepare_transform_code.photos import _planned_pages

        cells = {}
        for col, text in ((1, "A"), (2, "B"), (3, "C")):
            cells[(1, col)] = {"text": text}
        cells[(2, 1)] = {"text": "left"}
        cells[(2, 2)] = {"text": "mid"}
        cells[(4, 1)] = {"text": "meta"}
        grid = {
            "cols": [1, 2, 3],
            "rows": [1, 2, 3, 4],
            "header_rows": 1,
            "col_spans": [],
            "row_spans": [],
            "cells": cells,
        }
        planned = _planned_pages(grid, {1: 2000, 2: 2000, 3: 2000}, {1: 700, 2: 700, 3: 700, 4: 700})
        self.assertEqual([cols for cols, _rows in planned], [[1], [2], [1]])
        self.assertIn(2, planned[0][1])
        self.assertIn(4, planned[2][1])
        self.assertNotIn(3, planned[2][0])

    def test_sideways_scan_keeps_the_header_on_top(self):
        from PIL import Image

        from prepare_transform_code.photos import _black_ends, _upright

        folder = ROOT / "Шаблон 8"
        invoice = _upright(Image.open(next(folder.glob("*инвойс*"))).convert("RGB"))
        packing = _upright(Image.open(next(folder.glob("*пакинг*"))).convert("RGB"))
        self.assertGreater(invoice.size[1], invoice.size[0])
        self.assertGreater(packing.size[0], packing.size[1])
        for image in (invoice, packing):
            top, bottom = _black_ends(image)
            self.assertGreater(top, bottom)


class Shipment8Test(unittest.TestCase):
    def test_gloves_fixed_columns(self):
        result = load("Шаблон 8")
        roles = {doc["role"] for doc in result["documents"]}
        self.assertIn("specification", roles)
        self.assertIn("image", roles)
        self.assertEqual(result["columns"], list(LOT_FIELDS))
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 13)
        for lot in goods:
            for name in LOT_FIELDS:
                self.assertIn(name, lot)
            self.assertNotIn("PLN", lot)
            self.assertNotIn("USD", lot)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 520000)
        self.assertAlmostEqual(sum(lot["amount"] or 0 for lot in goods), 48465.12, places=2)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 655)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["contract"], "2")
        self.assertIn("ОКТЯБРЬ", result["buyer"])
        self.assertIn("YANTAI", result["seller"])
        first = goods[0]
        self.assertEqual(first["pieces"], 79920)
        self.assertEqual(first["packages"], 111)
        self.assertAlmostEqual(first["price"], 0.103)
        self.assertEqual(first["unit"], "пара")
        self.assertEqual(first["hs"], "6116920000")
        self.assertEqual(first["net"], None)
        self.assertEqual(first["gross"], None)


class PromptTest(unittest.TestCase):
    def test_rules_stay_free_of_shipment_facts(self):
        from prepare_transform_code.verdict import RULES, build_prompt

        for banned in ("Шаблон", "TIG200", "JH21024", "HUAH", "81A"):
            self.assertNotIn(banned, RULES)
        sample = {
            "documents": [{"name": "a.pdf", "role": "invoice", "line_count": 1}],
            "columns": ["description", "pieces"],
            "seller": "A",
            "buyer": "B",
            "lots": [{"description": "sample row", "pieces": 1}],
            "flags": [],
            "currency": "USD",
            "currencies": ["USD"],
            "invoice_no": "",
            "invoice_nos": [],
            "order_no": "",
            "order_nos": [],
            "contract": "",
        }
        text = build_prompt(sample)
        self.assertTrue(text.startswith(RULES))
        self.assertIn("sample row", text)
        self.assertIn("split", text)
        self.assertIn("add", text)
        self.assertEqual(sample["lots"][0]["description"], "sample row")

    def test_stronger_model_only_when_the_draft_is_empty(self):
        from prepare_transform_code.verdict import needs_stronger_model

        empty = {"lots": []}
        thin = [{"fields": {"description": "a", "pieces": 1, "price": 2}}]
        wide = [{"fields": {"description": "a", "vendor": "b", "pieces": 1, "price": 2, "amount": 3}}]
        self.assertTrue(needs_stronger_model(empty, thin))
        self.assertTrue(needs_stronger_model(empty, []))
        self.assertFalse(needs_stronger_model(empty, wide))
        seen = {"lots": [{"description": "leather", "pieces": 10, "amount": 5}]}
        self.assertFalse(needs_stronger_model(seen, thin))


class Shipment9Test(unittest.TestCase):
    def test_casters_from_specification(self):
        result = load("Шаблон 9")
        roles = {doc["role"] for doc in result["documents"]}
        self.assertIn("specification", roles)
        self.assertIn("image", roles)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 6)
        self.assertEqual(result["columns"], list(LOT_FIELDS))
        for lot in goods:
            for name in LOT_FIELDS:
                self.assertIn(name, lot)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 23680)
        self.assertAlmostEqual(sum(lot["amount"] or 0 for lot in goods), 4402.8, places=2)
        self.assertEqual([lot["vendor"] for lot in goods], ["J-34 ic", "J-34 cn", "J-34 cnt", "J-34 d", "J-34 e", "J-34 et"])
        self.assertEqual(goods[0]["pieces"], 20000)
        self.assertAlmostEqual(goods[0]["price"], 0.138)
        self.assertIn("цинк", goods[0]["finish"])
        self.assertEqual(goods[0]["hs"], "")
        self.assertEqual(goods[0]["net"], None)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["contract"], "HUAH-RS-2018")
        self.assertIsNone(result["invoice_no"])
        self.assertIsNone(result["order_no"])
        self.assertIn("HUANAN", result["seller"])
        self.assertIn("Респект", result["buyer"])
        self.assertNotIn("с одной стороны", result["seller"])


class Shipment10Test(unittest.TestCase):
    def test_boots_size_run_and_two_codes(self):
        result = load("Шаблон 10")
        roles = {doc["role"] for doc in result["documents"]}
        self.assertIn("specification", roles)
        self.assertEqual(sum(1 for doc in result["documents"] if doc["role"] == "scan"), 2)
        specs = [doc for doc in result["documents"] if doc["role"] == "specification"]
        self.assertEqual(len(specs), 1)
        self.assertEqual(specs[0]["line_count"], 23)
        self.assertEqual(result["columns"], list(LOT_FIELDS))
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 23)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 3817)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 491)
        self.assertAlmostEqual(sum(lot["amount"] or 0 for lot in goods), 166080.1, places=2)
        self.assertAlmostEqual(sum(lot["net"] or 0 for lot in goods), 4892, places=2)
        self.assertAlmostEqual(sum(lot["gross"] or 0 for lot in goods), 6974, places=2)
        self.assertEqual(result["currency"], "USD")
        self.assertNotIn("RUB", result["currencies"])
        self.assertEqual(result["contract"], "RK/RK/74372/120091/UKM/3208606.1")
        self.assertIsNone(result["invoice_no"])
        self.assertIsNone(result["order_no"])
        self.assertIn("AIRWAIR", result["seller"])
        self.assertIn("Обувь", result["buyer"])
        self.assertNotIn("AIRWAIR", result["buyer"])
        self.assertIn("hs_conflict", result["flags"])
        codes = [lot["model"].split()[0] for lot in goods]
        self.assertEqual(
            codes,
            [
                "11821260", "11822002", "14045001", "10072004", "10072600",
                "11821018", "11838600", "10085001", "11857001", "12270003",
                "13418002", "14286401", "14735450", "14735500", "14069002",
                "14070201", "14072201", "14335030", "14335201", "14340201",
                "14665001", "13512420", "14736450",
            ],
        )
        first = goods[0]
        self.assertEqual(first["pieces"], 144)
        self.assertEqual(first["packages"], 18)
        self.assertAlmostEqual(first["price"], 36.7)
        self.assertEqual(first["unit"], "пар")
        self.assertEqual(first["hs"], "6404199000")
        self.assertEqual(first["hs_alt"], "")
        self.assertIn("taupe", first["finish"])
        self.assertIn("3:11", first["size"])
        self.assertIn("6.5:10", first["size"])
        self.assertNotEqual(first["size"], "11")
        for lot in goods:
            self.assertEqual(lot["unit"], "пар")
            self.assertIn(":", lot["size"])
            total = 0
            for part in lot["size"].split(","):
                total += float(part.split(":")[1])
            self.assertEqual(total, lot["pieces"])
        vegan = next(lot for lot in goods if lot["model"].startswith("14045001"))
        self.assertEqual(vegan["hs"], "6403911300")
        self.assertEqual(vegan["hs_alt"], "6402919000")
        self.assertEqual(sum(1 for lot in goods if lot["hs_alt"]), 1)
        self.assertIn("натуральн", vegan["description"])


class Shipment11Test(unittest.TestCase):
    def test_style_invoice_under_letterhead(self):
        result = load("Шаблон 11")
        roles = {doc["role"]: doc["line_count"] for doc in result["documents"]}
        self.assertEqual(roles.get("invoice"), 14)
        self.assertEqual(roles.get("packing"), 14)
        self.assertEqual(roles.get("specification"), 14)
        self.assertEqual(result["columns"], list(LOT_FIELDS))
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 14)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 2804)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 277)
        self.assertAlmostEqual(sum(lot["amount"] or 0 for lot in goods), 50206.52, places=2)
        self.assertAlmostEqual(sum(lot["net"] or 0 for lot in goods), 2943.2, places=2)
        self.assertAlmostEqual(sum(lot["gross"] or 0 for lot in goods), 3554.2, places=2)
        self.assertAlmostEqual(sum(lot["volume"] or 0 for lot in goods), 27.158904, places=4)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["invoice_no"], "SG1251")
        self.assertEqual(result["invoice_nos"], ["SG1251"])
        self.assertEqual(result["contract"], "02/04/12")
        self.assertIn("SHINEGARDEN", result["seller"])
        self.assertNotIn("bank", result["seller"].lower())
        self.assertIn("XXI Century", result["buyer"])
        self.assertIn("hs_conflict", result["flags"])
        self.assertNotIn("unit_conflict", result["flags"])
        models = [lot["model"] for lot in goods]
        self.assertEqual(
            models,
            [
                "SH2077-4", "6008", "SH2077-4", "SH1151-5", "SH2063-2", "SH2063-4",
                "SH1527-7", "SH1527-9", "SH1527-3", "FW610105-2", "SG610105",
                "SG610105", "SH2012-3A", "SH1434-22",
            ],
        )
        first = goods[0]
        self.assertEqual(first["pieces"], 200)
        self.assertEqual(first["packages"], 20)
        self.assertAlmostEqual(first["price"], 16.43)
        self.assertEqual(first["unit"], "pairs")
        self.assertEqual(first["unit_net"], 11)
        self.assertAlmostEqual(first["volume"], 2.41542, places=4)
        self.assertEqual(first["hs"], "640399")
        self.assertEqual(first["hs_alt"], "6403911600")
        self.assertIn("черный", first["finish"])
        self.assertIn("40:15", first["size"])
        self.assertIn("Ботинки мужские", first["description"])
        self.assertNotIn("MEN'S AND LADY", first["description"])
        for lot in goods:
            self.assertFalse(lot["unit_conflict"])
            total = sum(float(part.split(":")[1]) for part in lot["size"].split(","))
            self.assertEqual(total, lot["pieces"])
        lady = goods[-1]
        self.assertEqual(lady["pieces"], 204)
        self.assertEqual(lady["packages"], 17)
        self.assertIn("36:17", lady["size"])
        self.assertIn("41:17", lady["size"])
        self.assertNotIn("42:", lady["size"])
        self.assertEqual(lady["hs"], "640399")
        self.assertEqual(lady["hs_alt"], "6403911800")
        self.assertIn("женские", lady["description"])
        coffee = goods[8]
        self.assertIn("светло", coffee["finish"])


class Shipment12Test(unittest.TestCase):
    def test_code_column_and_box_qty(self):
        result = load("Шаблон 12")
        roles = {doc["role"]: doc["line_count"] for doc in result["documents"] if doc["line_count"]}
        self.assertEqual(roles.get("invoice"), 17)
        self.assertEqual(roles.get("packing"), 17)
        self.assertEqual(roles.get("specification"), 17)
        self.assertEqual(sum(1 for doc in result["documents"] if doc["role"] == "scan"), 2)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 17)
        self.assertEqual(sum(lot["pieces"] or 0 for lot in goods), 4860)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 1824)
        self.assertAlmostEqual(sum(lot["amount"] or 0 for lot in goods), 42247, places=2)
        self.assertAlmostEqual(sum(lot["net"] or 0 for lot in goods), 10679.0792, places=2)
        self.assertAlmostEqual(sum(lot["gross"] or 0 for lot in goods), 11009.36, places=2)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["invoice_no"], "160305-160306")
        self.assertEqual(result["invoice_nos"], ["160305-160306"])
        self.assertEqual(result["contract"], "0702ST")
        self.assertIn("STAR MUTFAK", result["seller"])
        self.assertIn("RESPECT", result["buyer"])
        self.assertIn("packages_conflict", result["flags"])
        self.assertNotIn("unit_conflict", result["flags"])
        codes = [lot["vendor"] for lot in goods]
        self.assertEqual(
            codes,
            [
                "S-2211", "S-2213", "S-2214", "S-2251", "S-2281", "S-2282", "S-2283",
                "S-2284", "S-2285", "S-2286", "S-2287", "S-2288", "S-2289", "S-2290",
                "S-3013", "S-3014", "S-6611",
            ],
        )
        first = goods[0]
        self.assertEqual(first["pieces"], 850)
        self.assertEqual(first["packages"], 850)
        self.assertEqual(first["gross"], 4233)
        self.assertEqual(first["origin"], "TURKEY")
        self.assertEqual(first["unit"], "Pcs")
        self.assertEqual(first["hs"], "9403901000")
        self.assertIn("бутылочница", first["description"])
        bin_lot = goods[3]
        self.assertEqual(bin_lot["pieces"], 360)
        self.assertEqual(bin_lot["packages"], 30)
        self.assertIn("ведро", bin_lot["description"])
        conflict = next(lot for lot in goods if lot["packages_conflict"])
        self.assertEqual(conflict["vendor"], "S-2214")
        self.assertEqual(conflict["packages"], 180)
        self.assertEqual(conflict["pieces"], 170)
        self.assertEqual(sum(1 for lot in goods if lot["packages_conflict"]), 1)


class Shipment13Test(unittest.TestCase):
    def test_one_set_and_two_charges(self):
        result = load("Шаблон 13")
        roles = {doc["name"]: doc for doc in result["documents"]}
        invoice = next(doc for doc in result["documents"] if doc["role"] == "invoice")
        spec = next(doc for doc in result["documents"] if doc["role"] == "specification")
        self.assertEqual(invoice["line_count"], 0)
        self.assertEqual(spec["line_count"], 3)
        self.assertTrue(any(doc["role"] == "unknown" and doc["line_count"] == 0 for doc in roles.values()))
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 1)
        self.assertEqual(len(result["freights"]), 2)
        lot = goods[0]
        self.assertEqual(lot["vendor"], "4181")
        self.assertEqual(lot["pieces"], 960)
        self.assertAlmostEqual(lot["price"], 28.48)
        self.assertAlmostEqual(lot["amount"], 27340.8)
        self.assertEqual(lot["net"], 4)
        self.assertEqual(lot["unit"], "set / компл")
        self.assertEqual(lot["hs"], "9401908009")
        self.assertEqual(lot["hs_alt"], "9401909090")
        self.assertEqual(lot["origin"], "China")
        self.assertIn("реклайнер", lot["description"])
        self.assertEqual(result["currency"], "USD")
        self.assertNotIn("EUR", result["currencies"])
        self.assertEqual(result["invoice_no"], "C-17LPTZ021-1")
        self.assertEqual(result["contract"], "LPRS-19-04/17")
        self.assertIn("Leggett", result["seller"])
        self.assertIn("RESPECT", result["buyer"])
        self.assertIn("hs_conflict", result["flags"])
        freight_sum = sum(item["amount"] for item in result["freights"])
        self.assertAlmostEqual(freight_sum, 3432.52, places=2)
        self.assertAlmostEqual(lot["amount"] + freight_sum, 30773.32, places=2)


class Shipment14Test(unittest.TestCase):
    def test_stone_flowers_spec(self):
        result = load("Шаблон 14")
        roles = {doc["role"]: doc for doc in result["documents"]}
        self.assertEqual(roles["invoice"]["line_count"] if "invoice" in roles else None, None)
        scans = [doc for doc in result["documents"] if doc["role"] == "scan"]
        spec = next(doc for doc in result["documents"] if doc["role"] == "specification")
        self.assertEqual(len(scans), 2)
        self.assertEqual(spec["line_count"], 48)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 48)
        self.assertEqual(sum(lot["pieces"] for lot in goods), 1670)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 93)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 888.83, places=2)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["contract"], "0103-XDR-AC")
        self.assertIsNone(result["invoice_no"])
        self.assertIn("XUZHOU", result["seller"])
        self.assertIn("RESPECT", result["buyer"])
        self.assertNotEqual(result["buyer"], "Покупатель")
        self.assertNotIn("fails", result["seller"].lower())
        first = goods[0]
        self.assertEqual(first["vendor"], "F02PECAH")
        self.assertEqual(first["unit"], "pcs/шт")
        self.assertEqual(first["hs"], "6802999000")
        self.assertIn("China", first["origin"])
        self.assertIn("GUANGDONG", first["producer"])
        self.assertIn("цветы", first["description"])
        shared = [lot["vendor"] for lot in goods if lot["packages"] is None]
        self.assertEqual(shared, ["F111-6WE（FLOWER）", "F431-1FLOWER"])
        self.assertTrue(all(lot["packages"] != 0 for lot in goods))


class Shipment15Test(unittest.TestCase):
    def test_handles_spec_and_two_pallets(self):
        result = load("Шаблон 15")
        scans = [doc for doc in result["documents"] if doc["role"] == "scan"]
        spec = next(doc for doc in result["documents"] if doc["role"] == "specification")
        self.assertEqual(len(scans), 2)
        self.assertEqual(spec["line_count"], 53)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 53)
        self.assertEqual(sum(lot["pieces"] for lot in goods), 80905)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 177)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 25630.70, places=2)
        self.assertEqual(sum(lot["net"] for lot in goods), 1998)
        self.assertEqual(sum(lot["gross"] for lot in goods), 2201)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["contract"], "21V-210622-04IMP")
        self.assertIsNone(result["invoice_no"])
        self.assertIn("OZKARDESLER", result["seller"])
        self.assertIn("TD 21 VEK", result["buyer"])
        self.assertIn("hs_conflict", result["flags"])
        first = goods[0]
        self.assertEqual(first["vendor"], "5001-03")
        self.assertEqual(first["pieces"], 1500)
        self.assertEqual(first["packages"], 2)
        self.assertEqual(first["gross"], 34)
        self.assertEqual(first["hs"], "8302420000")
        self.assertEqual(first["hs_alt"], "8302.42.00.00.19")
        self.assertEqual(first["origin"], "TURKEY / ТУРЦИЯ")
        self.assertEqual(first["brand"], "OZKM")
        self.assertEqual(first["pallet"], "part")
        self.assertIn("Ручка", first["description"])
        plastic = next(lot for lot in goods if lot["vendor"] == "5115-03")
        self.assertEqual(plastic["hs"], "3926300000")
        self.assertEqual(plastic["unit"], "Pcs / шт")
        hanger = next(lot for lot in goods if lot["vendor"] == "1001-013")
        self.assertEqual(hanger["hs_alt"], "3924.90.00.00.19")
        self.assertEqual(hanger["pallet"], "1")
        shared = next(lot for lot in goods if lot["vendor"] == "5132-013")
        self.assertIsNone(shared["packages"])
        self.assertEqual(sum(1 for lot in goods if lot["pallet"] == "1"), 2)
        self.assertEqual(sum(1 for lot in goods if lot["pallet"] == "part"), 51)


class Shipment16Test(unittest.TestCase):
    def test_invoice_spec_and_packing(self):
        result = load("Шаблон 16")
        roles = {doc["role"]: doc["line_count"] for doc in result["documents"]}
        self.assertEqual(roles.get("invoice"), 34)
        self.assertEqual(roles.get("packing"), 34)
        self.assertEqual(roles.get("specification"), 34)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 34)
        self.assertEqual(result["flags"], [])
        self.assertEqual(sum(lot["pieces"] for lot in goods), 2333280)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 2522)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 475638.47, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 23447.96, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 24416.26, places=2)
        self.assertAlmostEqual(sum(lot["volume"] or 0 for lot in goods), 65.74, places=2)
        self.assertEqual(result["currency"], "CNY")
        self.assertEqual(result["contract"], "SHUNTEX/21V-2025")
        self.assertEqual(result["invoice_no"], "RU30006")
        self.assertEqual(result["invoice_nos"], ["RU30006"])
        self.assertIn("GUANGDONG SHUNTEX", result["seller"])
        self.assertIn("TD 21 VEK", result["buyer"])
        self.assertNotIn("DATE", result["buyer"])
        first = goods[0]
        self.assertEqual(first["vendor"], "0610609")
        self.assertEqual(first["model"], "25109.150L")
        self.assertEqual(first["pieces"], 240)
        self.assertEqual(first["packages"], 240)
        self.assertEqual(first["net"], 804)
        self.assertEqual(first["gross"], 876)
        self.assertAlmostEqual(first["volume"], 6.33, places=2)
        self.assertAlmostEqual(first["price"], 53.946, places=3)
        self.assertEqual(first["hs"], "9403990001")
        self.assertEqual(first["unit_net"], 3.3)
        self.assertEqual(first["size"], "")
        self.assertIn("комплект", first["unit"])
        self.assertIn("КИТАЙ", first["origin"])
        self.assertIn("Бутылочница", first["description"])
        bronze = next(lot for lot in goods if lot["vendor"] == "0904173")
        self.assertEqual(bronze["model"], "11017")
        self.assertEqual(bronze["pieces"], 2000)
        self.assertEqual(bronze["packages"], 20)
        self.assertEqual(bronze["hs"], "9403990001")
        self.assertEqual(bronze["unit_net"], 19.8)
        self.assertIn("бронза", bronze["description"])
        caster = next(lot for lot in goods if lot["vendor"] == "0903126")
        self.assertIsNone(caster["unit_net"])
        self.assertEqual(caster["hs"], "8302200000")
        self.assertTrue(all(lot["hs"] and lot["unit"] and lot["model"] for lot in goods))
        self.assertTrue(all(lot["size"] == "" for lot in goods))


class Shipment17Test(unittest.TestCase):
    def test_two_invoices_one_shipment(self):
        result = load("Шаблон 17")
        roles = {}
        for doc in result["documents"]:
            roles[doc["role"]] = roles.get(doc["role"], 0) + doc["line_count"]
        self.assertEqual(roles.get("invoice"), 4)
        self.assertEqual(roles.get("packing"), 4)
        self.assertEqual(roles.get("proforma"), 4)
        unknown = [doc for doc in result["documents"] if doc["role"] == "unknown" and doc["line_count"]]
        self.assertEqual(len(unknown), 1)
        self.assertEqual(unknown[0]["line_count"], 4)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 4)
        self.assertEqual(result["flags"], [])
        self.assertEqual(sum(lot["pieces"] for lot in goods), 1200)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 200)
        self.assertEqual(sum(lot["amount"] for lot in goods), 207000)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 2520, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 2760, places=2)
        self.assertEqual(result["currency"], "CNY")
        self.assertNotIn("USD", result["currencies"])
        self.assertNotIn("EUR", result["currencies"])
        self.assertEqual(result["invoice_no"], "MD260224")
        self.assertEqual(result["invoice_nos"], ["MD260224"])
        self.assertEqual(result["proforma_no"], "MD250915")
        self.assertEqual(result["contract"], "27/08")
        self.assertIn("GAOMI", result["seller"])
        self.assertIn("Vitoria", result["buyer"])
        first = goods[0]
        self.assertEqual(first["vendor"], "135532")
        self.assertEqual(first["model"], "REGENT")
        self.assertEqual(first["pieces"], 300)
        self.assertEqual(first["packages"], 50)
        self.assertEqual(first["price"], 150)
        self.assertEqual(first["amount"], 45000)
        self.assertEqual(first["net"], 607.5)
        self.assertEqual(first["gross"], 667.5)
        self.assertEqual(first["unit_net"], 1.9)
        self.assertEqual(first["hs"], "6403400000")
        self.assertEqual(first["unit"], "pairs")
        self.assertIn("38:10", first["size"])
        self.assertIn("42:65", first["size"])
        self.assertIn("чёрный", first["finish"])
        self.assertNotIn("BLK", first["finish"])
        self.assertIn("Обувь", first["description"])
        self.assertIn("FULL GRAIN", first["description"])
        self.assertTrue(all(lot["packages"] == 50 and lot["pieces"] == 300 for lot in goods))


class Shipment18Test(unittest.TestCase):
    def test_code_column_is_customs_code(self):
        result = load("Шаблон 18")
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 113)
        self.assertEqual(result["flags"], [])
        self.assertEqual(sum(lot["pieces"] for lot in goods), 43558)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 33364.31, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 9669.64, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 11376.05, places=2)
        self.assertEqual(sum(lot["packages"] or 0 for lot in goods), 740)
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["invoice_no"], "YT-20251024")
        self.assertEqual(result["contract"], "YT-02/04/2025")
        self.assertEqual(result["seller"], "YM TRANS GROUP CO., LIMITED")
        self.assertIn("Master", result["buyer"])
        self.assertNotIn("RECIPIENT", result["buyer"])
        self.assertNotIn("DELIVERY", result["seller"])
        first = goods[0]
        self.assertEqual(first["hs"], "8544429009")
        self.assertEqual(first["vendor"], "")
        self.assertEqual(first["pieces"], 30000)
        self.assertEqual(first["origin"], "CN")
        self.assertTrue(all(lot["hs"] and not lot["vendor"] for lot in goods))
        roles = {}
        for doc in result["documents"]:
            roles.setdefault(doc["role"], []).append(doc)
        self.assertEqual(roles["invoice"][0]["line_count"], 113)
        self.assertEqual(roles["packing"][0]["line_count"], 113)
        unknown = [doc for doc in result["documents"] if doc["role"] == "unknown" and doc["line_count"] == 113]
        self.assertGreaterEqual(len(unknown), 1)


class Shipment19Test(unittest.TestCase):
    def test_one_file_holds_invoice_and_packing(self):
        result = load("Шаблон 19")
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 1)
        self.assertEqual(result["flags"], [])
        self.assertEqual(result["currency"], "USD")
        self.assertEqual(result["invoice_no"], "INV-02/5PS-2024")
        self.assertEqual(result["contract"], "5PI-10/12")
        self.assertEqual(result["seller"], "5PS INTERNATIONAL")
        self.assertEqual(result["buyer"], "FRUITVILL LLC")
        self.assertNotIn("Master", result["buyer"])
        self.assertNotEqual(result["seller"], result["buyer"])
        lot = goods[0]
        self.assertIsNone(lot["pieces"])
        self.assertEqual(lot["packages"], 2150)
        self.assertEqual(lot["price"], 0.85)
        self.assertEqual(lot["amount"], 17850)
        self.assertEqual(lot["net"], 21000)
        self.assertEqual(lot["gross"], 21750)
        self.assertEqual(lot["hs"], "0805210000")
        self.assertEqual(lot["unit"], "kg")
        self.assertEqual(lot["origin"], "PAKISTAN")
        self.assertIn("FRESH MANDARINES", lot["description"])
        self.assertIn("Мандарины", lot["description"])
        combined = next(doc for doc in result["documents"] if "INV" in doc["name"])
        self.assertEqual(combined["line_count"], 1)
        self.assertIn("invoice", combined["roles"])
        self.assertIn("packing", combined["roles"])
        spec = next(doc for doc in result["documents"] if doc["role"] == "specification")
        self.assertEqual(spec["line_count"], 1)


class Shipment20Test(unittest.TestCase):
    def test_one_machine_across_packings(self):
        result = load("Шаблон 20")
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 1)
        self.assertEqual(result["flags"], [])
        self.assertEqual(result["currency"], "RUB")
        self.assertIn("RUB", result["currencies"])
        self.assertEqual(result["invoice_no"], "ANTALYA - 119000 REG")
        self.assertEqual(result["contract"], "22/11")
        self.assertIn("WELCOME YAPI", result["seller"])
        self.assertIn("Techexpo", result["buyer"])
        self.assertNotIn("Master", result["buyer"])
        self.assertNotIn("CONSINGNEE", result["seller"])
        lot = goods[0]
        self.assertEqual(lot["pieces"], 1)
        self.assertEqual(lot["packages"], 102)
        self.assertEqual(lot["amount"], 900000000)
        self.assertIsNone(lot["price"])
        self.assertEqual(lot["gross"], 185480)
        self.assertIsNone(lot["net"])
        self.assertEqual(lot["hs"], "8457109008")
        self.assertEqual(lot["model"], "B6S / 35240")
        self.assertEqual(lot["unit"], "pcs")
        self.assertEqual(lot["origin"], "UAE")
        self.assertIn("CNC Router", lot["description"])
        self.assertIn("Часть", lot["description"])
        self.assertNotIn("Package 1", lot["description"])
        self.assertEqual(sum(1 for doc in result["documents"] if doc["role"] == "packing"), 11)
        self.assertTrue(all(doc["line_count"] == 1 for doc in result["documents"] if doc["line_count"]))


class ForeignSheetTest(unittest.TestCase):
    def test_other_articles_stay_out(self):
        from prepare_transform_code.lines import Line
        from prepare_transform_code.shipment import _companions

        invoice = {"role": "invoice", "lines": [Line(vendor="135532", pieces=300)]}
        stranger = {"role": "invoice", "lines": [Line(vendor="47208", pieces=50)]}
        same = {"role": "proforma", "lines": [Line(vendor="135532", pieces=300, description="lining")]}
        spec_lines = []
        foreign = _companions([invoice, stranger, same], invoice, None, invoice["lines"], spec_lines)
        self.assertTrue(foreign)
        self.assertEqual(len(spec_lines), 1)
        self.assertEqual(spec_lines[0].description, "lining")


class FabricRollTest(unittest.TestCase):
    def test_headers(self):
        from prepare_transform_code.fields import is_row_index
        from prepare_transform_code.shipment import _contract

        self.assertEqual(column_of("Quality Code"), "vendor")
        self.assertEqual(column_of("Design Name"), "description")
        self.assertEqual(column_of("Net Mt."), "qty")
        self.assertEqual(column_of("Net Kg."), "net")
        self.assertEqual(column_of("Gross Kg."), "gross")
        self.assertEqual(column_of("Total USD"), "amount")
        self.assertEqual(column_of("Item"), "vendor")
        self.assertEqual(column_of("Cust PO"), "order_ref")
        self.assertIsNone(column_of("Design Code"))
        self.assertIsNone(column_of("Mill Code"))
        self.assertIsNone(column_of("Finishing Batch No"))
        self.assertIsNone(column_of("Report Size"))
        self.assertTrue(is_row_index("Item", "1"))
        self.assertFalse(is_row_index("Item", "FAB-1"))
        self.assertEqual(_contract("Contract No：AYD-1EXT от 23/05/18 FCA ISTAMBUL"), "AYD-1EXT")

    def test_one_invoice_line_splits_into_rolls(self):
        import tempfile

        from openpyxl import Workbook

        folder = Path(tempfile.mkdtemp())
        invoice = Workbook()
        sheet = invoice.active
        sheet.title = "Page1"
        sheet["A1"] = "COMMERCIAL INVOICE"
        sheet["A3"] = "Quantity\n(METERS)"
        sheet["B3"] = "Item"
        sheet["C3"] = "Description"
        sheet["D3"] = "H.S.Code"
        sheet["E3"] = "Unit Price\nUSD / METERS"
        sheet["F3"] = "Total USD"
        sheet["A4"] = 100
        sheet["B4"] = "FAB-1"
        sheet["C4"] = "FABRIC"
        sheet["D4"] = "5407.72.00.90.11"
        sheet["E4"] = 2
        sheet["F4"] = 200
        sheet["A5"] = "* The Goods are of Turkish Origin"
        sheet["E5"] = "Shipping:"
        sheet["F5"] = 0
        sheet["A6"] = "TOTAL:100"
        sheet["C6"] = "Roll:2"
        sheet["A8"] = "Contract No：AYD-1EXT от 23/05/18 FCA ISTAMBUL"
        sheet["A9"] = "Manufactured: ACME TEXTILE"
        invoice.save(folder / "invoice.xlsx")
        packing = Workbook()
        sheet = packing.active
        sheet.title = "Page1"
        sheet["A1"] = "PACKING LIST"
        headers = ["Design Name", "Quality Code", "Net Mt.", "Net Kg.", "Gross Kg.", "Item No", "Cust PO"]
        for col, name in enumerate(headers, start=1):
            sheet.cell(3, col, name)
        sheet["A4"] = "FAB"
        sheet["B4"] = "FAB-1"
        sheet["C4"] = 40
        sheet["D4"] = 1
        sheet["E4"] = 1.2
        sheet["F4"] = "RED"
        sheet["G4"] = "99"
        sheet["A5"] = "FAB"
        sheet["B5"] = "FAB-1"
        sheet["C5"] = 60
        sheet["D5"] = 2
        sheet["E5"] = 2.2
        sheet["F5"] = "BLUE"
        sheet["G5"] = "99"
        packing.save(folder / "packing.xlsx")
        result = analyze(folder)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 2)
        self.assertEqual([lot["pieces"] for lot in goods], [40, 60])
        self.assertEqual([lot["model"] for lot in goods], ["RED", "BLUE"])
        self.assertEqual(goods[0]["vendor"], "FAB-1")
        self.assertEqual(goods[0]["price"], 2)
        self.assertEqual(goods[0]["amount"], 80)
        self.assertEqual(goods[1]["amount"], 120)
        self.assertEqual(goods[0]["net"], 1)
        self.assertEqual(goods[0]["gross"], 1.2)
        self.assertEqual(goods[0]["packages"], 1)
        self.assertEqual(goods[0]["package_type"], "roll")
        self.assertEqual(goods[0]["unit"], "meters")
        self.assertEqual(goods[0]["order_ref"], "99")
        self.assertIn("FABRIC", goods[0]["description"])
        self.assertEqual(goods[0]["origin"], "Turkish")
        self.assertIn("ACME", goods[0]["producer"])
        self.assertEqual(result["contract"], "AYD-1EXT")
        self.assertEqual(sum(lot["pieces"] for lot in goods), 100)
        self.assertEqual(sum(lot["amount"] for lot in goods), 200)
        for child in folder.iterdir():
            child.unlink()
        folder.rmdir()


class LiveFabricTest(unittest.TestCase):
    """Боевые поставки тканей. В папке лежат ещё чужие файлы, в тест входят только три входа."""

    def test_exporter_is_seller(self):
        parties = party_after("EXPORTER: ACME TEXTILE\nIMPORTER: BUYCO LLC")
        self.assertIn("ACME", parties["seller"])
        self.assertIn("BUYCO", parties["buyer"])

    def test_aydin_rolls_from_real_files(self):
        result = _only(
            Path("documents/6_pravka/!Aydin"),
            (
                "AYDIN INVOICE 17255.xlsx",
                "AYDIN PL 17255.xlsx",
                "Копия Specification  Айдын 17255 N.pdf",
            ),
        )
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 16)
        self.assertAlmostEqual(sum(lot["pieces"] for lot in goods), 589.94, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 4802.1116, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 405.2, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 415.92, places=2)
        self.assertEqual(sum(lot["packages"] for lot in goods), 16)
        self.assertTrue(all("EDERAS" in lot["vendor"] for lot in goods))
        self.assertEqual(result["invoice_no"], "AYK2026000000012")
        self.assertEqual(result["currency"], "USD")

    def test_pehlivan_families_split_into_rolls(self):
        result = _only(
            Path("documents/6_pravka/!Pehlivan"),
            (
                "ИНВОЙС ПЕХЛИВАН 17255.pdf",
                "ПАКИНГ ПЕХЛИВАН 17255.xlsx",
                "Specification Пехливан 17255.pdf",
            ),
        )
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 53)
        self.assertAlmostEqual(sum(lot["pieces"] for lot in goods), 1473.8, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 14932.31, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 1170.8, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 1181.4, places=2)
        self.assertEqual(sum(lot["packages"] for lot in goods), 53)
        self.assertEqual(sum(1 for lot in goods if lot["vendor"] == "LIVERPOOL 600"), 11)
        self.assertEqual(sum(1 for lot in goods if not lot["hs"]), 11)
        self.assertTrue(all(lot["origin"] == "TURKISH" for lot in goods))
        self.assertTrue(all(lot["unit"] == "meters" for lot in goods))
        self.assertIn("PEHLIVAN", result["seller"].upper())
        self.assertIn("KONFEKSIYON", result["seller"].upper())
        self.assertIn("EXTERIO", result["buyer"].upper())
        self.assertEqual(result["invoice_no"], "PME2026000000002")
        self.assertEqual(result["contract"], "PEH-1EXT")
        self.assertIn("hs_conflict", result["flags"])

    def test_weavers_twenty_five_rolls(self):
        result = _only(
            Path("documents/5_pravka/!Weavers"),
            (
                "WEAVERS INVOICE.pdf",
                "WEAVERS PLIST.pdf",
                "Specification для 17768 WEAVERS.pdf",
                "Specification для 17768 WEAVERS.xlsx",
            ),
        )
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 25)
        self.assertAlmostEqual(sum(lot["pieces"] for lot in goods), 6344.30, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 28174.38, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 2301.33, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 2398.62, places=2)
        self.assertEqual(sum(lot["packages"] for lot in goods), 141)
        self.assertAlmostEqual(sum(lot["area"] or 0 for lot in goods), 8882.02, places=2)
        self.assertTrue(all(lot["hs"] for lot in goods))
        self.assertIn("WEAVERS", result["seller"].upper())
        self.assertIn("REGIONTEKSTIL", result["buyer"].upper())
        self.assertEqual(result["invoice_no"], "WI12026000000516")
        self.assertEqual(result["contract"], "1-BOY-SM")
        self.assertIn("hs_conflict", result["flags"])


class WeaversLongDesignTest(unittest.TestCase):
    def test_meters_stay_meters_and_ship_to_is_the_address(self):
        folder = next(
            path
            for path in (Path(__file__).resolve().parents[2] / "documents" / "я_тестирую").iterdir()
            if path.name.startswith("20_")
        )
        result = analyze(folder / "вход")
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        by_vendor = {lot["vendor"]: lot for lot in goods}
        self.assertEqual(len(goods), 4)
        self.assertEqual(by_vendor["DYER 789"]["pieces"], 206)
        self.assertEqual(by_vendor["DYER 965"]["pieces"], 212)
        self.assertEqual(by_vendor["LARDASO 100"]["pieces"], 221)
        self.assertEqual(by_vendor["LARDASO 490"]["pieces"], 197)
        self.assertAlmostEqual(sum(lot["pieces"] for lot in goods), 836)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 3218.60, places=2)
        self.assertEqual(sum(lot["packages"] for lot in goods), 16)
        self.assertIn("REGIONTEKSTIL", result["buyer"].upper())
        self.assertIn("Krasnogorsk", result["buyer_address"])
        self.assertIn("room 1", result["buyer_address"])
        self.assertIn("RUSSIA", result["buyer_address"])
        self.assertTrue(result["seller"].startswith("WEAVERS"))
        self.assertIn("EXW", result["delivery"])
        self.assertIn("BOY-SM", result["contract"])


class ArtNoSheetTest(unittest.TestCase):
    def test_article_column_and_buyer_to(self):
        folder = next(
            path
            for path in (Path(__file__).resolve().parents[2] / "documents" / "я_тестирую").iterdir()
            if path.name.startswith("12_")
        )
        result = _only(folder / "вход", ("инвойс_и_пакинг.xlsx",))
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        same = [lot for lot in goods if lot["vendor"] == "KD020"]
        self.assertEqual(len(same), 1)
        first = same[0]
        self.assertEqual(first["pieces"], 5040)
        self.assertEqual(first["packages"], 42)
        self.assertEqual(first["price"], 3.35)
        self.assertEqual(first["amount"], 16884)
        self.assertEqual(first["gross"], 935)
        self.assertEqual(first["net"], 918)
        self.assertAlmostEqual(first["volume"], 0.59)
        self.assertGreater(len(goods), 10)
        self.assertIn("ELEMENT", result["buyer"].upper())
        self.assertEqual(result["invoice_no"], "26BEET-002")
        self.assertEqual(result["currency"], "CNY")


class HmkInvoiceTest(unittest.TestCase):
    def test_broken_header_still_reads_the_line(self):
        folder = Path(__file__).resolve().parents[2] / "documents" / "4_pravka" / "для тест" / "bestway" / "_input"
        result = analyze(folder)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 13)
        self.assertEqual(result["flags"], [])
        self.assertEqual(result["currency"], "USD")
        self.assertIn("CNY", result["currencies"])
        self.assertEqual(result["invoice_no"], "NH-331004")
        self.assertEqual(result["invoice_date"], "15.03.2026")
        self.assertEqual(result["contract_date"], "01.10.2025")
        self.assertIn("01/10", result["contract"])
        self.assertEqual(result["contract"][2], "\u0421")
        self.assertEqual(result["container"], "TSRU8000815")
        self.assertIn("FCA Shanghai", result["delivery"])
        self.assertEqual(result["seller"], "HMK TRADING COMPANY LIMITED")
        self.assertEqual(result["buyer"], "LLC NECARGO")
        self.assertIn("RM1607", result["seller_address"])
        self.assertIn("Lobacheva", result["buyer_address"])
        self.assertEqual(result["director"], "Mr. Wang Jiandong")
        self.assertEqual(sum(lot["pieces"] for lot in goods), 8896)
        self.assertEqual(sum(lot["packages"] for lot in goods), 1487)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 15561.30, places=2)
        self.assertAlmostEqual(sum(lot["net_primary"] for lot in goods), 15561.30, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 17288.94, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 35765.25, places=2)
        first = goods[0]
        self.assertEqual(first["hs"], "9506620000")
        self.assertEqual(first["vendor"], "31021")
        self.assertEqual(first["model"], "31021")
        self.assertEqual(first["packages"], 10)
        self.assertEqual(first["unit"], "ШТ")
        self.assertEqual(first["brand"], "Bestway")
        self.assertEqual(first["producer"], "Bestway (Nantong) Recreation Corp.")
        self.assertEqual(first["pieces"], 360)
        self.assertAlmostEqual(first["price"], 0.1304)
        self.assertAlmostEqual(first["amount"], 46.94)


class FabricFamilyTest(unittest.TestCase):
    def test_design_rows_keep_the_family_price(self):
        folder = Path(__file__).resolve().parents[2] / "documents" / "4_pravka" / "для тест" / "18233" / "_input"
        result = analyze(folder)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        self.assertEqual(len(goods), 13)
        self.assertEqual(result["flags"], [])
        self.assertEqual(result["currency"], "CNY")
        self.assertEqual(result["invoice_no"], "ZFRMB26148-626-1")
        self.assertEqual(result["invoice_date"], "Aug.19,2026")
        self.assertEqual(result["contract"], "SM-LU2")
        self.assertEqual(result["contract_date"], "23/11/2018")
        self.assertEqual(result["delivery"], "EX-WORK HANGZHOU")
        self.assertIn("ZHONGFANG", result["seller"])
        self.assertIn("REGIONTEKSTIL", result["buyer"])
        self.assertIn("Krasnogorsk", result["buyer_address"])
        noble = next(lot for lot in goods if lot["model"] == "Noble 110")
        self.assertEqual(noble["packages"], 15)
        self.assertAlmostEqual(noble["pieces"], 611.5)
        self.assertAlmostEqual(noble["area"], 868.33)
        self.assertAlmostEqual(noble["width"], 1.42)
        self.assertEqual(noble["hs"], "5407610000")
        self.assertAlmostEqual(noble["price"], 43.8)
        self.assertAlmostEqual(noble["amount"], 26783.70)
        self.assertAlmostEqual(noble["net"], 580.93)
        self.assertAlmostEqual(noble["gross"], 597.27)
        self.assertEqual(noble["gsm"], 950)
        self.assertEqual(noble["unit"], "meters")
        self.assertEqual(noble["description"], "SOFA FABRIC Noble")
        self.assertAlmostEqual(sum(lot["pieces"] for lot in goods), 11671.6, places=2)
        self.assertEqual(sum(lot["packages"] for lot in goods), 263)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 6026.26, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 188362.5, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 6299.99, places=2)


class FamilyPriceBandTest(unittest.TestCase):
    def test_color_rows_take_the_family_price_and_the_spec_weight(self):
        from prepare_transform_code.join import build_lots
        from prepare_transform_code.lines import Line

        invoice = [
            Line(model="Velvet LUX 03", pieces=201.4, packages=4, area=285.988, width=1.42, hs="5801320000", unit="meters"),
            Line(model="Velvet LUX 32", pieces=200.4, packages=4, area=284.568, width=1.42, hs="5801320000", unit="meters"),
            Line(model="Velvet LUX 78", pieces=205.3, packages=4, area=291.526, width=1.42, hs="5801320000", unit="meters"),
            Line(model="SOFA FABRIC Velvet LUX", pieces=607.1, packages=12, price=21.67, amount=13155.86, area=862.082, width=1.42, hs="5801320000", unit="meters"),
            Line(model="Lazy Silver", pieces=517.8, packages=11, area=735.276, width=1.42, hs="5801330000", unit="meters"),
            Line(model="SOFA FABRIC Lazy", pieces=517.8, packages=11, price=15.18, amount=7860.2, area=735.276, width=1.42, hs="5801330000", unit="meters"),
            Line(model="Marseille Linen", pieces=2044.3, packages=40, area=2902.906, width=1.42, hs="5801360000", unit="meters"),
            Line(model="SOFA FABRIC Marseille", pieces=2044.3, packages=40, price=20.37, amount=41642.39, area=2902.906, width=1.42, hs="5801360000", unit="meters"),
        ]
        packing = [
            Line(model="SOFA FABRIC Velvet LUX", pieces=607.1, packages=12, net=279.27, gross=291.0),
            Line(model="SOFA FABRIC Lazy", pieces=517.8, packages=11, net=258.9, gross=270.0),
            Line(model="SOFA FABRIC Marseille", pieces=2044.3, packages=40, net=1226.58, gross=1271.0),
        ]
        spec = [
            Line(vendor="Velvet LUX 03", description="Upholstery fabric, velvet", pieces=201.4, packages=4, net=92.64, gross=96.55, hs="5801320000"),
            Line(vendor="Velvet LUX 32", pieces=200.4, packages=4, net=92.19, gross=96.1, hs="5801320000"),
            Line(vendor="Velvet LUX 78", pieces=205.3, packages=4, net=94.44, gross=98.35, hs="5801320000"),
            Line(vendor="Lazy Silver", pieces=517.8, packages=11, net=258.9, gross=270.0, hs="5801330000"),
            Line(vendor="Marseille Linen", pieces=2044.3, packages=40, net=1226.58, gross=1271.0, hs="5801360000"),
        ]
        lots, _freights = build_lots(invoice, packing, spec)
        self.assertEqual([lot["model"] for lot in lots], ["Velvet LUX 03", "Velvet LUX 32", "Velvet LUX 78", "Lazy Silver", "Marseille Linen"])
        velvet = lots[:3]
        self.assertEqual([lot["price"] for lot in velvet], [21.67, 21.67, 21.67])
        self.assertAlmostEqual(sum(lot["amount"] for lot in velvet), 13155.86, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in velvet), 279.27, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in velvet), 291.0, places=2)
        self.assertEqual(lots[0]["description"], "SOFA FABRIC Velvet LUX // Upholstery fabric, velvet")
        self.assertAlmostEqual(lots[3]["price"], 15.18)
        self.assertAlmostEqual(lots[3]["net"], 258.9)
        self.assertAlmostEqual(lots[4]["amount"], 41642.39)
        self.assertAlmostEqual(lots[4]["gross"], 1271.0)


class UpholsteryRollsTest(unittest.TestCase):
    def test_meters_stay_meters_and_both_codes_stay(self):
        folder = Path(__file__).resolve().parents[2] / "documents" / "5_pravka" / "!KDF" / "_input"
        result = analyze(folder)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        roles = {doc["role"]: doc["line_count"] for doc in result["documents"]}
        self.assertEqual(roles.get("invoice"), 6)
        self.assertEqual(roles.get("packing"), 6)
        self.assertEqual(len(goods), 6)
        self.assertEqual(result["invoice_no"], "TEK2026000000126")
        self.assertEqual(result["invoice_date"], "29/04/2026")
        self.assertEqual(result["contract"], "SM-KETS1")
        self.assertEqual(result["contract_date"], "24.01.2019")
        self.assertEqual(result["currency"], "USD")
        self.assertIn("Tekstop", result["seller"])
        self.assertNotIn("FOR THE GOODS", result["seller"])
        self.assertIn("Krasnogorsk", result["buyer_address"])
        self.assertNotIn("Ikitelli", result["buyer_address"])
        self.assertIn("FOB", result["delivery"])
        self.assertIn("FCA", result["delivery"])
        self.assertIn("hs_conflict", result["flags"])
        self.assertIn("delivery_conflict", result["flags"])
        self.assertIn("manufacturer_conflict", result["flags"])
        arne = goods[0]
        self.assertAlmostEqual(arne["pieces"], 98.8)
        self.assertEqual(arne["packages"], 2)
        self.assertAlmostEqual(arne["net"], 82.3)
        self.assertAlmostEqual(arne["gross"], 83.8)
        self.assertAlmostEqual(arne["price"], 8.36)
        self.assertAlmostEqual(arne["amount"], 825.97)
        self.assertEqual(arne["unit"], "meters")
        self.assertEqual(arne["hs"], "54.07.73.00.90.11")
        self.assertEqual(arne["hs_alt"], "5407730000")
        self.assertEqual(arne["vendor"], "ARNE (CATRIN-E2)")
        self.assertIn("Türkiye", arne["origin"])
        self.assertEqual(sum(lot["packages"] for lot in goods), 87)
        self.assertAlmostEqual(sum(lot["pieces"] for lot in goods), 3140, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 2638.8, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 2703.9, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 29655.84, places=2)


class FlockserLeatherTest(unittest.TestCase):
    def test_four_articles_and_comma_area(self):
        folder = Path(__file__).resolve().parents[2] / "documents" / "5_pravka" / "!Flockser"
        result = analyze(folder)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        roles = {}
        for doc in result["documents"]:
            roles.setdefault(doc["role"], []).append(doc)
        self.assertEqual(len(goods), 4)
        self.assertEqual(result["flags"], ["hs_conflict"])
        self.assertEqual(result["currency"], "USD")
        self.assertIn("EUR", result["currencies"])
        self.assertIsNone(result["invoice_no"])
        self.assertEqual(result["invoice_date"], "02.04.2026")
        self.assertEqual(result["contract"], "TUR-2-FL")
        self.assertEqual(result["contract_date"], "12.12.2018")
        self.assertIn("FLOKSER", result["seller"])
        self.assertIn("REGIONTEKSTIL", result["buyer"])
        self.assertIn("Krasnogorsk", result["buyer_address"])
        self.assertIn("EXW", result["delivery"])
        self.assertTrue(roles.get("scan"))
        self.assertTrue(all(doc["line_count"] == 0 for doc in roles["scan"]))
        copied = next(doc for doc in result["documents"] if doc["name"].startswith("Копия"))
        self.assertEqual(copied["role"], "specification")
        self.assertNotIn("invoice", copied["roles"])
        first = goods[0]
        self.assertEqual(first["vendor"], "NERGIS 001")
        self.assertAlmostEqual(first["pieces"], 516.3)
        self.assertEqual(first["packages"], 16)
        self.assertAlmostEqual(first["area"], 722.82)
        self.assertAlmostEqual(first["width"], 1.4)
        self.assertAlmostEqual(first["price"], 6.35)
        self.assertAlmostEqual(first["amount"], 3278.51)
        self.assertAlmostEqual(first["net"], 506.9)
        self.assertAlmostEqual(first["gross"], 515.7)
        self.assertEqual(first["hs"], "590310901000")
        self.assertEqual(first["hs_alt"], "3921120000")
        self.assertEqual(first["unit"], "meters")
        self.assertEqual(
            [lot["vendor"] for lot in goods],
            ["NERGIS 001", "NERGIS 305", "NERGIS 318", "NERGIS 901"],
        )
        self.assertEqual(sum(lot["packages"] for lot in goods), 97)
        self.assertAlmostEqual(sum(lot["pieces"] for lot in goods), 3236.8, places=2)
        self.assertAlmostEqual(sum(lot["area"] for lot in goods), 4531.52, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 3153.65, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 3207, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 20553.69, places=2)


class MattressInvoiceTest(unittest.TestCase):
    def test_packages_stay_out_of_pieces(self):
        folder = (
            Path(__file__).resolve().parents[2]
            / "documents"
            / "4_pravka"
            / "для тест"
            / "Матрац"
            / "_input"
        )
        result = analyze(folder)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        scan = next(doc for doc in result["documents"] if doc["role"] == "scan")
        self.assertEqual(scan["line_count"], 0)
        self.assertEqual(len(goods), 27)
        self.assertEqual(result["flags"], [])
        self.assertEqual(result["currency"], "USD")
        self.assertIn("CNY", result["currencies"])
        self.assertEqual(result["invoice_no"], "IDV3040C")
        self.assertEqual(result["invoice_date"], "09.01.2026")
        self.assertEqual(result["contract_date"], "01.10.2025")
        self.assertEqual(result["contract"][2], "\u0421")
        self.assertIn("HMK", result["seller"])
        self.assertNotIn("GOLDLUCK", result["seller"])
        self.assertIn("NECARGO", result["buyer"])
        self.assertIn("DAP Moscow", result["delivery"])
        self.assertEqual(result["container"], "CIMU0103406")
        self.assertEqual(result["director"], "Mr. Wang Jiandong")
        first = goods[0]
        self.assertEqual(first["vendor"], "64756")
        self.assertEqual(first["model"], "64756")
        self.assertEqual(first["pieces"], 750)
        self.assertEqual(first["packages"], 125)
        self.assertNotEqual(first["pieces"], first["packages"])
        self.assertEqual(first["unit"], "ШТ")
        self.assertEqual(first["hs"], "3926909200")
        self.assertEqual(first["brand"], "INTEX")
        self.assertIn("Intex Industries", first["producer"])
        self.assertAlmostEqual(first["net"], 1440)
        self.assertAlmostEqual(first["gross"], 1518.75)
        self.assertAlmostEqual(first["price"], 3.6864)
        self.assertAlmostEqual(first["amount"], 2764.8)
        self.assertEqual(sum(lot["pieces"] for lot in goods), 9964)
        self.assertEqual(sum(lot["packages"] for lot in goods), 1689)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 15928.17, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 16823.4, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 35319.79, places=2)
        self.assertTrue(any(doc["name"].startswith("INV") for doc in result["documents"]))


class BestwayWayTest(unittest.TestCase):
    def test_packages_and_postal_address(self):
        folder = (
            Path(__file__).resolve().parents[2]
            / "documents"
            / "4_pravka"
            / "для тест"
            / "besway 2"
            / "_input"
        )
        result = analyze(folder)
        goods = [lot for lot in result["lots"] if not lot["freight"]]
        spec = next(doc for doc in result["documents"] if doc["name"].startswith("SPEC"))
        self.assertEqual(spec["role"], "specification")
        self.assertEqual(len(goods), 14)
        self.assertEqual(result["flags"], [])
        self.assertEqual(result["currency"], "USD")
        self.assertIn("CNY", result["currencies"])
        self.assertEqual(result["invoice_no"], "NH-331005")
        self.assertEqual(result["invoice_date"], "15.03.2026")
        self.assertEqual(result["contract_date"], "01.10.2025")
        self.assertEqual(result["contract"][2], "\u0421")
        self.assertIn("HMK", result["seller"])
        self.assertNotIn("GOLDLUCK", result["seller"])
        self.assertIn("NECARGO", result["buyer"])
        self.assertIn("FCA Shanghai", result["delivery"])
        self.assertEqual(result["container"], "SORU4033371")
        self.assertEqual(result["director"], "Mr. Wang Jiandong")
        self.assertIn("RM1607", result["seller_address"])
        self.assertNotIn("Bank", result["seller_address"])
        self.assertIn("Lobacheva", result["buyer_address"])
        self.assertIn("room. 1", result["buyer_address"])
        self.assertNotIn("RM1607", result["buyer_address"])
        first = goods[0]
        self.assertEqual(first["vendor"], "310210")
        self.assertEqual(first["pieces"], 360)
        self.assertEqual(first["packages"], 10)
        self.assertNotEqual(first["pieces"], first["packages"])
        self.assertEqual(first["unit"], "ШТ")
        self.assertEqual(first["hs"], "9506620000")
        self.assertEqual(first["brand"], "Bestway")
        self.assertIn("Bestway (Nantong)", first["producer"])
        self.assertAlmostEqual(first["net"], 36.4)
        self.assertAlmostEqual(first["net_primary"], 36.4)
        self.assertAlmostEqual(first["gross"], 39.1)
        self.assertAlmostEqual(first["price"], 0.1362)
        self.assertAlmostEqual(first["amount"], 49.03)
        self.assertEqual(sum(lot["pieces"] for lot in goods), 9868)
        self.assertEqual(sum(lot["packages"] for lot in goods), 1511)
        self.assertAlmostEqual(sum(lot["net"] for lot in goods), 16366.93, places=2)
        self.assertAlmostEqual(sum(lot["gross"] for lot in goods), 17324.64, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in goods), 37775.75, places=2)


class MergedMeasureTest(unittest.TestCase):
    def test_merged_quantity_is_one_lot_and_shared_weight_stays_once(self):
        import tempfile

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Invoice and Packing list"
        ws["A1"] = "Invoice and Packing list"
        for col, name in enumerate(
            (
                "Art No.",
                "Color",
                "Quantity",
                "Unit",
                "Price (CNY)",
                "Amount (CNY)",
                "Gross Wt. (kg)",
                "Net Wt. (kg)",
                "Cartons",
                "Volume (m3)",
            ),
            1,
        ):
            ws.cell(2, col, name)
        ws["A3"] = "D680"
        ws["B3"] = "Black"
        ws["C3"] = 20
        ws["D3"] = "sets"
        ws["E3"] = 163.4
        ws["F3"] = 3268
        ws["G3"] = 33.5
        ws["H3"] = 32.5
        ws["I3"] = 2
        ws["J3"] = 0.04
        ws["B4"] = "PC"
        ws["G4"] = 19.5
        ws["H4"] = 18.5
        ws["I4"] = 2
        ws["J4"] = 0.02
        ws["G5"] = 116
        ws["H5"] = 112
        ws["I5"] = 4
        ws["J5"] = 0.5
        ws.merge_cells("A3:A5")
        ws.merge_cells("C3:C5")
        ws.merge_cells("D3:D5")
        ws.merge_cells("E3:E5")
        ws.merge_cells("F3:F5")
        ws["A6"] = "A519"
        ws["C6"] = 70
        ws["D6"] = "sets"
        ws["E6"] = 178.2
        ws["F6"] = 12474
        ws["G6"] = 425
        ws["H6"] = 370
        ws["I6"] = 2
        ws["J6"] = 2.48
        for row, art, qty, price, amount in (
            (7, "A711", 10, 160.2, 1602),
            (8, "A402S", 10, 106.95, 1069.5),
            (9, "A402", 10, 108.1, 1081),
        ):
            ws.cell(row, 1, art)
            ws.cell(row, 3, qty)
            ws.cell(row, 4, "sets")
            ws.cell(row, 5, price)
            ws.cell(row, 6, amount)
        ws.merge_cells("G6:G9")
        ws.merge_cells("H6:H9")
        ws.merge_cells("I6:I9")
        ws.merge_cells("J6:J9")
        dest = Path(tempfile.mkdtemp())
        wb.save(dest / "sheet.xlsx")
        lots = {lot["vendor"]: lot for lot in analyze(dest)["lots"]}
        self.assertEqual(set(lots), {"D680", "A519", "A711", "A402S", "A402"})
        d680 = lots["D680"]
        self.assertEqual(d680["pieces"], 20)
        self.assertAlmostEqual(d680["amount"], 3268)
        self.assertAlmostEqual(d680["packages"], 8)
        self.assertAlmostEqual(d680["net"], 163)
        self.assertAlmostEqual(d680["gross"], 169)
        self.assertAlmostEqual(d680["volume"], 0.56)
        self.assertEqual(lots["A519"]["packages"], 2)
        self.assertAlmostEqual(lots["A519"]["net"], 370)
        self.assertIsNone(lots["A711"]["packages"])
        self.assertIsNone(lots["A711"]["net"])
        self.assertIsNone(lots["A402"]["gross"])


class DescriptionEchoTest(unittest.TestCase):
    def test_echo_row_is_the_description_and_one_spec_row_keeps_its_weight(self):
        import tempfile

        from openpyxl import Workbook

        dest = Path(tempfile.mkdtemp())
        invoice = Workbook()
        sheet = invoice.active
        sheet.title = "Invoice"
        sheet["A1"] = "COMMERCIAL INVOICE"
        for col, name in enumerate(
            ("NO.", "DESIGN", "H.S. CODE", "PACKAGES", "QUANTITY", "UNIT M/PC", "UNIT PRICE(RMB)", "AMOUNT(RMB)"),
            1,
        ):
            sheet.cell(3, col, name)
        sheet["A4"] = 1
        sheet["B4"] = "Профиль О-30"
        sheet["C4"] = 3926909090
        sheet["D4"] = 3
        sheet["E4"] = 18000
        sheet["F4"] = "M"
        sheet["G4"] = 0.13
        sheet["H4"] = 2340
        sheet["B5"] = "Мебельный профиль Профиль О-30"
        sheet["D5"] = 3
        sheet["E5"] = 18000
        sheet["H5"] = 2340
        sheet["A6"] = 2
        sheet["B6"] = "MD813"
        sheet["C6"] = 7318230000
        sheet["D6"] = 1
        sheet["E6"] = 200100
        sheet["F6"] = "PC"
        sheet["G6"] = 0.049
        sheet["H6"] = 9804.9
        sheet["B7"] = "мебельная фурнитура MD813"
        sheet["D7"] = 1
        sheet["E7"] = 200100
        sheet["H7"] = 9804.9
        invoice.save(dest / "invoice.xlsx")

        spec = Workbook()
        body = spec.active
        body.title = "Sheet1"
        body["A1"] = "SPECIFICATION"
        for col, name in enumerate(
            ("ROLL No.", "PRODUCT NAME", "QUANTITY", "UNIT", "N.W", "G.W"),
            1,
        ):
            body.cell(3, col, name)
        for row, qty, net, gross in ((4, 6000, 28, 30), (5, 6000, 28, 30), (6, 6000, 28, 30)):
            body.cell(row, 2, "Профиль О-30")
            body.cell(row, 3, qty)
            body.cell(row, 4, "M")
            body.cell(row, 5, net)
            body.cell(row, 6, gross)
        body["B7"] = "MD813"
        body["C7"] = 200100
        body["D7"] = "PC"
        body["E7"] = 1311
        body["F7"] = 1344.8
        spec.save(dest / "spec.xlsx")

        lots = analyze(dest)["lots"]
        self.assertEqual(len(lots), 2)
        profile, fitting = lots
        self.assertEqual(profile["model"], "Профиль О-30")
        self.assertIn("Мебельный профиль", profile["description"])
        self.assertEqual(profile["pieces"], 18000)
        self.assertEqual(profile["packages"], 3)
        self.assertAlmostEqual(profile["amount"], 2340)
        self.assertAlmostEqual(profile["net"], 84)
        self.assertAlmostEqual(profile["gross"], 90)
        self.assertEqual(fitting["model"], "MD813")
        self.assertIn("фурнитура", fitting["description"])
        self.assertEqual(fitting["pieces"], 200100)
        self.assertAlmostEqual(fitting["net"], 1311)
        self.assertAlmostEqual(fitting["gross"], 1344.8)
        self.assertAlmostEqual(sum(lot["amount"] for lot in lots), 12144.9)


def _roll(vendor, meters, net, gross, model="", price=10.0, packages=1):
    return {
        "freight": False,
        "vendor": vendor,
        "model": model,
        "description": "FABRIC",
        "hs": "54077100",
        "hs_alt": "",
        "unit": "meters",
        "package_type": "ROLLS",
        "pieces": meters,
        "packages": packages,
        "price": price,
        "amount": round(price * meters, 2),
        "net": net,
        "gross": gross,
        "conflicts": {},
    }


class FoldRollsTest(unittest.TestCase):
    def test_rolls_of_one_design_become_one_row(self):
        from prepare_transform_code.rolls import fold_rolls

        lots = [
            _roll("LIV 600", 25.5, 22.0, 22.2, model="PM1"),
            _roll("LIV 600", 27.0, 23.3, 23.5, model="PM2"),
            _roll("OXF 1", 26.0, 20.8, 21.0, model="PM3", price=9.0),
        ]
        out, notes = fold_rolls(lots)
        self.assertEqual(len(out), 2)
        self.assertEqual(out[0]["packages"], 2)
        self.assertAlmostEqual(out[0]["pieces"], 52.5)
        self.assertAlmostEqual(out[0]["net"], 45.3)
        self.assertAlmostEqual(out[0]["gross"], 45.7)
        self.assertAlmostEqual(out[0]["amount"], 525.0)
        self.assertEqual(out[0]["model"], "")
        self.assertEqual(out[1]["packages"], 1)
        self.assertEqual(len(notes), 1)

    def test_repeated_model_splits_a_design_into_colours(self):
        from prepare_transform_code.rolls import fold_rolls

        lots = [
            _roll("EDERAS", 10.0, 5.0, 5.5, model="ERBA 27"),
            _roll("EDERAS", 11.0, 6.0, 6.5, model="ERBA 01"),
            _roll("EDERAS", 12.0, 7.0, 7.5, model="ERBA 27"),
            _roll("EDERAS", 13.0, 8.0, 8.5, model="ERBA 01"),
        ]
        out, _notes = fold_rolls(lots)
        self.assertEqual([lot["vendor"] for lot in out], ["EDERAS ERBA 27", "EDERAS ERBA 01"])
        self.assertEqual([lot["packages"] for lot in out], [2, 2])
        self.assertAlmostEqual(out[0]["pieces"], 22.0)

    def test_missing_weight_stays_empty_not_partial(self):
        from prepare_transform_code.rolls import fold_rolls

        lots = [_roll("A", 10.0, 5.0, 5.5), _roll("A", 11.0, None, 6.5)]
        out, _notes = fold_rolls(lots)
        self.assertIsNone(out[0]["net"])
        self.assertAlmostEqual(out[0]["gross"], 12.0)

    def test_lots_with_many_places_or_other_prices_are_left_alone(self):
        from prepare_transform_code.rolls import fold_rolls

        lots = [
            _roll("A", 10.0, 5.0, 5.5, packages=4),
            _roll("A", 11.0, 6.0, 6.5, packages=3),
            _roll("B", 10.0, 5.0, 5.5, price=1.0),
            _roll("B", 10.0, 5.0, 5.5, price=2.0),
        ]
        out, notes = fold_rolls(lots)
        self.assertEqual(len(out), 4)
        self.assertEqual(notes, [])

    def test_places_are_checked_against_the_specification(self):
        from prepare_transform_code.rolls import fold_rolls

        lots = [_roll("LIV 600", 25.5, 22.0, 22.2), _roll("LIV 600", 27.0, 23.3, 23.5)]
        same = fold_rolls(lots, [{"vendor": "LIV 600", "model": "", "packages": 2}])[0]
        self.assertFalse(same[0]["packages_conflict"])
        other = fold_rolls(lots, [{"vendor": "LIV 600", "model": "", "packages": 3}])[0]
        self.assertTrue(other[0]["packages_conflict"])

    def test_aydin_rolls_become_three_colour_rows(self):
        from prepare_transform_code.rolls import fold_rolls

        result = _only(
            Path("documents/6_pravka/!Aydin"),
            ("AYDIN INVOICE 17255.xlsx", "AYDIN PL 17255.xlsx", "Копия Specification  Айдын 17255 N.xlsx"),
        )
        lots, _notes = fold_rolls([lot for lot in result["lots"] if not lot["freight"]], result["spec_rows"])
        self.assertEqual(sorted(lot["packages"] for lot in lots), [5, 5, 6])
        self.assertAlmostEqual(sum(lot["pieces"] for lot in lots), 589.94, places=2)
        self.assertAlmostEqual(sum(lot["amount"] for lot in lots), 4802.11, places=2)
        self.assertFalse(any(lot["packages_conflict"] for lot in lots))
        self.assertEqual(sorted(round(lot["pieces"], 2) for lot in lots), [163.88, 207.27, 218.79])

    def test_pehlivan_rolls_become_twelve_designs(self):
        from prepare_transform_code.rolls import fold_rolls

        result = _only(
            Path("documents/6_pravka/!Pehlivan"),
            (
                "Specification Пехливан 17255.xlsx",
                "ИНВОЙС ПЕХЛИВАН 17255.pdf",
                "ПАКИНГ ПЕХЛИВАН 17255.xlsx",
            ),
        )
        lots, _notes = fold_rolls([lot for lot in result["lots"] if not lot["freight"]], result["spec_rows"])
        self.assertEqual(len(lots), 12)
        self.assertEqual(sum(lot["packages"] for lot in lots), 53)
        self.assertAlmostEqual(sum(lot["pieces"] for lot in lots), 1473.8, places=2)
        self.assertAlmostEqual(sum(lot["net"] for lot in lots), 1170.8, places=2)
        liverpool = next(lot for lot in lots if lot["vendor"] == "LIVERPOOL 600")
        self.assertEqual(liverpool["packages"], 11)
        self.assertAlmostEqual(liverpool["pieces"], 314.3, places=2)
        self.assertFalse(any(lot["packages_conflict"] for lot in lots))


class DocumentTablesTest(unittest.TestCase):
    def test_each_file_keeps_its_own_rows(self):
        result = _only(
            Path("documents/6_pravka/!Aydin"),
            ("AYDIN INVOICE 17255.xlsx", "AYDIN PL 17255.xlsx", "Копия Specification  Айдын 17255 N.xlsx"),
        )
        tables = result["document_tables"]
        packing = next(table for name, table in tables.items() if "PL" in name)
        invoice = next(table for name, table in tables.items() if "INVOICE" in name)
        self.assertEqual(len(packing["rows"]), 16)
        self.assertEqual(len(invoice["rows"]), 1)
        self.assertEqual(result["currency_printed"], "USD")


def _only(folder, names):
    import shutil
    import tempfile

    dest = Path(tempfile.mkdtemp())
    try:
        for name in names:
            shutil.copy(folder / name, dest / name)
        return analyze(dest)
    finally:
        for child in dest.iterdir():
            child.unlink()
        dest.rmdir()


if __name__ == "__main__":
    unittest.main()
