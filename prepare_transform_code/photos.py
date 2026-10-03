"""Фото страниц PDF и листов Excel для разбора. Ячейку не режем: перенос только по границе клетки."""

import re
from pathlib import Path

from openpyxl.utils import get_column_letter

_PDF_SCALE = 140 / 72
_PDF_MAX_SIDE = 2000
_PAGE_MAX_W = 2400
_PAGE_MAX_H = 1500
_ROW_CAP = 2500
_COL_CAP = 80
_FONT_SIZE = 15
_LINE_H = 18
_PAD = 4


def render_folder(folder, dest):
    folder = Path(folder)
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    written = []
    for path in sorted(folder.iterdir(), key=lambda item: item.name.lower()):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        suffix = path.suffix.lower()
        slug = _slug(path)
        if suffix == ".pdf":
            written.extend(_pdf(path, dest, slug))
        elif suffix in {".xlsx", ".xlsm", ".xls"}:
            written.extend(_excel(path, dest, slug))
        elif suffix in {".jpg", ".jpeg", ".png"}:
            written.extend(_raster(path, dest, slug))
    return written


def _slug(path):
    low = path.stem.lower()
    role = ""
    if "proforma" in low or "проформ" in low:
        role = "proforma"
    elif "dispatch" in low or "диспат" in low:
        role = "invoice_dispatch"
    elif "pack" in low or "пакинг" in low or "packing" in low:
        role = "packing"
    elif "invoice" in low or "инвойс" in low:
        role = "invoice"
    elif "spec" in low or "спец" in low:
        role = "specification"
    extra = re.sub(r"[^\w]+", "_", path.stem, flags=re.UNICODE).strip("_")
    extra = re.sub(
        r"(?i)(proforma|проформ\w*|dispatch|диспат\w*|packing|пакинг|pack|invoice|инвойс|specification|спецификац\w*|spec)",
        "",
        extra,
    ).strip("_")
    if role and extra and extra.lower() not in {"ok", "pdf", "new", "fin", "final"}:
        return f"{role}_{extra[:24]}"
    if role:
        return role
    cleaned = re.sub(r"[^\w]+", "_", path.stem, flags=re.UNICODE).strip("_")
    return cleaned[:40] or "sheet"


def _pdf(path, dest, slug):
    import pypdfium2 as pdfium

    written = []
    pdf = pdfium.PdfDocument(str(path))
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                image = page.render(scale=_PDF_SCALE).to_pil().convert("RGB")
                image = _upright(image, _page_words(path, index))
                image = _fit(image, _PDF_MAX_SIDE)
                out = dest / f"{slug}_p{index + 1:02d}of{len(pdf):02d}.jpg"
                image.save(out, format="JPEG", quality=82, optimize=True)
                written.append(out)
            finally:
                page.close()
    finally:
        pdf.close()
    return written


def _raster(path, dest, slug):
    from PIL import Image

    image = Image.open(path).convert("RGB")
    image = _upright(image)
    image = _fit(image, 2400)
    out = dest / f"{slug}.jpg"
    image.save(out, format="JPEG", quality=82, optimize=True)
    return [out]


def _page_words(path, index):
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        if index >= len(pdf.pages):
            return []
        return pdf.pages[index].extract_words() or []


def _upright(image, words=None):
    """Скан на боку повернуть так, чтобы строки шли горизонтально.

    Страница, где текстовый слой уже лежит строками, совпадает с просмотрщиком.
    Её не крутить: плотная таблица внизу бланка не значит, что лист вверх ногами.
    """
    from PIL import Image

    if _horizontal_text(words):
        return image
    sideways = image.transpose(Image.ROTATE_90)
    if _row_variation(sideways) > _row_variation(image) * 1.12:
        image = sideways
    if not words:
        top, bottom = _black_ends(image)
        if bottom > top * 1.4:
            image = image.transpose(Image.ROTATE_180)
    return image


def _horizontal_text(words):
    """Слова шире своей высоты. Кадр уже прямой, эвристика краски его не трогает."""
    if not words or len(words) < 12:
        return False
    wide = 0
    for word in words:
        width = float(word.get("x1", 0)) - float(word.get("x0", 0))
        height = float(word.get("bottom", 0)) - float(word.get("top", 0))
        if width > max(height, 1) * 1.15 and len(str(word.get("text") or "")) >= 2:
            wide += 1
    return wide >= max(8, len(words) * 0.35)


def _black_ends(image):
    """Чёрная краска сверху и снизу. Печать обычно цветная и сюда не попадает."""
    small = image.convert("RGB").resize((240, 320))
    width, height = small.size
    pixels = small.load()
    band = max(8, int(height * 0.16))

    def count(y0, y1):
        total = 0
        for y in range(y0, y1):
            for x in range(width):
                red, green, blue = pixels[x, y]
                if red < 80 and green < 80 and blue < 80:
                    total += 1
        return total

    return count(0, band), count(height - band, height)


def _row_variation(image):
    gray = image.convert("L").resize((120, 160))
    width, height = gray.size
    pixels = gray.load()
    bands = []
    for y in range(height):
        bands.append(sum(1 for x in range(width) if pixels[x, y] < 170))
    return sum(abs(bands[index] - bands[index - 1]) for index in range(1, height))


def _fit(image, limit):
    width, height = image.size
    longest = max(width, height)
    if longest <= limit:
        return image
    scale = limit / longest
    return image.resize((max(1, int(width * scale)), max(1, int(height * scale))))


def _font():
    from PIL import ImageFont

    for candidate in (
        r"C:\Windows\Fonts\arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, _FONT_SIZE)
    return ImageFont.load_default()


def _excel(path, dest, slug):
    if path.suffix.lower() == ".xls":
        sheets = _xls_sheets(path)
    else:
        sheets = _xlsx_sheets(path)
    written = []
    for sheet_slug, tiles in sheets:
        total = len(tiles)
        for index, image in enumerate(tiles, start=1):
            out = dest / f"{slug}_{sheet_slug}_p{index:02d}of{total:02d}.jpg"
            image.save(out, format="JPEG", quality=82, optimize=True)
            written.append(out)
    return written


def _xlsx_sheets(path):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    book = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    sheets = []
    for ws in book.worksheets:
        values = cached[ws.title] if ws.title in cached.sheetnames else None
        grid = _grid(ws, values)
        widths = {col: _col_width(ws, col, get_column_letter) for col in grid["cols"]}
        sheets.append((_sheet_slug(ws.title), _tiles(grid, widths, ws)))
    return sheets


def _xls_sheets(path):
    try:
        import xlrd
    except ImportError:
        return []
    try:
        book = xlrd.open_workbook(str(path), formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(str(path), formatting_info=False)
    sheets = []
    for sheet in book.sheets():
        grid = _xls_grid(sheet)
        widths = {col: _text_width(grid, col) for col in grid["cols"]}
        sheets.append((_sheet_slug(sheet.name), _tiles(grid, widths, None)))
    return sheets


def _text_width(grid, col):
    longest = 0
    for row in grid["rows"]:
        text = grid["cells"].get((row, col), {}).get("text", "")
        longest = max(longest, len(text.replace("\n", " ")))
    if longest > 40:
        return 340
    if longest > 16:
        return 180
    return 88


def _sheet_slug(name):
    cleaned = re.sub(r"[^\w]+", "_", str(name), flags=re.UNICODE).strip("_")
    return (cleaned[:24] or "sheet").lower()


def _tiles(grid, widths, ws):
    from PIL import Image

    if not grid["cols"]:
        image = Image.new("RGB", (400, 80), "white")
        return [image]
    font = _font()
    col_w = dict(widths)
    _widen_headers(grid, col_w, font)
    row_h = {row: _row_height(ws, grid, row, col_w, font) for row in grid["rows"]}
    planned = _planned_pages(grid, col_w, row_h)
    tiles = []
    total = len(planned)
    for index, (cols, rows) in enumerate(planned, start=1):
        tiles.append(_draw(grid, cols, rows, col_w, row_h, font, index, total))
    return tiles


def _planned_pages(grid, col_w, row_h):
    """Слева направо, потом следующая полоса строк. Пустой угол (только шапка) не снимать."""
    col_groups = _chunks(grid["cols"], col_w, _PAGE_MAX_W, grid["col_spans"])
    row_groups = _chunks(grid["rows"], row_h, _PAGE_MAX_H, grid["row_spans"])
    header_rows = _repeat_rows(grid)
    header_ids = set(header_rows)
    planned = []
    for r0, r1 in row_groups:
        rows = grid["rows"][r0:r1]
        if r0 > 0:
            rows = list(dict.fromkeys(header_rows + rows))
        for c0, c1 in col_groups:
            cols = grid["cols"][c0:c1]
            if _tile_has_body(grid, cols, rows, header_ids):
                planned.append((cols, rows))
    return planned


def _tile_has_body(grid, cols, rows, header_ids):
    for row in rows:
        if row in header_ids:
            continue
        for col in cols:
            text = grid["cells"].get((row, col), {}).get("text", "").strip()
            if text and text not in {"0", "0.0"}:
                return True
    return False


def _grid(ws, values):
    hidden = set()
    max_col = min(ws.max_column or 1, _COL_CAP)
    max_row = min(ws.max_row or 1, _ROW_CAP)
    for col in range(1, max_col + 1):
        dim = ws.column_dimensions[get_column_letter(col)]
        if dim.hidden or dim.width == 0:
            hidden.add(col)
    merges = []
    for item in ws.merged_cells.ranges:
        merges.append((item.min_row, item.min_col, item.max_row, item.max_col))
    covered = set()
    for r1, c1, r2, c2 in merges:
        for row in range(r1, r2 + 1):
            for col in range(c1, c2 + 1):
                if (row, col) != (r1, c1):
                    covered.add((row, col))
    last_row = 1
    last_col = 1
    cells = {}
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            if col in hidden:
                continue
            text = _display(ws.cell(row, col).value, values.cell(row, col).value if values is not None else None)
            wrap = bool(ws.cell(row, col).alignment.wrap_text)
            if text:
                last_row = max(last_row, row)
                last_col = max(last_col, col)
                cells[(row, col)] = {"text": text, "wrap": wrap}
    cols = [col for col in range(1, last_col + 1) if col not in hidden]
    rows = list(range(1, last_row + 1))
    col_spans = []
    row_spans = []
    for r1, c1, r2, c2 in merges:
        col_indexes = [index for index, col in enumerate(cols) if c1 <= col <= c2]
        row_indexes = [index for index, row in enumerate(rows) if r1 <= row <= r2]
        if len(col_indexes) > 1:
            col_spans.append((col_indexes[0], col_indexes[-1]))
        if len(row_indexes) > 1:
            row_spans.append((row_indexes[0], row_indexes[-1]))
    return {
        "cells": cells,
        "covered": covered,
        "merges": merges,
        "cols": cols,
        "rows": rows,
        "col_spans": col_spans,
        "row_spans": row_spans,
        "header_rows": _header_rows(cells, rows, cols),
    }


def _repeat_rows(grid):
    """На продолжении повторяется строка колонок. Число в тесте — верхние строки, как раньше."""
    raw = grid.get("header_rows") or []
    if isinstance(raw, int):
        return grid["rows"][:raw]
    return [row for row in raw if row in grid["rows"]]


def _display(formula, cached):
    value = formula
    if isinstance(formula, str) and formula.startswith("=") and cached is not None:
        value = cached
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if abs(value - round(value)) < 1e-9 and abs(value) >= 100:
            return str(int(round(value)))
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value).replace("\r", "")


def _header_rows(cells, rows, cols):
    """Строка колонок, и соседняя, если шапка в две строки. Письмо над таблицей не повторять."""
    from prepare_transform_code.fields import column_of

    scores = []
    for row in rows[:12]:
        score = 0
        for col in cols:
            cell = cells.get((row, col))
            if cell and column_of(cell["text"]):
                score += 1
        scores.append(score)
    if not scores or max(scores) < 2:
        return rows[:1]
    best = scores.index(max(scores))
    picked = {rows[best]}
    for neighbor in (best - 1, best + 1):
        if 0 <= neighbor < len(scores) and scores[neighbor] >= 2:
            picked.add(rows[neighbor])
    return sorted(picked)


def _widen_headers(grid, col_w, font):
    """Узкая колонка не должна рвать слово шапки посередине: расширяем до самого длинного слова."""
    repeated = _repeat_rows(grid)
    if not repeated:
        return
    header = repeated[-1]
    for col in grid["cols"]:
        text = grid["cells"].get((header, col), {}).get("text", "")
        if not text:
            continue
        words = text.replace("/", " ").replace("\\", " ").split()
        longest = max((_length(font, word) for word in words), default=0)
        col_w[col] = int(max(col_w[col], min(180, longest + _PAD * 4)))


def _col_width(ws, col, get_column_letter):
    width = ws.column_dimensions[get_column_letter(col)].width
    if not width:
        width = 11
    return max(48, min(380, int(width * 8)))


def _row_height(ws, grid, row, col_w, font):
    height = 22
    if ws is not None:
        excel_h = ws.row_dimensions[row].height
        if excel_h:
            height = max(height, int(excel_h * 1.4))
    for r1, c1, r2, c2 in grid["merges"]:
        if r1 != row:
            continue
        text = grid["cells"].get((r1, c1), {}).get("text", "")
        width = sum(col_w.get(col, 0) for col in range(c1, c2 + 1)) - _PAD * 2
        height = max(height, _text_height(text, font, max(40, width)))
    for col in grid["cols"]:
        if (row, col) in grid["covered"]:
            continue
        if any(r1 == row and c1 == col for r1, c1, r2, c2 in grid["merges"]):
            continue
        cell = grid["cells"].get((row, col))
        if not cell:
            continue
        width = _span_width(grid, row, col, col_w) if not cell.get("wrap") else col_w[col]
        height = max(height, _text_height(cell["text"], font, width - _PAD * 2))
    return height


def _span_width(grid, row, col, col_w):
    """Текст без переноса в Excel вылезает в пустые клетки справа, колонку не раздувая."""
    width = 0
    for item in grid["cols"]:
        if item < col:
            continue
        if item != col:
            other = grid["cells"].get((row, item))
            if other and str(other.get("text") or "").strip():
                break
        width += col_w.get(item, 0)
    return max(width, col_w.get(col, 48))


def _text_height(text, font, width):
    lines = _wrap(text, font, max(20, width))
    return max(22, len(lines) * _LINE_H + _PAD * 2)


def _chunks(keys, sizes, limit, spans):
    groups = []
    start = 0
    used = 0
    index = 0
    count = len(keys)
    while index < count:
        end = index
        for left, right in spans:
            if left <= index <= right:
                end = max(end, right)
        width = sum(sizes[keys[cursor]] for cursor in range(index, end + 1))
        if start < index and used + width > limit:
            groups.append((start, index))
            start = index
            used = 0
            continue
        used += width
        index = end + 1
    if start < count:
        groups.append((start, count))
    return groups or [(0, count)]


def _draw(grid, cols, rows, col_w, row_h, font, index, total):
    from PIL import Image, ImageDraw

    width = 16 + sum(col_w[col] for col in cols)
    height = 36 + sum(row_h[row] for row in rows)
    canvas = Image.new("RGB", (max(width, 80), max(height, 80)), "white")
    draw = ImageDraw.Draw(canvas)
    title = f"photo {index}/{total}  rows {_span_label(rows)}  cols {cols[0]}-{cols[-1]}"
    draw.text((8, 8), title, fill="black", font=font)
    y = 32
    row_y = {}
    for row in rows:
        row_y[row] = y
        y += row_h[row]
    col_x = {}
    x = 8
    for col in cols:
        col_x[col] = x
        x += col_w[col]
    boxes = []
    drawn = set()
    for row in rows:
        for col in cols:
            if (row, col) in drawn or (row, col) in grid["covered"]:
                continue
            merge = _merge_at(grid["merges"], row, col)
            if merge:
                r1, c1, r2, c2 = merge
                span_cols = [item for item in cols if c1 <= item <= c2]
                span_rows = [item for item in rows if r1 <= item <= r2]
                if not span_cols or not span_rows or span_cols[0] != col or span_rows[0] != row:
                    continue
                box_w = sum(col_w[item] for item in span_cols)
                box_h = sum(row_h[item] for item in span_rows)
                for item_row in span_rows:
                    for item_col in span_cols:
                        drawn.add((item_row, item_col))
            else:
                box_w = col_w[col]
                box_h = row_h[row]
                drawn.add((row, col))
            boxes.append((row, col, col_x[col], row_y[row], box_w, box_h, merge))
    for row, col, x0, y0, box_w, box_h, _merge in boxes:
        draw.rectangle((x0, y0, x0 + box_w, y0 + box_h), outline="#b0b0b0", fill="white")
    for row, col, x0, y0, box_w, _box_h, merge in boxes:
        cell = grid["cells"].get((row, col), {})
        text = cell.get("text", "")
        if not text:
            continue
        text_w = box_w
        if not merge and not cell.get("wrap"):
            text_w = _span_width(grid, row, col, col_w)
        for line_index, line in enumerate(_wrap(text, font, text_w - _PAD * 2)):
            draw.text((x0 + _PAD, y0 + _PAD + line_index * _LINE_H), line, fill="black", font=font)
    return canvas


def _merge_at(merges, row, col):
    for item in merges:
        r1, c1, r2, c2 = item
        if r1 == row and c1 == col:
            return item
    return None


def _xls_grid(sheet):
    max_row = min(sheet.nrows, _ROW_CAP)
    max_col = min(sheet.ncols, _COL_CAP)
    merges = []
    for rlo, rhi, clo, chi in getattr(sheet, "merged_cells", []) or []:
        merges.append((rlo + 1, clo + 1, rhi, chi))
    covered = set()
    for r1, c1, r2, c2 in merges:
        for row in range(r1, r2 + 1):
            for col in range(c1, c2 + 1):
                if (row, col) != (r1, c1):
                    covered.add((row, col))
    cells = {}
    last_row = 1
    last_col = 1
    for row in range(max_row):
        for col in range(max_col):
            value = sheet.cell_value(row, col)
            text = _display(value, None)
            if not text:
                continue
            last_row = max(last_row, row + 1)
            last_col = max(last_col, col + 1)
            cells[(row + 1, col + 1)] = {"text": text, "wrap": True}
    cols = list(range(1, last_col + 1))
    rows = list(range(1, last_row + 1))
    col_spans = []
    row_spans = []
    for r1, c1, r2, c2 in merges:
        col_indexes = [index for index, col in enumerate(cols) if c1 <= col <= c2]
        row_indexes = [index for index, row in enumerate(rows) if r1 <= row <= r2]
        if len(col_indexes) > 1:
            col_spans.append((col_indexes[0], col_indexes[-1]))
        if len(row_indexes) > 1:
            row_spans.append((row_indexes[0], row_indexes[-1]))
    return {
        "cells": cells,
        "covered": covered,
        "merges": merges,
        "cols": cols,
        "rows": rows,
        "col_spans": col_spans,
        "row_spans": row_spans,
        "header_rows": _header_rows(cells, rows, cols),
    }


def _span_label(rows):
    parts = []
    start = prev = rows[0]
    for row in rows[1:]:
        if row == prev + 1:
            prev = row
            continue
        parts.append(f"{start}-{prev}" if start != prev else str(start))
        start = prev = row
    parts.append(f"{start}-{prev}" if start != prev else str(start))
    return ",".join(parts)


def _wrap(text, font, width):
    if not text:
        return [""]
    lines = []
    for paragraph in str(text).split("\n"):
        current = ""
        for word in paragraph.split(" ") or [""]:
            for piece in _fit_token(word, font, width):
                trial = piece if not current else current + " " + piece
                if _length(font, trial) <= width or not current:
                    current = trial
                else:
                    lines.append(current)
                    current = piece
        lines.append(current)
    return lines or [""]


def _fit_token(word, font, width):
    if not word or _length(font, word) <= width:
        return [word]
    parts = []
    buf = ""
    for char in word:
        trial = buf + char
        if _length(font, trial) <= width or not buf:
            buf = trial
        else:
            parts.append(buf)
            buf = char
    if buf:
        parts.append(buf)
    return parts or [word]


def _length(font, text):
    if hasattr(font, "getlength"):
        return font.getlength(text)
    return len(text) * _FONT_SIZE * 0.6
