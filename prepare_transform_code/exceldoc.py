from openpyxl import load_workbook

import re

from prepare_transform_code.fields import (
    column_of,
    stamp_header,
    is_freight,
    is_header_row,
    is_row_index,
    is_shipper_code,
    is_size_label,
    is_stop_label,
    role_of,
    roles_of,
)
from prepare_transform_code.lines import (
    Line,
    assign_cell,
    attach_pallet,
    fold_parts,
    hs_text,
    is_bare_total,
    is_pallet_only,
    pull_article,
    settle_pallets,
)
from prepare_transform_code.numbers import currency_of, parse_number

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None


def read_excel(path):
    suffix = path.suffix.lower()
    if suffix == ".xls":
        sheets = _xls(path)
    else:
        sheets = _xlsx(path)
    docs = []
    for sheet in sheets:
        preview = _preview(sheet)
        role = _role(sheet["name"], sheet["rows"])
        roles = _opening_roles(sheet["rows"])
        if role not in roles and role not in {"unknown", "draft", "gtd_form", "customs_appendix"}:
            roles.insert(0, role)
        lines = [] if role in {"draft", "gtd_form", "customs_appendix"} else _lines(sheet["rows"])
        docs.append(
            {
                "kind": "excel",
                "sheet": sheet["name"],
                "role": role,
                "roles": roles,
                "text": preview,
                "currency": currency_of(preview),
                "readable": True,
                "lines": lines,
            }
        )
    return docs


def _preview(sheet):
    rows = sheet["rows"]
    if len(rows) <= 40:
        body = rows
    else:
        # Шапка и подвал. Тело таблицы в текст сторон и номеров не кладём.
        body = rows[:15] + rows[-25:]
    extra = _label_pairs(rows)
    # Пары «подпись — значение» раньше плоского хвоста: подпись внизу листа без имени не затирает фирму.
    text = sheet["name"] + "\n"
    if extra:
        text += extra + "\n"
    return text + _join(body)


def _role(name, rows):
    head = (name + "\n" + _join(rows[:4])).upper()
    whole = _join(rows).upper()
    if "ПРИЛОЖЕНИЕ К ИНВОЙСУ" in head or name.upper() == "ГТД":
        return "customs_appendix"
    if "ЧЕРНОВИК" in whole:
        return "draft"
    if "ФОРМА ДЛЯ ГТД" in head:
        return "gtd_form"
    # Роль берётся из заголовка листа. Слово invoice в шапке колонки заголовок не заменяет.
    opened = _opening_role(rows)
    if opened != "unknown":
        return opened
    found = role_of(head)
    if found != "unknown":
        return found
    if _calculation(rows):
        return "draft"
    return "unknown"


def _opening_roles(rows):
    """Все роли, названные в шапке листа. Инвойс и пакинг в одном заголовке — обе."""
    lines = []
    for row in rows[:12]:
        text = " ".join(str(item) for item in row if item not in (None, ""))
        if text:
            lines.append(text)
    return roles_of("\n".join(lines))


def _calculation(rows):
    """Колонки проходной, дороги и цены с транспортом — расчёт, не товарный лист."""
    blob = _join(rows[:8]).lower()
    return any(
        mark in blob
        for mark in ("проходн", "инвойс без дороги", "цена за штуку округл", "стоимость с транспортом")
    )


def _opening_role(rows):
    """Заголовок может стоять ниже шапки фирмы. Первая найденная роль и есть роль листа."""
    opened = []
    for row in rows[:12]:
        cell = next((str(item) for item in row if item not in (None, "")), "")
        if not cell:
            continue
        opened.append(cell)
        found = role_of("\n".join(opened))
        if found != "unknown":
            return found
    return "unknown"


def _footer_label(cell):
    """Подпись подвала, не шапка колонки товара."""
    text = " ".join(str(cell or "").replace("\n", " ").split())
    if not text or len(text) > 140:
        return False
    # Короткая шапка колонки — не подвал. «Country of origin:» с двоеточием в ряду других подписей — подвал.
    low = text.lower()
    marks = (
        "contract", "контракт", "seller", "продавец", "buyer", "покупатель",
        "origin", "происхожд", "delivery", "поставк", "payment", "оплат",
        "swift", "bank", "банк", "address", "адрес",
    )
    if not any(mark in low for mark in marks):
        return False
    if column_of(text) and ":" not in text:
        return False
    return True


def _label_pairs(rows):
    """Подпись и значение в одной колонке. Плоский порядок клеток значение не подменяет."""
    chunks = []
    for index, row in enumerate(rows):
        labels = [(col, cell) for col, cell in enumerate(row) if _footer_label(cell)]
        if not labels or all(column_of(cell) for _, cell in labels):
            continue
        for col, cell in labels:
            for right in range(col + 1, min(len(row), col + 5)):
                other = row[right]
                if other in (None, ""):
                    continue
                if _footer_label(other):
                    break
                chunks.append(str(cell))
                chunks.append(str(other))
                break
        if index + 1 >= len(rows) or len(labels) < 2:
            continue
        nxt = rows[index + 1]
        for col, cell in labels:
            if col >= len(nxt) or nxt[col] in (None, ""):
                continue
            if _footer_label(nxt[col]):
                continue
            chunks.append(str(cell))
            chunks.append(str(nxt[col]))
    return "\n".join(chunks)


def _join(rows):
    chunks = []
    for row in rows:
        chunks.extend(str(cell) for cell in row if cell not in (None, ""))
    return "\n".join(chunks)


def _lines(rows):
    header_idx, mapping = _header(rows)
    if header_idx is None:
        return []
    pair_unit = _qty_unit(rows[header_idx][mapping["qty"]]) if "qty" in mapping else ""
    header_cells = rows[header_idx]
    lines = []
    scale = _label_scale(header_cells, mapping)
    carried = ""
    body = rows[header_idx + 1 :]
    for row in body:
        texts = [str(cell).strip() for cell in row if cell not in (None, "")]
        if not texts:
            continue
        if is_header_row(row):
            # Вторая строка той же шапки (перевод) колонками не заменяет первую.
            # Повтор шапки уже после товара начинает новую размерную сетку.
            if lines:
                mapped, _score = _map_row(row)
                if _header_ready(mapped):
                    mapping = mapped
                    header_cells = row
                    pair_unit = _qty_unit(row[mapping["qty"]]) if "qty" in mapping else ""
                    scale = _label_scale(row, mapping)
            continue
        if any(is_stop_label(cell) for cell in texts):
            rolls = _roll_total(texts)
            if rolls and len(lines) == 1 and lines[0].packages is None:
                lines[0].packages = rolls
                lines[0].package_type = lines[0].package_type or "roll"
            break
        fresh = _row_scale(row, mapping)
        if fresh:
            # Новая подпись размеров (мужские, затем женские) сменяет сетку ниже.
            scale = fresh
            continue
        if len(texts) == 1 and not any(ch.isdigit() for ch in texts[0]):
            continue
        line = Line(source="excel")
        for name, col in mapping.items():
            if not isinstance(col, int) or col >= len(row):
                continue
            field = "description" if name == "description_2" else name
            if field == "vendor" and is_row_index(header_cells[col] if col < len(header_cells) else "", row[col]):
                continue
            assign_cell(line, field, row[col], header_is_package=False)
        if _section_title(line.description):
            line.description = ""
        if scale:
            run = _size_run(row, scale)
            if run:
                line.size = run
        if pair_unit and not line.unit:
            line.unit = pair_unit
        _unit_beside_divisor(line, row, mapping)
        if _same_code(line.hs, line.hs_alt):
            line.hs_alt = ""
        if line.hs_alt and mapping.get("_hs_alt_shipper"):
            line.hs_alt_shipper = True
        pull_article(line)
        # Короткое имя варианта написано один раз, ниже клетка пустая. Длинное описание на соседние строки не переносим.
        if (
            not line.description
            and _short_label(carried)
            and line.vendor
            and line.pieces is not None
        ):
            line.description = carried
        if line.description and line.vendor and line.pieces is not None and _short_label(line.description):
            carried = line.description
        if is_freight(line.description):
            line.freight = True
        if line.pieces is None and line.price and line.amount:
            pass
        # Колонка мест вдруг держит заводской код: ниже уже другая таблица (рулоны), не те же колонки.
        if lines and _places_became_code(row, mapping):
            break
        useful = any(
            getattr(line, name) is not None
            for name in ("pieces", "packages", "price", "amount", "gross", "net", "gross_with_pallet")
        )
        if is_bare_total(line, lines):
            break
        if is_pallet_only(line):
            attach_pallet(lines, line)
            continue
        if _blank_charge(line):
            continue
        if lines and _package_part(line, lines[-1]):
            _add_measures(lines[-1], line)
            continue
        if useful or (line.vendor and line.description):
            lines.append(line)
    folded = fold_parts(lines)
    settle_pallets(folded)
    return folded


def _places_became_code(row, mapping):
    """В колонке мест стоит код вроде YS950700VLT0690110 — шапка выше уже не про эту таблицу."""
    col = mapping.get("packages")
    if not isinstance(col, int) or col >= len(row):
        return False
    compact = re.sub(r"\s+", "", str(row[col] or "").strip())
    # 24 и 24 rolls начинаются с цифры. Заводской код начинается с буквы, внутри цифры.
    return bool(re.match(r"[A-Za-z]", compact) and re.search(r"\d", compact) and len(compact) >= 8)


def _package_part(line, previous):
    """Строка без своего количества и цены, тот же артикул или пустой, — укладка лота выше."""
    if line.pieces is not None or line.price is not None or line.amount is not None:
        return False
    if not any(getattr(line, name) is not None for name in ("packages", "net", "gross", "volume")):
        return False
    if line.vendor and previous.vendor and _same_vendor(line.vendor, previous.vendor):
        return True
    return not line.vendor


def _same_vendor(left, right):
    return re.sub(r"[^a-z0-9]+", "", str(left).lower()) == re.sub(r"[^a-z0-9]+", "", str(right).lower())


def _add_measures(head, part):
    for name in ("packages", "net", "gross", "volume"):
        value = getattr(part, name)
        if value is None:
            continue
        current = getattr(head, name)
        setattr(head, name, value if current is None else current + value)


def _same_code(left, right):
    digits_left = "".join(ch for ch in hs_text(left) if ch.isdigit())
    digits_right = "".join(ch for ch in hs_text(right) if ch.isdigit())
    return bool(digits_left) and digits_left == digits_right


_UNIT_WORD = re.compile(
    r"^(?:pcs|pc|psc|шт|set|компл|pairs|pair|пар|пара)"
    r"(?:\s*/\s*(?:pcs|pc|psc|шт|set|компл|pairs|pair|пар|пара))?$",
    re.I,
)


def _unit_beside_divisor(line, row, mapping):
    """Число под шапкой единицы — делитель цены. Слово pcs или шт в соседней клетке — единица."""
    col = mapping.get("unit")
    if not isinstance(col, int) or parse_number(line.unit) is None:
        return
    for neighbor in (col + 1, col - 1):
        if neighbor < 0 or neighbor >= len(row) or neighbor in mapping.values():
            continue
        text = " ".join(str(row[neighbor] or "").replace("\n", " ").split())
        if _UNIT_WORD.fullmatch(text):
            line.unit = text
            return


def _qty_unit(header):
    """Единица сидит в шапке количества: пары или метры."""
    pair = _pair_unit(header)
    if pair:
        return pair
    text = str(header or "").lower()
    if re.search(r"meters?|metres?|(?<![a-z])mt(?![a-z])", text):
        return "meters"
    return ""


def _roll_total(texts):
    """«Roll:16» в строке итога — сколько мест, не ещё один товар."""
    for text in texts:
        match = re.search(r"\brolls?\s*:?\s*(\d+)\b", str(text), re.I)
        if match:
            return float(match.group(1))
    return None


def _short_label(text):
    """Имя варианта вроде «LIVERPOOL 600»: несколько слов и цифра. Длинное описание товара сюда не входит."""
    words = str(text or "").split()
    return 1 <= len(words) <= 4 and any(ch.isdigit() for ch in text or "")


def _blank_charge(line):
    """Сбор без артикула и без количества, сумма ноль. Это не товар.
    Хвост, где в местах только 1 и больше ничего нет, тоже не товар."""
    if line.vendor or line.description or line.model or line.pieces is not None or line.price is not None:
        return False
    if line.net is not None or line.gross is not None:
        return False
    if line.packages not in (None, 1):
        return False
    return line.amount in (None, 0)


def _pair_unit(header):
    """«Кол-во пар» — пары, отдельной колонки единицы нет."""
    text = str(header or "")
    best = None
    for match in re.finditer(r"(?<![A-Za-zА-Яа-яЁё])(pairs?|пар)(?![A-Za-zА-Яа-яЁё])", text, re.I):
        if best is None or match.start() < best[0]:
            word = match.group(1).lower()
            best = (match.start(), "pairs" if word == "pair" else word)
    return best[1] if best else ""


def _row_scale(row, mapping):
    """Строка номеров размеров: цены и количества нет, подписи 4 и больше."""
    size_col = mapping.get("size")
    if not isinstance(size_col, int):
        return None
    for key in ("price", "amount", "qty"):
        col = mapping.get(key)
        if isinstance(col, int) and col < len(row) and parse_number(row[col]) not in (None, 0):
            return None
    scale = []
    for col in range(size_col, len(row)):
        number = parse_number(row[col]) if row[col] not in (None, "") else None
        if number is None or number <= 0 or number > 70:
            break
        scale.append((col, _num_token(number)))
    if len(scale) < 4:
        return None
    return scale


def _section_title(text):
    """«MEN'S AND LADY'S SHOES» в первой клетке слитой колонки — не описание лота."""
    folded = " ".join(str(text or "").lower().replace("'", "").replace("’", "").split())
    words = re.findall(r"[a-zа-яё]+", folded)
    if not words:
        return False
    allowed = {
        "men", "mens", "man", "lady", "ladys", "ladies", "women", "womens",
        "and", "shoes", "shoe", "обувь",
    }
    return set(words) <= allowed and ("shoe" in folded or "обувь" in folded)


def _size_run(row, scale):
    parts = []
    for col, token in scale:
        if col >= len(row) or row[col] in (None, ""):
            continue
        count = parse_number(row[col])
        if count is None or abs(count) < 1e-9:
            continue
        parts.append(f"{token}:{_num_token(count)}")
    # В подписи уже есть запятая (180-185, XL). Пары тогда делит точка с запятой.
    if any("," in part.split(":", 1)[0] for part in parts):
        return "; ".join(parts)
    return ", ".join(parts)


def _printed_label(cell):
    return " ".join(str(cell).replace("\n", " ").split())


def _label_scale(row, mapping):
    """Колонки с подписью размера как на бланке. Две одинаковые подписи — две клетки."""
    used = {col for col in mapping.values() if isinstance(col, int)}
    found = []
    for col, cell in enumerate(row):
        if col in used or cell in (None, ""):
            continue
        if not is_size_label(cell):
            continue
        found.append((col, _printed_label(cell)))
    return found or None


def _num_token(number):
    if abs(number - round(number)) < 1e-9:
        return str(int(round(number)))
    return f"{number:.2f}".rstrip("0").rstrip(".")


def _bare_measurement(cell):
    """Measurement без CBM — габарит коробки. Volume и Measurement (CBM) — объём строки."""
    raw = " ".join(str(cell or "").lower().replace("\n", " ").split())
    text = " ".join(re.sub(r"\([^)]*\)", " ", raw).split())
    if not re.fullmatch(r"measure(?:ment)?", text.replace(" ", "")):
        return False
    return re.search(r"cbm|m3|м3|куб", raw) is None


def _map_row(row):
    mapped = {}
    score = 0
    for col, cell in enumerate(row):
        name = column_of(cell)
        if not name:
            continue
        if name == "hs" and is_shipper_code(cell):
            if "hs_alt" not in mapped:
                mapped["hs_alt"] = col
                # Строка-маркер, не число: цикл по колонкам берёт только int.
                mapped["_hs_alt_shipper"] = "1"
            continue
        if name == "hs" and "hs" in mapped and "hs_alt" not in mapped:
            mapped["hs_alt"] = col
            continue
        if name == "description" and "description" in mapped and "description_2" not in mapped:
            mapped["description_2"] = col
            continue
        if name == "vendor" and "vendor" in mapped and "model" not in mapped:
            # «Маркировка» с названием фирмы — печать, не вторая модель.
            if not stamp_header(cell):
                mapped["model"] = col
            continue
        if name in mapped:
            if name == "volume" and _bare_measurement(row[mapped["volume"]]) and not _bare_measurement(cell):
                mapped["volume"] = col
            continue
        mapped[name] = col
        score += 2 if name in {"description", "qty", "vendor", "price"} else 1
    if "hs" not in mapped and "hs_alt" in mapped:
        mapped["hs"] = mapped.pop("hs_alt")
        mapped.pop("_hs_alt_shipper", None)
    return mapped, score


def _header_ready(mapped):
    identity = {"description", "vendor", "model"} & set(mapped)
    measures = {"qty", "price", "packages"} & set(mapped)
    return bool(identity and measures)


def _header(rows):
    best = None
    best_score = 0
    best_idx = None
    for index, row in enumerate(rows[:40]):
        mapped, score = _map_row(row)
        if score > best_score and _header_ready(mapped):
            best = mapped
            best_score = score
            best_idx = index
    return best_idx, best or {}


def _xlsx(path):
    wb = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        values = cached[ws.title] if ws.title in cached.sheetnames else None
        rows = []
        last = min(ws.max_row or 1, 2500)
        width = min(ws.max_column or 1, 40)
        for r in range(1, last + 1):
            row = []
            for c in range(1, width + 1):
                value = ws.cell(r, c).value
                if isinstance(value, str) and value.startswith("=") and values is not None:
                    number = values.cell(r, c).value
                    if number is not None:
                        value = number
                row.append(value)
            rows.append(row)
        merges = [
            (item.min_row - 1, item.max_row, item.min_col - 1, item.max_col)
            for item in ws.merged_cells.ranges
        ]
        _fill_merges(rows, merges)
        sheets.append({"name": ws.title, "rows": rows})
    return sheets


def _xls(path):
    if xlrd is None:
        return []
    try:
        book = xlrd.open_workbook(path, formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(path)
    sheets = []
    for sheet in book.sheets():
        rows = []
        width = min(sheet.ncols, 40)
        for r in range(min(sheet.nrows, 2500)):
            rows.append([sheet.cell_value(r, c) for c in range(width)])
        merges = list(getattr(sheet, "merged_cells", []) or [])
        _fill_merges(rows, merges)
        sheets.append({"name": sheet.name, "rows": rows})
    return sheets


def _fill_merges(rows, merges):
    """Пустая клетка под текстом в той же колонке слитого диапазона его получает.
    Число вниз не копируется: количество и вес, закрывающие несколько строк, не становятся числом каждой из них."""
    for rlo, rhi, clo, chi in merges:
        if rhi - rlo < 2:
            continue
        for c in range(clo, chi):
            source = None
            for r in range(rlo, rhi):
                if r >= len(rows):
                    break
                row = rows[r]
                if c < len(row) and row[c] not in (None, ""):
                    source = row[c]
                    break
            if source in (None, "") or _pure_number(source):
                continue
            for r in range(rlo, rhi):
                if r >= len(rows):
                    break
                row = rows[r]
                while len(row) <= c and len(row) < 40:
                    row.append(None)
                if c < len(row) and row[c] in (None, ""):
                    row[c] = source


def _pure_number(value):
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    text = str(value).strip()
    if not text or re.search(r"[A-Za-zА-Яа-яЁё]", text):
        return False
    return parse_number(text) is not None
