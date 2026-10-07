# -*- coding: utf-8 -*-
"""Снимок поставки для разбора. Путь меняется аргументом, новый файл не создавать.

python study_tools/dump_shipment.py "documents/7_pravka/Дегон"
"""
import argparse
import zipfile
from pathlib import Path

import openpyxl
import xlrd
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
DEFAULT_OUT = HERE / "out" / "dump.txt"


def cell_s(value, limit):
    if value is None:
        return ""
    text = " ".join(str(value).replace("\n", " | ").replace("\r", "").split())
    if len(text) > limit:
        return text[: limit - 1] + "…"
    return text


def write_grid(fh, rows, cols, getter, limit, row_offset=0):
    for row in range(rows):
        parts = []
        for col in range(cols):
            text = cell_s(getter(row, col), limit)
            if text:
                parts.append(f"C{col + 1}:{text}")
        if parts:
            fh.write(f"R{row_offset + row + 1}: " + " || ".join(parts) + "\n")


def dump_xlsx(path, fh, max_rows, max_cols, limit):
    fh.write(f"\n######## XLSX {path}\n")
    try:
        book = openpyxl.load_workbook(path, data_only=True)
    except Exception as error:
        fh.write(f"OPEN FAIL {error}\n")
        return
    for name in book.sheetnames:
        sheet = book[name]
        fh.write(
            f"\n--- sheet {name!r} dims={sheet.dimensions} "
            f"max_r={sheet.max_row} max_c={sheet.max_column}\n"
        )
        merges = list(sheet.merged_cells.ranges)
        fh.write(f"merges {len(merges)}\n")
        for merge in merges[:60]:
            fh.write(f"  merge {merge}\n")
        if len(merges) > 60:
            fh.write(f"  ... +{len(merges) - 60} merges\n")
        last = sheet.max_row or 0
        width = min(sheet.max_column or 0, max_cols)

        def getter(row, col, sheet=sheet):
            return sheet.cell(row + 1, col + 1).value

        write_grid(fh, min(last, max_rows), width, getter, limit)
        if last > max_rows:
            fh.write(f"... truncated, total rows {last}\n")
            start = max(max_rows, last - 8)
            fh.write(f"-- tail from R{start + 1}\n")
            write_grid(
                fh,
                last - start,
                width,
                lambda row, col, start=start: getter(start + row, col),
                limit,
                row_offset=start,
            )
    book.close()


def dump_xls(path, fh, max_rows, max_cols, limit):
    fh.write(f"\n######## XLS {path}\n")
    try:
        book = xlrd.open_workbook(str(path), formatting_info=False)
    except Exception as error:
        fh.write(f"OPEN FAIL {error}\n")
        return
    for name in book.sheet_names():
        sheet = book.sheet_by_name(name)
        fh.write(f"\n--- sheet {name!r} rows={sheet.nrows} cols={sheet.ncols}\n")
        merges = getattr(sheet, "merged_cells", []) or []
        fh.write(f"merges {len(merges)}\n")
        for top, bottom, left, right in list(merges)[:60]:
            fh.write(f"  merge r{top + 1}-{bottom} c{left + 1}-{right}\n")
        width = min(sheet.ncols, max_cols)

        def getter(row, col, sheet=sheet):
            return sheet.cell_value(row, col)

        write_grid(fh, min(sheet.nrows, max_rows), width, getter, limit)
        if sheet.nrows > max_rows:
            fh.write(f"... truncated, total rows {sheet.nrows}\n")
            start = max(max_rows, sheet.nrows - 8)
            fh.write(f"-- tail from R{start + 1}\n")
            write_grid(
                fh,
                sheet.nrows - start,
                width,
                lambda row, col, start=start: getter(start + row, col),
                limit,
                row_offset=start,
            )


def read_note(path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


HEADER_HINTS = (
    "no.",
    "№",
    "item",
    "art",
    "design",
    "description",
    "наименован",
    "qty",
    "quantity",
    "количест",
    "hs",
    "code",
    "код",
    "price",
    "цена",
    "amount",
    "сумма",
    "brand",
    "mark",
    "marca",
    "marque",
    "товарн",
    "торгов",
    "марка",
    "вкус",
    "flavour",
    "flavor",
    "manufacturer",
    "изготов",
    "производ",
    "origin",
    "country",
    "стран",
    "unit",
    "единиц",
    "weight",
    "нетто",
    "брутто",
    "net",
    "gross",
    "package",
    "упаков",
    "carton",
    "мест",
    "color",
    "цвет",
    "size",
    "размер",
    "model",
    "модел",
    "meters",
    "rolls",
    "width",
    "cbm",
    "volume",
    "finish",
    "тн вэд",
    "тз",
    "артикул",
    "commodity",
    "photos",
)


def _norm_header(text):
    return " ".join(str(text).replace("\n", " ").replace("\r", " ").split())


def _header_score(values):
    hits = 0
    nonempty = 0
    for value in values:
        if value is None or str(value).strip() == "":
            continue
        nonempty += 1
        folded = str(value).casefold()
        if any(hint in folded for hint in HEADER_HINTS):
            hits += 1
    if nonempty < 3 or hits < 2:
        return 0
    return hits


def _row_values(count, getter):
    return [getter(col) for col in range(count)]


def write_header_rows(fh, label, rows):
    if not rows:
        return
    seen = set()
    for row_no, values in rows:
        cells = [_norm_header(value) for value in values if value not in (None, "")]
        key = tuple(cell.casefold() for cell in cells)
        if key in seen:
            continue
        seen.add(key)
        fh.write(f"{label} R{row_no}: " + " | ".join(cells) + "\n")


def dump_xlsx_headers(path, fh, max_rows, max_cols):
    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as error:
        fh.write(f"OPEN FAIL {path} {error}\n")
        return
    try:
        for name in book.sheetnames:
            sheet = book[name]
            width = min(sheet.max_column or 0, max_cols) or max_cols
            found = []
            for row_no, row in enumerate(sheet.iter_rows(max_row=max_rows, max_col=width, values_only=True), start=1):
                values = list(row)
                if _header_score(values) >= 2:
                    found.append((row_no, values))
            write_header_rows(fh, f"XLSX {path} | {name!r}", found)
    finally:
        book.close()


def dump_xls_headers(path, fh, max_rows, max_cols):
    try:
        book = xlrd.open_workbook(str(path), formatting_info=False)
    except Exception as error:
        fh.write(f"OPEN FAIL {path} {error}\n")
        return
    for name in book.sheet_names():
        sheet = book.sheet_by_name(name)
        width = min(sheet.ncols, max_cols)
        found = []
        for row_no in range(min(sheet.nrows, max_rows)):
            values = [sheet.cell_value(row_no, col) for col in range(width)]
            if _header_score(values) >= 2:
                found.append((row_no + 1, values))
        write_header_rows(fh, f"XLS {path} | {name!r}", found)


def dump_pdf_headers(path, fh, pages):
    try:
        reader = PdfReader(str(path))
    except Exception as error:
        fh.write(f"OPEN FAIL {path} {error}\n")
        return
    total = len(reader.pages)
    chosen = list(range(min(pages, total)))
    if total > pages:
        chosen.append(total - 1)
    hits = []
    for index in chosen:
        text = reader.pages[index].extract_text() or ""
        for raw in text.splitlines():
            line = _norm_header(raw)
            if not line:
                continue
            folded = line.casefold()
            if any(hint in folded for hint in HEADER_HINTS) and len(line) <= 180:
                hits.append(f"p{index + 1}: {line}")
    seen = []
    for line in hits:
        if line not in seen:
            seen.append(line)
    if seen:
        fh.write(f"PDF {path} pages={total}\n")
        for line in seen[:40]:
            fh.write(f"  {line}\n")


def dump_pdf(path, fh, pages):
    fh.write(f"\n#### PDF {path}\n")
    try:
        reader = PdfReader(str(path))
    except Exception as error:
        fh.write(f"OPEN FAIL {error}\n")
        return
    total = len(reader.pages)
    fh.write(f"pages {total}\n")
    chosen = list(range(min(pages, total)))
    if total > pages:
        chosen.append(total - 1)
    seen = set()
    for index in chosen:
        if index in seen:
            continue
        seen.add(index)
        text = reader.pages[index].extract_text() or ""
        fh.write(f"\n--- page {index + 1}/{total} ---\n")
        if len(text) > 4000:
            text = text[:4000] + "\n...[cut]\n"
        fh.write(text + "\n")


def collect_files(roots, no_pdf, unzip_dir):
    excel = []
    pdfs = []
    notes = []
    pool = []
    for root in roots:
        pool.extend([root] if root.is_file() else list(root.rglob("*")))
    for path in list(pool):
        if path.suffix.lower() != ".zip":
            continue
        dest = unzip_dir / path.parent.name
        dest.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(path) as archive:
            archive.extractall(dest)
    if unzip_dir.exists():
        pool.extend(list(unzip_dir.rglob("*")))
    seen = set()
    for path in pool:
        if not path.is_file() or path.name.startswith("~$"):
            continue
        key = str(path.resolve()).lower()
        if key in seen:
            continue
        seen.add(key)
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm", ".xls"}:
            excel.append(path)
        elif suffix == ".pdf" and not no_pdf:
            pdfs.append(path)
        elif suffix == ".txt":
            notes.append(path)
    excel.sort(key=lambda item: str(item).lower())
    pdfs.sort(key=lambda item: str(item).lower())
    notes.sort(key=lambda item: str(item).lower())
    return excel, pdfs, notes


def dump_headers(roots, excel, pdfs, notes, out, max_rows, max_cols, pdf_pages):
    with out.open("w", encoding="utf-8") as handle:
        handle.write("HEADER SCAN " + " | ".join(str(root) for root in roots) + "\n")
        for path in notes:
            handle.write(f"\n===== NOTE {path} =====\n")
            handle.write(read_note(path) + "\n")
        handle.write("\n===== EXCEL HEADERS =====\n")
        for path in excel:
            print("headers", path.name.encode("ascii", "replace").decode("ascii"))
            try:
                if path.suffix.lower() == ".xls":
                    dump_xls_headers(path, handle, max_rows, max_cols)
                else:
                    dump_xlsx_headers(path, handle, max_rows, max_cols)
            except Exception as error:
                handle.write(f"OPEN FAIL {path} {error}\n")
        handle.write("\n===== PDF HEADER LINES =====\n")
        for path in pdfs:
            print("pdf-headers", path.name.encode("ascii", "replace").decode("ascii"))
            try:
                dump_pdf_headers(path, handle, pdf_pages)
            except Exception as error:
                handle.write(f"OPEN FAIL {path} {error}\n")
    print("wrote", out, "bytes", out.stat().st_size)


def main():
    parser = argparse.ArgumentParser(description="Снимок Excel и PDF одной папки поставки")
    parser.add_argument("path", nargs="+", help="Папка поставки или файл")
    parser.add_argument("--out", default=str(DEFAULT_OUT), help="Куда писать снимок, файл перезаписывается")
    parser.add_argument("--rows", type=int, default=40)
    parser.add_argument("--cols", type=int, default=30)
    parser.add_argument("--cell", type=int, default=80, help="Сколько знаков клетки оставить")
    parser.add_argument("--pdf-pages", type=int, default=2)
    parser.add_argument("--no-pdf", action="store_true")
    parser.add_argument("--headers", action="store_true", help="Только шапки таблиц, без полной сетки")
    args = parser.parse_args()

    roots = [Path(item).resolve() for item in args.path]
    out = Path(args.out).resolve()
    unzip_dir = out.parent / "unzip"
    out.parent.mkdir(parents=True, exist_ok=True)
    if unzip_dir.exists():
        for old in unzip_dir.rglob("*"):
            if old.is_file():
                old.unlink()

    excel, pdfs, notes = collect_files(roots, args.no_pdf, unzip_dir)
    if args.headers:
        dump_headers(roots, excel, pdfs, notes, out, args.rows, args.cols, args.pdf_pages)
        return

    with out.open("w", encoding="utf-8") as handle:
        handle.write("ROOT " + " | ".join(str(root) for root in roots) + "\n")
        for path in notes:
            handle.write(f"\n===== NOTE {path} =====\n")
            handle.write(read_note(path) + "\n")
        for path in excel:
            print("dump", path.name.encode("ascii", "replace").decode("ascii"))
            if path.suffix.lower() == ".xls":
                dump_xls(path, handle, args.rows, args.cols, args.cell)
            else:
                dump_xlsx(path, handle, args.rows, args.cols, args.cell)
        for path in pdfs:
            print("pdf", path.name.encode("ascii", "replace").decode("ascii"))
            dump_pdf(path, handle, args.pdf_pages)
    print("wrote", out, "bytes", out.stat().st_size)


if __name__ == "__main__":
    main()
