# -*- coding: utf-8 -*-
from pathlib import Path
from openpyxl import load_workbook

OUT = Path("study_tools/out/moi_dump.txt")
lines = []


def log(*parts):
    lines.append(" ".join(str(p) for p in parts))


def pack_sum(ws, pkg_col):
    total = 0.0
    n = 0
    empties = 0
    dashes = 0
    for r in range(1, (ws.max_row or 0) + 1):
        v = ws.cell(r, pkg_col).value
        if v is None or v == "":
            empties += 1
            continue
        if str(v).strip() in {"-", "—", "–"}:
            dashes += 1
            continue
        try:
            total += float(v)
            n += 1
        except (TypeError, ValueError):
            if str(v).upper().startswith("TOTAL") or "ИТОГ" in str(v).upper():
                log("  TOTAL_ROW", r, [ws.cell(r, c).value for c in range(1, 12)])
    return total, n, empties, dashes


def dump_book(path: Path):
    log("===", path.as_posix())
    wb = load_workbook(path, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        log("SHEET", sn, "rows", ws.max_row, "cols", ws.max_column)
        # header-ish
        for r in range(1, min(35, (ws.max_row or 1) + 1)):
            vals = [ws.cell(r, c).value for c in range(1, min(14, (ws.max_column or 1) + 1))]
            blob = " | ".join(str(v) for v in vals if v not in (None, ""))
            if any(
                k in blob.upper()
                for k in ("PACKAGE", "NO.", "№", "TOTAL", "ИТОГ", "CODE", "QTY", "МЕСТА")
            ):
                log(f"  R{r}", vals)
        # find PACKAGE col
        pkg_col = None
        for r in range(1, min(40, (ws.max_row or 1) + 1)):
            for c in range(1, (ws.max_column or 1) + 1):
                v = str(ws.cell(r, c).value or "")
                if "PACKAGE" in v.upper() or "МЕСТА" in v.upper() or "CARTON" in v.upper():
                    pkg_col = c
                    break
            if pkg_col:
                break
        if pkg_col:
            total, n, empties, dashes = pack_sum(ws, pkg_col)
            log(f"  PKG_COL {pkg_col} sum={total} filled={n} empty={empties} dashes={dashes}")
        for r in range(max(1, (ws.max_row or 1) - 2), (ws.max_row or 1) + 1):
            log(f"  TAIL{r}", [ws.cell(r, c).value for c in range(1, min(12, (ws.max_column or 1) + 1))])


root = Path("documents/7_pravka")
for folder in sorted(root.rglob("*")):
    if not folder.is_dir():
        continue
    if "мои" not in folder.name.casefold():
        continue
    if "261006" not in str(folder) and "LOGDV_7_8200" not in str(folder):
        continue
    log("\nFOLDER", folder.as_posix())
    for f in sorted(folder.glob("*.xlsx")):
        dump_book(f)

OUT.write_text("\n".join(lines), encoding="utf-8")
print("wrote", OUT, "lines", len(lines))
